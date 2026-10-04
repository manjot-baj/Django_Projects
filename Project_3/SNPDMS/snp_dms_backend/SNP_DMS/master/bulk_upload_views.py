import os
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework import views
from django.http import HttpResponse
from openpyxl import load_workbook
from SNP_DMS.settings.base import BASE_DIR
from master.models import Location, Site
from master.models_two import SealNo
from master.functions import extract_excel_data
import pandas as pd
from django.utils import timezone
import datetime
from depot.models import ContainerStock, GateOut
from django.db import transaction
from common.error_logging import ErrorLogging
from common.exceptions import ValidationError, ResourceNotFound, AlreadyExists
from rest_framework import status


class MasterSealNoUploadSampleFile(views.APIView):
    permission_classes = (IsAuthenticated,)

    def get(self, request, *args, **kwargs):
        try:
            temp_file_path = os.path.join(
                BASE_DIR, "sample_stock/sample_seal_no_upload_master.xlsx"
            )
            with open(temp_file_path, "rb") as temp:
                file_response = HttpResponse(
                    temp.read(), content_type=f"application/xlsx"
                )
                file_response["Content-Disposition"] = (
                    f'attachment; filename="sample_seal_no_upload_master.xlsx"'
                )
            return file_response
        except AlreadyExists as e:
            return Response(
                {"errorMsg": str(e.message)}, status=status.HTTP_400_BAD_REQUEST
            )
        except ValidationError as e:
            return Response(
                {"errorMsg": str(e.message)}, status=status.HTTP_400_BAD_REQUEST
            )
        except ResourceNotFound as e:
            return Response(
                {"errorMsg": str(e.message)}, status=status.HTTP_404_NOT_FOUND
            )
        except Exception as e:
            ErrorLogging().log_error()
            return Response(
                {"errorMsg": "An unexpected error occurred, please try again later."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )

    def post(self, request, *args, **kwargs):
        try:
            data = request.data["file"]
            result = extract_excel_data(input_excel=data)

            if result == "Header Not Found" or result is False:
                raise ValidationError("Data file is corrupted, unable to import data")
            else:
                return Response(result, status=200)
        except AlreadyExists as e:
            return Response(
                {"errorMsg": str(e.message)}, status=status.HTTP_400_BAD_REQUEST
            )
        except ValidationError as e:
            return Response(
                {"errorMsg": str(e.message)}, status=status.HTTP_400_BAD_REQUEST
            )
        except ResourceNotFound as e:
            return Response(
                {"errorMsg": str(e.message)}, status=status.HTTP_404_NOT_FOUND
            )
        except Exception as e:
            ErrorLogging().log_error()
            return Response(
                {"errorMsg": "An unexpected error occurred, please try again later."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )


class MasterSealNoImport(views.APIView):
    permission_classes = (IsAuthenticated,)

    def post(self, request, *args, **kwargs):

        try:
            data = request.data["importable_data"]
            site = Site.objects.get(name=request.data["site"])
            location = Location.objects.get(name=request.data["location"])
            with transaction.atomic():
                if not site or not location:
                    raise ResourceNotFound("Invalid Site or Location")

                for each in data:
                    number = each["number"]
                    seal_box_number = each["seal_box_number"]
                    line = each["line"]
                    in_date_str = each["in_date"]
                    in_time_str = each["in_time"]
                    in_date = datetime.datetime.strptime(in_date_str, "%Y_%m_%d").date()
                    in_time = datetime.datetime.strptime(in_time_str, "%H_%M").time()
                    in_date_time = datetime.datetime.combine(
                        in_date, in_time
                    ).astimezone(timezone.get_current_timezone())

                    if (
                        SealNo.objects.filter(number=number).exists()
                        or ContainerStock.objects.filter(seal_no=number).exists()
                        or GateOut.objects.filter(seal_no=number).exists()
                    ):
                        raise AlreadyExists("Seal NO already exists")

                    seal_no_entry = SealNo.create(
                        number=number,
                        seal_box_number=seal_box_number,
                        location=location,
                        site=site,
                        line=line,
                        in_date=in_date_time,
                    )
                    seal_no_entry.save()
            return Response(
                {"successMsg": "Seal No Imported"}, status=status.HTTP_201_CREATED
            )
        except AlreadyExists as e:
            return Response(
                {"errorMsg": str(e.message)}, status=status.HTTP_400_BAD_REQUEST
            )
        except ValidationError as e:
            return Response(
                {"errorMsg": str(e.message)}, status=status.HTTP_400_BAD_REQUEST
            )
        except ResourceNotFound as e:
            return Response(
                {"errorMsg": str(e.message)}, status=status.HTTP_404_NOT_FOUND
            )
        except Exception as e:
            ErrorLogging().log_error()
            return Response(
                {"errorMsg": "An unexpected error occurred, please try again later."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )


class MasterRejectedSealNoDataFile(views.APIView):
    """Post Function will download and return the rejected data in a xls file"""

    permission_classes = (IsAuthenticated,)

    def post(self, request, *args, **kwargs):
        try:
            data = request.data["rejected_data"]
            faults = request.data["faults"]
            seal_no = []
            seal_box_number = []
            line = []
            in_date = []
            in_time = []

            for each in data:
                seal_no.append(each["number"])
                seal_box_number.append(each["seal_box_number"])
                line.append(each["line"])
                in_date.append(each["in_date"])
                in_time.append(each["in_time"])

            seal_df_data = {
                "number": seal_no,
                "seal_box_number": seal_box_number,
                "line": line,
                "in_date": in_date,
                "in_time": in_time,
            }
            if not os.path.exists(os.path.join(BASE_DIR, "temp/sample_stock/")):
                os.makedirs(os.path.join(BASE_DIR, "temp/sample_stock/"))
            new_temp_file_path = os.path.join(
                BASE_DIR, f"temp/sample_stock/sample_seal_no_upload_master.xlsx"
            )
            temp_file_path = os.path.join(
                BASE_DIR, "sample_stock/sample_seal_no_upload_master.xlsx"
            )
            old_temp_data = None
            with open(temp_file_path, "rb") as temp:
                old_temp_data = temp.read()
            with open(new_temp_file_path, "wb") as f:
                f.write(old_temp_data)

            # New Version code
            book = load_workbook(new_temp_file_path)
            with pd.ExcelWriter(
                new_temp_file_path,
                engine="openpyxl",
                mode="a",
                if_sheet_exists="replace",
            ) as writer:

                seal_df = pd.DataFrame(seal_df_data)
                seal_df.to_excel(writer, sheet_name="seal_number", index=False)

                faults_df = pd.DataFrame(faults)
                faults_df.to_excel(writer, sheet_name="faults", index=False)

            # Old Version
            # book = load_workbook(new_temp_file_path)
            # writer = pd.ExcelWriter(new_temp_file_path, engine="openpyxl")
            # writer.book = book
            # writer.sheets = dict((ws.title, ws) for ws in book.worksheets)
            # seal_df = pd.DataFrame(seal_df_data)
            # faults_df = pd.DataFrame(faults)
            # seal_df.to_excel(
            #     writer, sheet_name="seal_number", startrow=0, startcol=0, index=False
            # )
            # faults_df.to_excel(
            #     writer, sheet_name="faults", startrow=0, startcol=0, index=False
            # )
            # fault_sheet = book.get_sheet_by_name("faults")
            # writer.save()

            with open(new_temp_file_path, "rb") as temp:
                file_response = HttpResponse(
                    temp.read(), content_type=f"application/xlsx"
                )
                file_response["Content-Disposition"] = (
                    f'attachment; filename="rejected_sample_seal_no_upload_master.xlsx"'
                )
                os.remove(new_temp_file_path)
            return file_response
        except AlreadyExists as e:
            return Response(
                {"errorMsg": str(e.message)}, status=status.HTTP_400_BAD_REQUEST
            )
        except ValidationError as e:
            return Response(
                {"errorMsg": str(e.message)}, status=status.HTTP_400_BAD_REQUEST
            )
        except ResourceNotFound as e:
            return Response(
                {"errorMsg": str(e.message)}, status=status.HTTP_404_NOT_FOUND
            )
        except Exception as e:
            ErrorLogging().log_error()
            return Response(
                {"errorMsg": "An unexpected error occurred, please try again later."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )
