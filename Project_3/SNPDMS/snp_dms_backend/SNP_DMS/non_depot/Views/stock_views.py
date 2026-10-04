import os
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework import views
from django.http import HttpResponse
from openpyxl import load_workbook
from SNP_DMS.settings.base import BASE_DIR
from account.models import AccountUser
from master.models import Location, Site, ContainerSize, ContainerType
from master.models_two import Client, ClientAbbreviation
from non_depot.functions import (
    extract_excel_data,
    create_non_depot_stock_sheet,
    stock_sheet_df,
)
from depot.functions_two import *
import pandas as pd
from django.utils import timezone
from non_depot.models import (
    NonDepotContainer,
    NonDepotContainerStock,
    NonDepotGateIn,
    NonDepotContainerInOutRecord,
)
import datetime
import traceback, logging


class NonDepotStockUploadSampleFile(views.APIView):
    permission_classes = (IsAuthenticated,)

    def get(self, request, *args, **kwargs):
        try:
            temp_file_path = os.path.join(
                BASE_DIR, "sample_stock/sample_stock_upload_non_depot.xlsx"
            )
            with open(temp_file_path, "rb") as temp:
                file_response = HttpResponse(
                    temp.read(), content_type=f"application/xlsx"
                )
                file_response["Content-Disposition"] = (
                    f'attachment; filename="sample_stock_upload_non_depot.xlsx"'
                )
            return file_response

        except:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return Response({"errorMsg": f"Data Not Found"}, status=200)

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data = request.data["file"]
            user = request.user
            app_user = AccountUser.objects.get(username=user.username)
            try:
                location_str = request.data["location"]
                site_str = request.data["site"]
                location = Location.objects.get(name=location_str)
                site = Site.objects.get(name=site_str)
            except:
                location = app_user.location
                site = app_user.site
            result = extract_excel_data(input_excel=data, location=location, site=site)

            if result == "Header Not Found" or result is False:
                return Response(
                    {"errorMsg": "Data file is corrupted, unable to import data"},
                    status=200,
                )
            else:
                return Response(result, status=200)
        except Exception as e:
            return Response({"errorMsg": f"Invalid Data Provided [ {e} ]"}, status=200)


class NonDepotStockImport(views.APIView):
    permission_classes = (IsAuthenticated,)

    def post(self, request, *args, **kwargs):

        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data = request.data["importable_data"]

            user = request.user
            app_user = AccountUser.objects.get(username=user.username)
            try:
                location_str = request.data["location"]
                site_str = request.data["site"]
                location = Location.objects.get(name=location_str)
                site = Site.objects.get(name=site_str)
            except:
                location = app_user.location
                site = app_user.site

            automatic_mnr_status_change = site.automatic_mnr_status_change

            for each in data:
                client_str = each["client"]
                shipping_line_str = each["shipping_line"]
                type_str = each["type"]
                size_str = each["size"]
                container_no_str = each["container_no"]
                if each["gross_wt"] == "":
                    gross_wt_str = 0
                else:
                    gross_wt_str = str(int(float(each["gross_wt"])))
                if each["tare_wt"] == "":
                    tare_wt_str = 0
                else:
                    tare_wt_str = str(int(float(each["tare_wt"])))
                manufacturing_date_str = each["manufacturing_date"]
                in_date_str = each["in_date"]
                in_time_str = each["in_time"]
                condition_str = each["condition"]
                grade_str = each["grade"]
                dock_destuff_str = each["dock_destuff"]
                mode_str = each["mode"]
                client_object = Client.objects.get(
                    name=client_str, location=location, site=site
                )
                shipping_line = list(
                    ClientAbbreviation.objects.filter(
                        client=client_object, name=shipping_line_str
                    )
                )[0]
                type_object = ContainerType.objects.get(name=type_str)
                size_object = ContainerSize.objects.get(name=str(int(float(size_str))))
                manufacturing_date = datetime.datetime.strptime(
                    manufacturing_date_str, "%d_%m_%Y"
                ).date()
                in_date = datetime.datetime.strptime(in_date_str, "%d_%m_%Y").date()
                in_time = datetime.datetime.strptime(in_time_str, "%H_%M").time()

                # Container Check
                try:
                    container_object = NonDepotContainer.objects.get(
                        container_no=container_no_str, location=location, site=site
                    )

                    if container_object.status == "IN":
                        return Response(
                            {"errorMsg": "Container Already In"}, status=200
                        )
                    else:
                        container_object.status = "IN"
                        container_object.is_available = False
                        container_object.save()
                except:
                    container_no_response = check_digit(container_no=container_no_str)

                    #     Adding Non Depot Container
                    container_object = NonDepotContainer.create(
                        client=client_object,
                        type_object=type_object,
                        size_object=size_object,
                        container_no=container_no_str,
                        payload=None,
                        gross_wt=gross_wt_str,
                        tare_wt=tare_wt_str,
                        manufacturing_date=manufacturing_date,
                        shipping_line=shipping_line,
                        dock_destuff=dock_destuff_str,
                        mode=mode_str,
                        condition=condition_str,
                        grade=grade_str,
                        automatic_mnr_status_change=automatic_mnr_status_change,
                    )
                    container_object.save()
                    container_object.location = location
                    container_object.site = site
                    container_object.save()

                    # Adding Non Depot GateIn
                gate_in_object = NonDepotGateIn.create(
                    container=container_object,
                    in_date=in_date,
                    in_time=in_time,
                )
                gate_in_object.save()

                # Adding Non Depot Container Stock
                NonDepotContainerStock.create(
                    gate_in=gate_in_object, container=container_object
                ).save()

                # Adding Non Depot Container InOut Record
                container_in_out_record = NonDepotContainerInOutRecord(
                    container=container_object, gate_in=gate_in_object, gate_out=None
                )
                container_in_out_record.save()

            return Response({"successMsg": "Stock Imported"}, status=200)
        except Exception as e:
            return Response({"errorMsg": f"Invalid Data Provided [ {e} ]"}, status=200)


class NonDepotRejectedStockDataFile(views.APIView):
    """Post Function will download and return the rejected data in a xls file"""

    permission_classes = (IsAuthenticated,)

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data = request.data["rejected_data"]
            faults = request.data["faults"]
            shipping_line = []
            client = []
            c_type = []
            c_size = []
            container_no = []
            gross_wt = []
            tare_wt = []
            manufacturing_date = []
            in_date = []
            in_time = []
            condition = []
            grade = []
            mode = []
            dock_destuff = []

            for each in data:
                shipping_line.append(each["shipping_line"])
                client.append(each["client"])
                c_type.append(each["type"])
                c_size.append(each["size"])
                container_no.append(each["container_no"])
                if each["gross_wt"] == "":
                    each["gross_wt"] = "_"
                    gross_wt.append(each["gross_wt"])
                else:
                    gross_wt.append(each["gross_wt"])
                if each["tare_wt"] == "":
                    each["tare_wt"] = "_"
                    tare_wt.append(each["tare_wt"])
                else:
                    tare_wt.append(each["tare_wt"])
                manufacturing_date.append(each["manufacturing_date"])
                in_date.append(each["in_date"])
                in_time.append(each["in_time"])
                condition.append(each["condition"])
                if each["grade"] == "":
                    each["grade"] = "_"
                    grade.append(each["grade"])
                else:
                    grade.append(each["grade"])
                mode.append(each["mode"])
                if each["dock_destuff"] == "":
                    each["dock_destuff"] = "_"
                    dock_destuff.append(each["dock_destuff"])
                else:
                    dock_destuff.append(each["dock_destuff"])

            stock_df_data = {
                "shipping_line": shipping_line,
                "client": client,
                "type": c_type,
                "size": c_size,
                "container_no": container_no,
                "gross_wt": gross_wt,
                "tare_wt": tare_wt,
                "manufacturing_date": manufacturing_date,
                "in_date": in_date,
                "in_time": in_time,
                "condition": condition,
                "grade": grade,
                "mode": mode,
                "dock_destuff": dock_destuff,
            }
            if not os.path.exists(os.path.join(BASE_DIR, "temp/sample_stock/")):
                os.makedirs(os.path.join(BASE_DIR, "temp/sample_stock/"))
            new_temp_file_path = os.path.join(
                BASE_DIR, f"temp/sample_stock/sample_stock_upload_non_depot.xlsx"
            )
            temp_file_path = os.path.join(
                BASE_DIR, "sample_stock/sample_stock_upload_non_depot.xlsx"
            )
            old_temp_data = None
            with open(temp_file_path, "rb") as temp:
                old_temp_data = temp.read()
            with open(new_temp_file_path, "wb") as f:
                f.write(old_temp_data)

            # New Version code
            book = load_workbook(new_temp_file_path)
            with pd.ExcelWriter(
                new_temp_file_path,
                engine="openpyxl",
                mode="a",
                if_sheet_exists="replace",
            ) as writer:

                stock_df = pd.DataFrame(stock_df_data)
                stock_df.to_excel(writer, sheet_name="stock", index=False)

                faults_df = pd.DataFrame(faults)
                faults_df.to_excel(writer, sheet_name="faults", index=False)

            # Old Version
            # book = load_workbook(new_temp_file_path)
            # writer = pd.ExcelWriter(new_temp_file_path, engine="openpyxl")
            # writer.book = book
            # writer.sheets = dict((ws.title, ws) for ws in book.worksheets)
            # stock_df = pd.DataFrame(stock_df_data)
            # faults_df = pd.DataFrame(faults)
            # stock_df.to_excel(
            #     writer, sheet_name="stock", startrow=0, startcol=0, index=False
            # )
            # faults_df.to_excel(
            #     writer, sheet_name="faults", startrow=0, startcol=0, index=False
            # )
            # fault_sheet = book.get_sheet_by_name("faults")
            # writer.save()

            with open(new_temp_file_path, "rb") as temp:
                file_response = HttpResponse(
                    temp.read(), content_type=f"application/xlsx"
                )
                file_response["Content-Disposition"] = (
                    f'attachment; filename="rejected_sample_stock_upload_non_depot.xlsx"'
                )
                os.remove(new_temp_file_path)
            return file_response
        except Exception as e:
            return Response({"errorMsg": f"Invalid Data Provided [ {e} ]"}, status=200)


def get_merge_format(workbook):
    merge_format = workbook.add_format(
        {
            "bold": 2,
            "align": "center",
            "valign": "vcenter",
            "font_color": "#e0a92a",
            "font_size": 18,
        }
    )
    merge_format3 = workbook.add_format(
        {
            "bold": 1,
            "font_size": 11,
        }
    )
    return merge_format, merge_format3


class StockSheetDownload(views.APIView):

    permission_classes = (IsAuthenticated,)

    def params(self, data):
        site = Site.objects.get(name=data["site"])
        location = Location.objects.get(name=data["location"])
        parameters = {
            "container__location": location,
            "container__site": site,
        }
        if data["ref_code"]:
            parameters["container__client__ref_code"] = data["ref_code"]
        if data["client"]:
            parameters["container__client__name"] = data["client"]

        if data["container_no"]:
            parameters["container__container_no__in"] = data["container_no"]

        if data["stage"]:
            parameters["stage"] = data["stage"]

        if data["from_date"] and data["to_date"]:
            parameters["gate_in__in_date__range"] = [
                data["from_date"],
                data["to_date"],
            ]

        if data["ref_code"]:
            parameters["container__client__ref_code"] = data["ref_code"]

        if "status" in data:
            if data["status"]:
                parameters["status"] = data["status"]

        if data["out_history"] == "True":
            parameters["container_status"] = "OUT"
        else:
            parameters["container_status"] = "IN"

        return location, site, parameters

    def post(self, request, *args, **kwargs):

        try:

            location, site, param = self.params(data=request.data)

            dt = datetime.datetime.now().astimezone(timezone.get_current_timezone())
            date = dt.date().strftime("%Y%m%d")
            time = dt.time().strftime("%H%M%S")

            queryset = (
                NonDepotContainerStock.objects.select_related(
                    "container",
                    "container__client",
                    "container__type",
                    "container__size",
                    "gate_in",
                )
                .filter(**param)
                .values(
                    "container__client__name",
                    "container__client__ref_code",
                    "gate_in__in_date",
                    "container__container_no",
                    "container__size__name",
                    "container__type__name",
                    "stage",
                    "estimate_status",
                    "status",
                    "container__condition",
                    "is_estimate_westim_sent",
                    "is_repair_destim_sent",
                )
                .order_by("-gate_in__in_date")
            )

            df_data = stock_sheet_df(queryset)

            temp_file_path = create_non_depot_stock_sheet(
                location, site, df_data, date, time
            )
            with open(temp_file_path, "rb") as temp:
                file_response = HttpResponse(
                    temp.read(), content_type=f"application/xlsx"
                )
                file_response["Content-Disposition"] = (
                    f'attachment; filename="non_depot_stock_sheet_{date}_{time}.xlsx"'
                )
                os.remove(temp_file_path)
            return file_response

        except:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return None
