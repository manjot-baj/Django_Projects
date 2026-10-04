from rest_framework import views
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from ..models import *
from master.models_two import Client, ClientAbbreviation, SealNo
from master.models import ContainerType, ContainerSize, Location, Transporter, Site
from account.models import AccountUser
import datetime
from ..functions import (
    extract_excel_data,
    generic_stock_sheet_main_data,
    create_generic_stock_sheet_wb,
    create_allotment_track_record,
)
from ..filter_functions import *
from ..functions_two import *
from django.http import HttpResponse
from SNP_DMS.settings.base import BASE_DIR
import os
import pandas as pd
from openpyxl import load_workbook
from django.utils import timezone
from django.core.paginator import Paginator
from report.functions2 import user_data
from edi.cycle_functions import zim_repair_edi_mail_process
from account.permissions import HasAllowedRoles
from django.db import transaction


class StockOfContainer(views.APIView):
    """
    the Post function will give you all container in stock
    and put function will update the data
    """

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data = request.data
            pg_no = request.data["pg_no"]
            on_page_data = request.data["on_page_data"]
            location = request.data["location"]
            site = request.data["site"]
            ref_code = data.get("ref_code", None)
            queued_recently = data.get("queued_recently", None)
            queryset = None

            if data["out_history"] == "True":
                if location == "ALL":
                    queryset = (
                        ContainerStock.objects.select_related(
                            "container",
                            "container__client",
                            "container__type",
                            "container__size",
                            "gate_in",
                        )
                        .filter(container_status="OUT")
                        .order_by("-gate_in__in_date")
                    )

                elif site == "ALL":
                    queryset = (
                        ContainerStock.objects.select_related(
                            "container",
                            "container__client",
                            "container__type",
                            "container__size",
                            "container__location",
                            "gate_in",
                        )
                        .filter(
                            container_status="OUT", container__location__name=location
                        )
                        .order_by("-gate_in__in_date")
                    )

                else:
                    queryset = (
                        ContainerStock.objects.select_related(
                            "container",
                            "container__client",
                            "container__type",
                            "container__size",
                            "container__location",
                            "container__site",
                            "gate_in",
                        )
                        .filter(
                            container_status="OUT",
                            container__location__name=location,
                            container__site__name=site,
                        )
                        .order_by("-gate_in__in_date")
                    )
            else:
                if data["do_not_lift_queue"] == "True":
                    if location == "ALL":
                        queryset = (
                            ContainerStock.objects.select_related(
                                "container",
                                "container__client",
                                "container__type",
                                "container__size",
                                "gate_in",
                            )
                            .filter(
                                container_status="IN",
                                container__in_do_not_lift_queue=True,
                            )
                            .order_by("-gate_in__in_date")
                        )
                    elif site == "ALL":
                        queryset = (
                            ContainerStock.objects.select_related(
                                "container",
                                "container__client",
                                "container__type",
                                "container__size",
                                "container__location",
                                "gate_in",
                            )
                            .filter(
                                container_status="IN",
                                container__location__name=location,
                                container__in_do_not_lift_queue=True,
                            )
                            .order_by("-gate_in__in_date")
                        )
                    else:
                        queryset = (
                            ContainerStock.objects.select_related(
                                "container",
                                "container__client",
                                "container__type",
                                "container__size",
                                "container__location",
                                "container__site",
                                "gate_in",
                            )
                            .filter(
                                container_status="IN",
                                container__location__name=location,
                                container__site__name=site,
                                container__in_do_not_lift_queue=True,
                            )
                            .order_by("-gate_in__in_date")
                        )
                else:
                    if location == "ALL":
                        queryset = (
                            ContainerStock.objects.select_related(
                                "container",
                                "container__client",
                                "container__type",
                                "container__size",
                                "gate_in",
                            )
                            .filter(
                                container_status="IN",
                                container__in_do_not_lift_queue=False,
                            )
                            .order_by("-gate_in__in_date")
                        )
                        if queued_recently is True:
                            queryset = queryset.filter(
                                container__queued_recently=queued_recently
                            )

                    elif site == "ALL":
                        queryset = (
                            ContainerStock.objects.select_related(
                                "container",
                                "container__client",
                                "container__type",
                                "container__size",
                                "container__location",
                                "gate_in",
                            )
                            .filter(
                                container_status="IN",
                                container__location__name=location,
                                container__in_do_not_lift_queue=False,
                            )
                            .order_by("-gate_in__in_date")
                        )
                        if queued_recently is True:
                            queryset = queryset.filter(
                                container__queued_recently=queued_recently
                            )
                    else:
                        queryset = (
                            ContainerStock.objects.select_related(
                                "container",
                                "container__client",
                                "container__type",
                                "container__size",
                                "container__location",
                                "container__site",
                                "gate_in",
                            )
                            .filter(
                                container_status="IN",
                                container__location__name=location,
                                container__site__name=site,
                                container__in_do_not_lift_queue=False,
                            )
                            .order_by("-gate_in__in_date")
                        )
                        if queued_recently is True:
                            queryset = queryset.filter(
                                container__queued_recently=queued_recently
                            )

            if ref_code is not None:
                if not len(ref_code) == 0:
                    queryset = queryset.filter(container__client__ref_code=ref_code)

            if not len(data["client"]) == 0:
                queryset = queryset.filter(container__client__name=data["client"])

            if not len(data["container_no"]) == 0:
                queryset = queryset.filter(
                    container__container_no__in=data["container_no"]
                )

            if not len(data["booking_no"]) == 0:
                queryset = queryset.filter(booking_no=data["booking_no"])

            if not len(data["status"]) == 0:
                queryset = queryset.filter(status=data["status"])

            gate_in_date = data["gate_in_date"]
            if not len(gate_in_date["from"]) == 0 and not len(gate_in_date["to"]) == 0:
                queryset = queryset.filter(
                    gate_in__in_date__range=[gate_in_date["from"], gate_in_date["to"]]
                )

            gate_out_date = data["gate_out_date"]
            if (
                not len(gate_out_date["from"]) == 0
                and not len(gate_out_date["to"]) == 0
            ):
                queryset = queryset.filter(
                    gate_out__out_date__range=[
                        gate_out_date["from"],
                        gate_out_date["to"],
                    ]
                )

            available_date = data["available_date"]
            if (
                not len(available_date["from"]) == 0
                and not len(available_date["to"]) == 0
            ):
                queryset = queryset.filter(
                    available_date__range=[available_date["from"], available_date["to"]]
                )

            allotment_date = data["allotment_date"]
            if (
                not len(allotment_date["from"]) == 0
                and not len(allotment_date["to"]) == 0
            ):
                queryset = queryset.filter(
                    allotment_date__range=[allotment_date["from"], allotment_date["to"]]
                )

            # Pagination
            paginator = Paginator(queryset, on_page_data)
            no_of_data_count = paginator.count
            no_of_pages = paginator.num_pages
            current_page = paginator.page(pg_no)
            on_page_data_count = (
                current_page.end_index() - current_page.start_index() + 1
            )
            prev_page = ""
            if current_page.has_previous():
                prev_page = current_page.previous_page_number()
            next_page = ""
            if current_page.has_next():
                next_page = current_page.next_page_number()

            response_data = []
            count = 0
            for each in current_page:
                data_dict = each.get_stock()
                data_dict["sr_no"] = current_page.start_index() + count
                response_data.append(data_dict)
                count += 1

            return Response(
                {
                    "no_of_data": no_of_data_count,
                    "on_page_data": on_page_data_count,
                    "total_pages": no_of_pages,
                    "prev_page": prev_page,
                    "next_page": next_page,
                    "data": response_data,
                },
                status=200,
            )
        except Exception as e:
            return Response({"errorMsg": f"Data Not Found [ {e} ]"}, status=200)

    def put(self, request, pk, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data = request.data
            stock_object = ContainerStock.objects.get(pk=pk)
            allotment_removed = False
            if len(data["grade"]) == 0:
                grade = None
            else:
                grade = data["grade"]
            if len(data["status"]) == 0:
                status = None
            else:
                status = data["status"]

            available_date_str = data["available_date"]
            if len(available_date_str) == 0:
                available_date = None
            else:
                try:
                    available_date = datetime.datetime.strptime(
                        available_date_str, "%d/%m/%Y"
                    ).date()
                except:
                    available_date = None

            if not stock_object.status == status and status is not None:
                if stock_object.status == "Survey Pending":
                    stock_object.survey_pending_out_date_time = (
                        datetime.datetime.now().astimezone(
                            timezone.get_current_timezone()
                        )
                    )
                    stock_object.save()
                    if status == "Estimate Pending":
                        stock_object.estimate_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Approval Pending":
                        stock_object.approval_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Approved":
                        stock_object.approved_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Under Repairing":
                        stock_object.under_repair_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Empty Alloted":
                        stock_object.empty_allotment_date = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        ).date()
                        stock_object.empty_allotment_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status in ["Available", "Without_Repair_Available"]:
                        stock_object.available_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    else:
                        pass
                elif stock_object.status == "Estimate Pending":
                    if status == "Survey Pending":
                        stock_object.estimate_pending_in_date_time = None
                        stock_object.survey_pending_out_date_time = None
                        stock_object.survey_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Approval Pending":
                        stock_object.estimate_pending_out_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.approval_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Approved":
                        stock_object.estimate_pending_out_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.approved_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Under Repairing":
                        stock_object.estimate_pending_out_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.under_repair_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Empty Alloted":
                        stock_object.estimate_pending_out_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.empty_allotment_date = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        ).date()
                        stock_object.empty_allotment_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status in ["Available", "Without_Repair_Available"]:
                        stock_object.estimate_pending_out_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.available_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    else:
                        pass
                elif stock_object.status == "Approval Pending":
                    if status == "Survey Pending":
                        stock_object.approval_pending_in_date_time = None
                        stock_object.estimate_pending_out_date_time = None
                        stock_object.estimate_pending_in_date_time = None
                        stock_object.survey_pending_out_date_time = None
                        stock_object.survey_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Estimate Pending":
                        stock_object.approval_pending_in_date_time = None
                        stock_object.estimate_pending_out_date_time = None
                        stock_object.estimate_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Approved":
                        stock_object.approval_pending_out_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.approved_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Under Repairing":
                        stock_object.approval_pending_out_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.under_repair_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Empty Alloted":
                        stock_object.approval_pending_out_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.empty_allotment_date = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        ).date()
                        stock_object.empty_allotment_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status in ["Available", "Without_Repair_Available"]:
                        stock_object.approval_pending_out_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.available_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    else:
                        pass
                elif stock_object.status == "Approved":
                    if status == "Survey Pending":
                        stock_object.approved_in_date_time = None
                        stock_object.approval_pending_out_date_time = None
                        stock_object.approval_pending_in_date_time = None
                        stock_object.estimate_pending_out_date_time = None
                        stock_object.estimate_pending_in_date_time = None
                        stock_object.survey_pending_out_date_time = None
                        stock_object.survey_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Estimate Pending":
                        stock_object.approved_in_date_time = None
                        stock_object.approval_pending_out_date_time = None
                        stock_object.approval_pending_in_date_time = None
                        stock_object.estimate_pending_out_date_time = None
                        stock_object.estimate_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Approval Pending":
                        stock_object.approved_in_date_time = None
                        stock_object.approval_pending_out_date_time = None
                        stock_object.approval_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Under Repairing":
                        stock_object.approved_out_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.under_repair_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Empty Alloted":
                        stock_object.approved_out_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.empty_allotment_date = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        ).date()
                        stock_object.empty_allotment_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status in ["Available", "Without_Repair_Available"]:
                        stock_object.approved_out_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.available_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    else:
                        pass
                elif stock_object.status == "Under Repairing":
                    if status == "Survey Pending":
                        stock_object.under_repair_in_date_time = None
                        stock_object.approved_out_date_time = None
                        stock_object.approved_in_date_time = None
                        stock_object.approval_pending_out_date_time = None
                        stock_object.approval_pending_in_date_time = None
                        stock_object.estimate_pending_out_date_time = None
                        stock_object.estimate_pending_in_date_time = None
                        stock_object.survey_pending_out_date_time = None
                        stock_object.survey_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Estimate Pending":
                        stock_object.under_repair_in_date_time = None
                        stock_object.approved_out_date_time = None
                        stock_object.approved_in_date_time = None
                        stock_object.approval_pending_out_date_time = None
                        stock_object.approval_pending_in_date_time = None
                        stock_object.estimate_pending_out_date_time = None
                        stock_object.estimate_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Approval Pending":
                        stock_object.under_repair_in_date_time = None
                        stock_object.approved_out_date_time = None
                        stock_object.approved_in_date_time = None
                        stock_object.approval_pending_out_date_time = None
                        stock_object.approval_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Approved":
                        stock_object.under_repair_in_date_time = None
                        stock_object.approved_out_date_time = None
                        stock_object.approved_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Empty Alloted":
                        stock_object.under_repair_out_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.empty_allotment_date = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        ).date()
                        stock_object.empty_allotment_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status in ["Available", "Without_Repair_Available"]:
                        stock_object.under_repair_out_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.available_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    else:
                        pass
                elif stock_object.status == "Empty Alloted":
                    if status == "Survey Pending":
                        stock_object.empty_allotment_date = None
                        stock_object.empty_allotment_in_date_time = None
                        stock_object.under_repair_out_date_time = None
                        stock_object.under_repair_in_date_time = None
                        stock_object.approved_out_date_time = None
                        stock_object.approved_in_date_time = None
                        stock_object.approval_pending_out_date_time = None
                        stock_object.approval_pending_in_date_time = None
                        stock_object.estimate_pending_out_date_time = None
                        stock_object.estimate_pending_in_date_time = None
                        stock_object.survey_pending_out_date_time = None
                        stock_object.survey_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Estimate Pending":
                        stock_object.empty_allotment_date = None
                        stock_object.empty_allotment_in_date_time = None
                        stock_object.under_repair_out_date_time = None
                        stock_object.under_repair_in_date_time = None
                        stock_object.approved_out_date_time = None
                        stock_object.approved_in_date_time = None
                        stock_object.approval_pending_out_date_time = None
                        stock_object.approval_pending_in_date_time = None
                        stock_object.estimate_pending_out_date_time = None
                        stock_object.estimate_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Approval Pending":
                        stock_object.empty_allotment_date = None
                        stock_object.empty_allotment_in_date_time = None
                        stock_object.under_repair_out_date_time = None
                        stock_object.under_repair_in_date_time = None
                        stock_object.approved_out_date_time = None
                        stock_object.approved_in_date_time = None
                        stock_object.approval_pending_out_date_time = None
                        stock_object.approval_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Approved":
                        stock_object.empty_allotment_date = None
                        stock_object.empty_allotment_in_date_time = None
                        stock_object.under_repair_out_date_time = None
                        stock_object.under_repair_in_date_time = None
                        stock_object.approved_out_date_time = None
                        stock_object.approved_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Under Repairing":
                        stock_object.empty_allotment_date = None
                        stock_object.empty_allotment_in_date_time = None
                        stock_object.under_repair_out_date_time = None
                        stock_object.under_repair_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status in ["Available", "Without_Repair_Available"]:
                        stock_object.empty_allotment_out_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.available_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    else:
                        pass
                elif stock_object.status in ["Available", "Without_Repair_Available"]:
                    if status == "Survey Pending":
                        stock_object.available_in_date_time = None
                        stock_object.empty_allotment_out_date_time = None
                        stock_object.empty_allotment_date = None
                        stock_object.empty_allotment_in_date_time = None
                        stock_object.under_repair_out_date_time = None
                        stock_object.under_repair_in_date_time = None
                        stock_object.approved_out_date_time = None
                        stock_object.approved_in_date_time = None
                        stock_object.approval_pending_out_date_time = None
                        stock_object.approval_pending_in_date_time = None
                        stock_object.estimate_pending_out_date_time = None
                        stock_object.estimate_pending_in_date_time = None
                        stock_object.survey_pending_out_date_time = None
                        stock_object.survey_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Estimate Pending":
                        stock_object.available_in_date_time = None
                        stock_object.empty_allotment_out_date_time = None
                        stock_object.empty_allotment_date = None
                        stock_object.empty_allotment_in_date_time = None
                        stock_object.under_repair_out_date_time = None
                        stock_object.under_repair_in_date_time = None
                        stock_object.approved_out_date_time = None
                        stock_object.approved_in_date_time = None
                        stock_object.approval_pending_out_date_time = None
                        stock_object.approval_pending_in_date_time = None
                        stock_object.estimate_pending_out_date_time = None
                        stock_object.estimate_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Approval Pending":
                        stock_object.available_in_date_time = None
                        stock_object.empty_allotment_out_date_time = None
                        stock_object.empty_allotment_date = None
                        stock_object.empty_allotment_in_date_time = None
                        stock_object.under_repair_out_date_time = None
                        stock_object.under_repair_in_date_time = None
                        stock_object.approved_out_date_time = None
                        stock_object.approved_in_date_time = None
                        stock_object.approval_pending_out_date_time = None
                        stock_object.approval_pending_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Approved":
                        stock_object.available_in_date_time = None
                        stock_object.empty_allotment_out_date_time = None
                        stock_object.empty_allotment_date = None
                        stock_object.empty_allotment_in_date_time = None
                        stock_object.under_repair_out_date_time = None
                        stock_object.under_repair_in_date_time = None
                        stock_object.approved_out_date_time = None
                        stock_object.approved_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Under Repairing":
                        stock_object.available_in_date_time = None
                        stock_object.empty_allotment_out_date_time = None
                        stock_object.empty_allotment_date = None
                        stock_object.empty_allotment_in_date_time = None
                        stock_object.under_repair_out_date_time = None
                        stock_object.under_repair_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    elif status == "Empty Alloted":
                        stock_object.available_in_date_time = None
                        stock_object.empty_allotment_out_date_time = None
                        stock_object.empty_allotment_date = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        ).date()
                        stock_object.empty_allotment_in_date_time = (
                            datetime.datetime.now().astimezone(
                                timezone.get_current_timezone()
                            )
                        )
                        stock_object.save()
                    else:
                        pass
            else:
                pass

            seal_no = data["seal_no"]
            if seal_no == "":
                seal_no = None

            if not stock_object.seal_no == seal_no:
                if (
                    stock_object.seal_no is not None
                    and SealNo.objects.filter(number=stock_object.seal_no).exists()
                ):
                    old_seal_no_object = SealNo.objects.get(number=stock_object.seal_no)
                    old_seal_no_object.is_available = True
                    old_seal_no_object.in_use = False
                    old_seal_no_object.container_no = None
                    old_seal_no_object.in_use_date = None
                    if (
                        stock_object.container_status == "OUT"
                        and stock_object.container.status == "OUT"
                    ):
                        old_seal_no_object.out_date = None
                        old_seal_no_object.is_lock = False
                    old_seal_no_object.save()

                if seal_no is not None:
                    seal_no_object = SealNo.objects.get(number=seal_no)
                    stock_object.seal_no = seal_no_object.number
                    stock_object.save()
                    seal_no_object.is_available = False
                    seal_no_object.in_use = True
                    seal_no_object.container_no = stock_object.container.container_no
                    seal_no_object.in_use_date = datetime.datetime.now().astimezone(
                        timezone.get_current_timezone()
                    )
                    if (
                        stock_object.container_status == "OUT"
                        and stock_object.container.status == "OUT"
                    ):
                        seal_no_object.is_lock = True
                        stock_object.gate_out.seal_no = seal_no_object.number
                        out_history_object = GateOutHistory.objects.get(
                            container=stock_object.container,
                            gate_out=stock_object.gate_out,
                        )
                        seal_no_object.in_use_date = out_history_object.date
                        seal_no_object.out_date = out_history_object.date
                        stock_object.gate_out.save()
                    seal_no_object.save()
                else:
                    stock_object.seal_no = seal_no
                    if (
                        stock_object.container_status == "OUT"
                        and stock_object.container.status == "OUT"
                    ):
                        stock_object.gate_out.seal_no = seal_no
                        stock_object.gate_out.save()
                    stock_object.save()

            # if not stock_object.seal_no == seal_no:
            #     try:
            #         stock_object_with_given_seal_no = ContainerStock.objects.get(
            #             seal_no=seal_no
            #         )
            #         if stock_object_with_given_seal_no:
            #             return Response(
            #                 {"errorMsg": "Given Seal no is in Use"}, status=200
            #             )
            #     except:
            #         pass
            # else:
            #     pass

            if data["is_alloted"] == "False" and not len(data["booking_no"]) == 0:
                booking_no = data["booking_no"]
                container_no = data["container_no"]
                container_object = Container.objects.get(container_no=container_no)
                booking_object = ContainerAllotment.objects.get(booking_no=booking_no)
                booking_object.remove_container(container=container_object)
                stock_object.allotment_date = None
                stock_object.status = "Available"
                stock_object.booking_no = None
                stock_object.save()
                allotment_removed = True
            else:
                pass

            if allotment_removed is True:
                if (
                    stock_object.seal_no is not None
                    and SealNo.objects.filter(number=stock_object.seal_no).exists()
                ):
                    old_seal_no_object = SealNo.objects.get(number=stock_object.seal_no)
                    old_seal_no_object.is_available = True
                    old_seal_no_object.in_use = False
                    old_seal_no_object.container_no = None
                    old_seal_no_object.in_use_date = None
                    old_seal_no_object.out_date = None
                    old_seal_no_object.is_lock = False
                    old_seal_no_object.save()

                stock_object.seal_no = None
                stock_object.save()
                if stock_object.gate_out is not None:
                    stock_object.gate_out = None
                    stock_object.gate_out.save()

            if stock_object.available_date is not None:
                if status in ["Available", "Alloted", "Without_Repair_Available"]:
                    stock_object.available_date = available_date
                    available_time = datetime.datetime.now().time()
                    stock_object.available_in_date_time = datetime.datetime.combine(
                        available_date, available_time
                    ).astimezone(timezone.get_current_timezone())
                    stock_object.save()
                else:
                    stock_object.make_not_available()
            else:
                if status in ["Available", "Without_Repair_Available"]:
                    stock_object.make_available()
                    if (
                        stock_object.container.client.ref_code.upper() == "ZIM"
                        and not stock_object.is_zim_repair_sent
                        and status == "Available"
                    ):
                        zim_repair_edi_mail_process(obj=stock_object)
                        stock_object.is_zim_repair_sent = True
                        stock_object.save()

            stock_object.gate_in.grade = grade
            stock_object.gate_in.save()
            stock_object.grade = grade
            stock_object.status = status
            stock_object.save()
            return Response({"successMsg": "Data Update"}, status=200)
        except Exception as e:
            return Response({"errorMsg": f"Data Not Found [ {e} ]"}, status=200)


class AddOrRemoveContainerInQueue(views.APIView):
    """Post function will add container to queue"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            container_list = request.data["container_list"]
            for each in container_list:
                try:
                    container = Container.objects.get(container_no=each, status="IN")
                    if container.in_do_not_lift_queue is False:
                        container.in_do_not_lift_queue = True
                        container.queued_recently = True
                    else:
                        container.in_do_not_lift_queue = False
                    container.save()
                except:
                    pass
            return Response({"successMsg": "Data Saved"}, status=200)
        except Exception as e:
            return Response({"errorMsg": f"Data Not Found [ {e} ]"}, status=200)


class CollectContainerStock(views.APIView):
    """post function will collect and return container"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            container_list = request.data
            main_container_list = []
            location = ""
            site = ""
            for each in container_list:
                stock_object = ContainerStock.objects.get(pk=each)
                container = stock_object.container
                location = container.location.name
                site = container.site.name
                if stock_object.status in ["Available", "Without_Repair_Available"]:
                    main_container_list.append(container.container_no)
            return Response(
                {
                    "container_list": main_container_list,
                    "no_of_container": str(len(main_container_list)),
                    "location": location,
                    "site": site,
                },
                status=200,
            )
        except:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return Response({"errorMsg": f"Data Not Found"}, status=200)


class AllotmentOfContainer(views.APIView):
    """Post function will allot a container on new booking no or existing one"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data = request.data
            container_list = data["container_list"]
            main_container_list = []
            stock_list = []
            usa_containers = []
            for each in container_list:
                try:
                    if (
                        int(
                            Container.objects.filter(
                                container_no=each,
                                status="IN",
                                location__name=data["location"],
                                site__name=data["site"],
                            ).count()
                        )
                        == 1
                    ):
                        container = Container.objects.get(
                            container_no=each,
                            status="IN",
                            location__name=data["location"],
                            site__name=data["site"],
                        )
                        stock_data_object = ContainerStock.objects.get(
                            container=container, container_status="IN"
                        )
                        if stock_data_object.usa_approval_container:
                            usa_containers.append(each)
                        stock_list.append(stock_data_object)
                        main_container_list.append(container)
                except:
                    pass
            alert = ""
            if usa_containers:
                alert = f"USA Approved Containers [{', '.join(usa_containers)}]"

            if len(main_container_list) == 0:
                return Response(
                    {"errorMsg": "Please provide IN Stock Containers"}, status=200
                )

            booking_no = data["booking_no"]
            if len(booking_no) == 0:
                return Response({"errorMsg": "Please provide booking no"}, status=200)

            new_allotment = [s for s in stock_list if not s.status == "Alloted"]

            try:
                allot_object = ContainerAllotment.objects.get(booking_no=booking_no)
                if allot_object:
                    if allot_object.remaining <= 0:
                        return Response(
                            {
                                "errorMsg": "No Remaining Quantity of Container Left on this booking no !!!"
                            },
                            status=200,
                        )

                    if len(new_allotment) > allot_object.remaining:
                        new_allot_qty = len(new_allotment) + int(
                            allot_object.container.all().count()
                        )
                        return Response(
                            {
                                "errorMsg": f"Remaining Quantity is [{allot_object.remaining}] and you are alloting [{len(new_allotment)}] more containers, Please Update Allot Quantity to [{new_allot_qty}] !!!"
                            },
                            status=200,
                        )

                    else:
                        for stock_object in new_allotment:
                            allot_object.add_container(container=stock_object.container)
                            stock_object.allotment_date = allot_object.booking_date
                            stock_object.available_out_date_time = (
                                datetime.datetime.now().astimezone(
                                    timezone.get_current_timezone()
                                )
                            )
                            stock_object.allotment_in_date_time = (
                                datetime.datetime.now().astimezone(
                                    timezone.get_current_timezone()
                                )
                            )
                            stock_object.status = "Alloted"
                            stock_object.booking_no = booking_no
                            stock_object.save()
                            create_allotment_track_record(
                                location=stock_object.container.location,
                                site=stock_object.container.site,
                                booking_date=allot_object.booking_date,
                                booking_no=booking_no,
                                container_no=stock_object.container.container_no,
                                stock_id=stock_object.pk,
                            )
                        return Response(
                            {"successMsg": "Data Saved", "alert": alert}, status=200
                        )
                else:
                    pass
            except:
                booking_date_str = data["booking_date"]
                if len(booking_date_str) == 0:
                    booking_date = (
                        datetime.datetime.now()
                        .astimezone(timezone.get_current_timezone())
                        .date()
                    )
                else:
                    try:
                        booking_date = datetime.datetime.strptime(
                            booking_date_str, "%Y-%m-%d"
                        ).date()
                    except:
                        booking_date = (
                            datetime.datetime.now()
                            .astimezone(timezone.get_current_timezone())
                            .date()
                        )

                booking_party = data["booking_party"]
                if len(booking_party) == 0:
                    booking_party = None

                validity_date_str = data["validity_date"]
                if len(validity_date_str) == 0:
                    validity_date = None
                else:
                    try:
                        validity_date = datetime.datetime.strptime(
                            validity_date_str, "%Y-%m-%d"
                        ).date()
                    except:
                        validity_date = None

                quantity = data["quantity"]
                if (
                    len(quantity) == 0
                    or quantity.replace(".", "", 1).isnumeric() is False
                ):
                    quantity = 1
                else:
                    quantity = int(quantity)

                remarks = data["remarks"]
                if len(remarks) == 0:
                    remarks = None

                allot = ContainerAllotment.create(
                    booking_date=booking_date,
                    booking_no=booking_no,
                    booking_party=booking_party,
                    validity_date=validity_date,
                    quantity=quantity,
                    remarks=remarks,
                )
                allot.save()

                for stock_object in new_allotment:
                    allot.add_container(container=stock_object.container)
                    stock_object.allotment_date = allot.booking_date
                    stock_object.available_out_date_time = (
                        datetime.datetime.now().astimezone(
                            timezone.get_current_timezone()
                        )
                    )
                    stock_object.allotment_in_date_time = (
                        datetime.datetime.now().astimezone(
                            timezone.get_current_timezone()
                        )
                    )
                    stock_object.status = "Alloted"
                    stock_object.booking_no = booking_no
                    stock_object.save()
                    create_allotment_track_record(
                        location=stock_object.container.location,
                        site=stock_object.container.site,
                        booking_date=allot.booking_date,
                        booking_no=booking_no,
                        container_no=stock_object.container.container_no,
                        stock_id=stock_object.pk,
                    )

                return Response(
                    {"successMsg": "Data Saved", "alert": alert}, status=200
                )
        except Exception as e:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return Response({"errorMsg": f"Data Not Found [ {e} ]"}, status=200)


class RemoveAllotment(views.APIView):
    """Post function will remove the  containers from booking"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            for pk in request.data:
                stock_object = ContainerStock.objects.get(pk=pk)
                booking_no = stock_object.booking_no
                container_object = stock_object.container
                booking_object = ContainerAllotment.objects.get(booking_no=booking_no)
                booking_object.remove_container(container=container_object)

                if AllotmentTracker.objects.filter(
                    location=stock_object.container.location,
                    site=stock_object.container.site,
                    booking_no=booking_no,
                    container_no=stock_object.container.container_no,
                    stock_id=stock_object.pk,
                ).exists():
                    booking_track_record = AllotmentTracker.objects.filter(
                        location=stock_object.container.location,
                        site=stock_object.container.site,
                        booking_no=booking_no,
                        container_no=stock_object.container.container_no,
                        stock_id=stock_object.pk,
                    ).latest("pk")
                    booking_track_record.booking_canceled = True
                    booking_track_record.save()

                if (
                    stock_object.seal_no is not None
                    and SealNo.objects.filter(number=stock_object.seal_no).exists()
                ):
                    seal_no_object = SealNo.objects.get(number=stock_object.seal_no)
                    seal_no_object.is_available = True
                    seal_no_object.in_use = False
                    seal_no_object.container_no = None
                    seal_no_object.in_use_date = None
                    seal_no_object.out_date = None
                    seal_no_object.is_lock = False
                    seal_no_object.save()
                stock_object.allotment_date = None
                stock_object.available_out_date_time = None
                stock_object.allotment_in_date_time = None
                stock_object.status = "Available"
                stock_object.booking_no = None
                stock_object.seal_no = None
                stock_object.save()
            return Response({"successMsg": "Data Saved"}, status=200)
        except Exception as e:
            return Response({"errorMsg": f"Data Not Found [ {e} ]"}, status=200)


class UpdateAllotmentOfContainer(views.APIView):
    """Post function will update booking no quantity"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, pk, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data = request.data
            allot_object = ContainerAllotment.objects.get(pk=pk)
            quantity = data["quantity"]
            if len(quantity) == 0 or quantity.replace(".", "", 1).isnumeric() is False:
                quantity = 1
            else:
                quantity = int(quantity)

            if not allot_object.quantity == quantity:
                if quantity > int(allot_object.container.all().count()):
                    remaining = quantity - int(allot_object.container.all().count())
                else:
                    quantity = int(allot_object.container.all().count())
                    remaining = quantity - int(allot_object.container.all().count())

                if remaining <= 0:
                    return Response(
                        {
                            "errorMsg": "Given Quantity is less than or equal to  the containers "
                            "present on this booking no"
                        },
                        status=200,
                    )
                allot_object.quantity = quantity
                allot_object.remaining = remaining
                allot_object.save()
            else:
                pass

            return Response({"successMsg": "Data Saved"}, status=200)
        except Exception as e:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return Response({"errorMsg": f"Data Not Found [ {e} ]"}, status=200)


# class UpdateAllotmentOfContainer(views.APIView):
#     """Post function will update booking no"""

#     permission_classes = (IsAuthenticated,)

#     def post(self, request, pk, *args, **kwargs):
#         if not request.data:
#             return Response({"errorMsg": "Please provide Data"}, status=200)
#         try:
#             data = request.data

#             booking_no = data["booking_no"]
#             if len(booking_no) == 0:
#                 return Response({"errorMsg": "Please provide booking no"}, status=200)

#             allot_object = ContainerAllotment.objects.get(pk=pk)

#             if not allot_object.booking_no == booking_no:
#                 try:
#                     booking_object = ContainerAllotment.objects.get(
#                         booking_no=booking_no
#                     )
#                     if booking_object:
#                         return Response(
#                             {"errorMsg": "Provided booking no is in Use"}, status=200
#                         )
#                 except:
#                     pass

#             booking_date_str = data["booking_date"]
#             if len(booking_date_str) == 0:
#                 booking_date = (
#                     datetime.datetime.now()
#                     .astimezone(timezone.get_current_timezone())
#                     .date()
#                 )
#             else:
#                 try:
#                     booking_date = datetime.datetime.strptime(
#                         booking_date_str, "%Y-%m-%d"
#                     ).date()
#                 except:
#                     booking_date = (
#                         datetime.datetime.now()
#                         .astimezone(timezone.get_current_timezone())
#                         .date()
#                     )

#             booking_party = data["booking_party"]
#             if len(booking_party) == 0:
#                 booking_party = None

#             validity_date_str = data["validity_date"]
#             if len(validity_date_str) == 0:
#                 validity_date = None
#             else:
#                 try:
#                     validity_date = datetime.datetime.strptime(
#                         validity_date_str, "%Y-%m-%d"
#                     ).date()
#                 except:
#                     validity_date = None

#             quantity = data["quantity"]
#             if len(quantity) == 0 or quantity.replace(".", "", 1).isnumeric() is False:
#                 quantity = 1
#             else:
#                 quantity = int(quantity)

#             if not allot_object.quantity == quantity:
#                 if quantity > int(allot_object.container.all().count()):
#                     remaining = quantity - int(allot_object.container.all().count())
#                 else:
#                     quantity = int(allot_object.container.all().count())
#                     remaining = quantity - int(allot_object.container.all().count())

#                 if remaining <= 0:
#                     return Response(
#                         {
#                             "errorMsg": "Given Quantity is less than or equal to  the containers "
#                             "present on this booking no"
#                         },
#                         status=200,
#                     )
#                 allot_object.quantity = quantity
#                 allot_object.remaining = remaining
#                 allot_object.save()
#             else:
#                 pass

#             remarks = data["remarks"]
#             if len(remarks) == 0:
#                 remarks = None

#             if allot_object.container.all():
#                 for each in allot_object.container.all():
#                     if each.status == "OUT":
#                         stock_object = ContainerStock.objects.get(
#                             container=each,
#                             container_status="OUT",
#                             booking_no=allot_object.booking_no,
#                         )
#                         gate_out_object = GateOut.objects.get(
#                             container=each, booking_no=allot_object.booking_no
#                         )
#                         gate_out_object.booking_no = booking_no
#                         gate_out_object.booking_date = booking_date
#                         gate_out_object.booking_party = booking_party
#                         gate_out_object.save()
#                         stock_object.booking_no = booking_no
#                         stock_object.save()
#                     else:
#                         stock_object = ContainerStock.objects.get(
#                             container=each, container_status="IN"
#                         )
#                         stock_object.booking_no = booking_no
#                         stock_object.save()
#             else:
#                 pass

#             allot_object.booking_date = booking_date
#             allot_object.booking_no = booking_no
#             allot_object.booking_party = booking_party
#             allot_object.validity_date = validity_date
#             allot_object.remarks = remarks
#             allot_object.save()

#             # allotment
#             container_list = data.get("container_list", None)
#             if container_list is not None:
#                 main_container_list = []
#                 stock_list = []
#                 for each in container_list:
#                     try:
#                         if (
#                             int(
#                                 Container.objects.filter(
#                                     container_no=each,
#                                     status="IN",
#                                     location__name=data["location"],
#                                     site__name=data["site"],
#                                 ).count()
#                             )
#                             == 1
#                         ):
#                             container = Container.objects.get(
#                                 container_no=each,
#                                 status="IN",
#                                 location__name=data["location"],
#                                 site__name=data["site"],
#                             )
#                             stock_data_object = ContainerStock.objects.get(
#                                 container=container, container_status="IN"
#                             )
#                             stock_list.append(stock_data_object)
#                             main_container_list.append(container)
#                     except:
#                         pass

#                 if len(main_container_list) == 0:
#                     return Response(
#                         {"errorMsg": "Please provide IN Stock Containers"}, status=200
#                     )

#                 if allot_object.remaining <= 0:
#                     return Response(
#                         {
#                             "errorMsg": "No Remaining Quantity of Container Left on this booking no "
#                         },
#                         status=200,
#                     )
#                 else:
#                     for each in main_container_list:
#                         allot_object.add_container(container=each)
#                     for stock_object in stock_list:
#                         stock_object.allotment_date = allot_object.booking_date
#                         stock_object.available_out_date_time = (
#                             datetime.datetime.now().astimezone(
#                                 timezone.get_current_timezone()
#                             )
#                         )
#                         stock_object.allotment_in_date_time = (
#                             datetime.datetime.now().astimezone(
#                                 timezone.get_current_timezone()
#                             )
#                         )
#                         stock_object.status = "Alloted"
#                         stock_object.booking_no = booking_no
#                         stock_object.save()

#             return Response({"successMsg": "Data Saved"}, status=200)
#         except Exception as e:
#             error_log = logging.getLogger("error_log")
#             error_log.error(traceback.format_exc())
#             return Response({"errorMsg": f"Data Not Found [ {e} ]"}, status=200)


class BookingNumberDetails(views.APIView):
    """Post function will get the booking no data"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data = request.data
            booking_no = data["booking_no"]
            booking_detail = ContainerAllotment.objects.get(
                booking_no=booking_no
            ).get_allotment()
            return Response(booking_detail, status=200)
        except Exception as e:
            return Response({"errorMsg": f"Data Not Found [ {e} ]"}, status=200)


class StockUploadSampleFile(views.APIView):
    """Get Function will give the sample file"""

    """ Post Function will give the extracted data """

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def get(self, request, *args, **kwargs):
        try:
            temp_file_path = os.path.join(
                BASE_DIR, "sample_stock/sample_stock_upload.xlsx"
            )
            with open(temp_file_path, "rb") as temp:
                file_response = HttpResponse(
                    temp.read(), content_type=f"application/xlsx"
                )
                file_response["Content-Disposition"] = (
                    f'attachment; filename="sample_stock_upload.xlsx"'
                )
            return file_response
        except:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return Response({"errorMsg": f"Data Not Found"}, status=200)

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data = request.data["file"]
            user = request.user
            app_user = AccountUser.objects.get(username=user.username)
            try:
                location_str = request.data["location"]
                site_str = request.data["site"]
                location = Location.objects.get(name=location_str)
                site = Site.objects.get(name=site_str)
            except:
                location = app_user.location
                site = app_user.site
            result = extract_excel_data(input_excel=data, location=location, site=site)
            if result == "Header Not Found" or result is False:
                return Response(
                    {"errorMsg": "Data file is corrupted, unable to import data"},
                    status=200,
                )
            else:
                return Response(result, status=200)
        except Exception as e:
            return Response({"errorMsg": f"Invalid Data Provided [ {e} ]"}, status=200)


class StockImport(views.APIView):
    """Post Function will Import the data"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            with transaction.atomic():
                data = request.data["importable_data"]
                user = request.user
                app_user = AccountUser.objects.get(username=user.username)
                try:
                    location_str = request.data["location"]
                    site_str = request.data["site"]
                    location = Location.objects.get(name=location_str)
                    site = Site.objects.get(name=site_str)
                except:
                    location = app_user.location
                    site = app_user.site

                automatic_mnr_status_change = site.automatic_mnr_status_change

                for each in data:
                    client_str = each["client"]
                    shipping_line_str = each["shipping_line"]
                    type_str = each["type"]
                    size_str = each["size"]
                    container_no_str = each["container_no"]
                    if each["gross_wt"] == "":
                        gross_wt_str = 0
                    else:
                        gross_wt_str = str(int(float(each["gross_wt"])))
                    if each["tare_wt"] == "":
                        tare_wt_str = 0
                    else:
                        tare_wt_str = str(int(float(each["tare_wt"])))
                    manufacturing_date_str = each["manufacturing_date"]
                    in_date_str = each["in_date"]
                    in_time_str = each["in_time"]
                    condition_str = each["condition"]
                    grade_str = each["grade"]
                    arrived_str = each["arrived"]
                    client_object = Client.objects.get(
                        name=client_str, location=location, site=site
                    )
                    shipping_line = list(
                        ClientAbbreviation.objects.filter(
                            client=client_object, name=shipping_line_str
                        )
                    )[0]
                    type_object = ContainerType.objects.get(name=type_str)
                    size_object = ContainerSize.objects.get(
                        name=str(int(float(size_str)))
                    )
                    manufacturing_date = datetime.datetime.strptime(
                        manufacturing_date_str, "%d_%m_%Y"
                    ).date()
                    in_date = datetime.datetime.strptime(in_date_str, "%d_%m_%Y").date()
                    in_time = datetime.datetime.strptime(in_time_str, "%H_%M").time()
                    transporter_name = each["transporter"]
                    transporter = Transporter.objects.get(
                        name=transporter_name, location=location, site=site
                    )
                    vehicle_no = each["vehicle_no"]
                    if vehicle_no == "":
                        vehicle_no = None
                    shipper = each["shipper"]
                    if shipper == "":
                        shipper = None
                    vessel = each["vessel"]
                    if vessel == "":
                        vessel = None
                    voyage = each["voyage"]
                    if voyage == "":
                        voyage = ""
                    cargo = each["cargo"]
                    if cargo == "":
                        cargo = None

                    #     Check
                    try:
                        container_object = Container.objects.get(
                            container_no=container_no_str, location=location, site=site
                        )
                        if container_object.status == "IN":
                            return Response(
                                {"errorMsg": "Container Already In"}, status=200
                            )
                        else:
                            container_object.status = "IN"
                            container_object.is_available = False
                            container_object.save()
                    except:
                        container_no_response = check_digit(
                            container_no=container_no_str
                        )
                        #     Adding Container
                        container_object = Container.create(
                            client=client_object,
                            type_object=type_object,
                            size_object=size_object,
                            container_no=container_no_str,
                            payload=None,
                            gross_wt=gross_wt_str,
                            tare_wt=tare_wt_str,
                            manufacturing_date=manufacturing_date,
                            shipping_line=shipping_line,
                            leased_box=False,
                            is_valid=container_no_response,
                            in_do_not_lift_queue=False,
                            queued_recently=False,
                            automatic_mnr_status_change=automatic_mnr_status_change,
                            do_not_lift_remarks=None,
                            location=location,
                            site=site,
                        )
                        container_object.save()

                    #     Adding EIR
                    eir_object = Eir.create(
                        container=container_object,
                        done_by=app_user,
                        eir_date=in_date,
                        eir_time=in_time,
                        offload_date=in_date,
                        offload_time=in_time,
                        eir_img=save_default_eir_img(),
                        eir_amount=float(0),
                        repair_amount=float(0),
                        entry_type="IN",
                    )
                    eir_object.save()
                    eir_object.save_eir_no(location=location)
                    eir_object.upload_eir_img()

                    #     Adding Gate In
                    gate_in_object = GateIn.create(
                        container=container_object,
                        in_date=in_date,
                        in_time=in_time,
                        condition=condition_str,
                        grade=grade_str,
                        arrived=arrived_str,
                        consignee=None,
                        shipper=shipper,
                        source=None,
                        from_location_code=None,
                        from_port_code=None,
                        vessel_name=vessel,
                        voyage_no=voyage,
                        do_ref=None,
                        cargo=cargo,
                        export_cargo_type=None,
                        transporter_name=transporter,
                        vehicle_no=vehicle_no,
                        bl_no=None,
                        driver_name=None,
                        driver_license=None,
                        driver_mobile_no=None,
                        carrier_code=None,
                        location=location,
                        remarks=None,
                        from_location_name_code=None,
                    )
                    gate_in_object.save()

                    # gate in time and eir time 15 mins difference
                    temp_date_time = datetime.datetime.combine(
                        gate_in_object.in_date, gate_in_object.in_time
                    ) - datetime.timedelta(minutes=15)
                    temp_date_time = temp_date_time.astimezone(
                        timezone.get_current_timezone()
                    )
                    eir_object.eir_date = temp_date_time.date()
                    eir_object.eir_time = temp_date_time.time()
                    eir_object.save()

                    serial_no = str(gate_in_object.pk).zfill(5)
                    gate_in_object.gate_pass_no = (
                        f"GIP/{location.code}/{temp_date_time.date().year}/{serial_no}"
                    )
                    gate_in_object.save()

                    # Adding  Stock IN
                    stock_object = ContainerStock.create(
                        gate_in=gate_in_object,
                        container=container_object,
                        grade=grade_str,
                    )
                    stock_object.survey_pending_in_date_time = (
                        datetime.datetime.combine(in_date, in_time).astimezone(
                            timezone.get_current_timezone()
                        )
                    )
                    stock_object.save()

                    if str(automatic_mnr_status_change) == "False":
                        stock_object.status = "Estimate Pending"
                        stock_object.survey_pending_out_date_time = temp_date_time
                        stock_object.estimate_pending_in_date_time = temp_date_time
                        stock_object.save()

                    #     Adding Handling
                    handling_objects = Handling.create(
                        container=container_object,
                        apply_charges="Line",
                        customer_name=None,
                        invoice_date=in_date,
                        receipt_date=in_date,
                        lolo_type="OFF",
                        payment_type="None",
                        lolo_amount=float(0),
                        remark=None,
                        entry_type="IN",
                    )
                    handling_objects.save()
                    handling_objects.save_invoice_no(location=location)
                    handling_objects.save_receipt_no(location=location)

                    #  Adding Gate IN history
                    gih = GateInHistory(
                        created_at=datetime.datetime.now().astimezone(
                            timezone.get_current_timezone()
                        ),
                        date=datetime.datetime.combine(in_date, in_time).astimezone(
                            timezone.get_current_timezone()
                        ),
                        container=container_object,
                        eir=eir_object,
                        gate_in=gate_in_object,
                        lolo=handling_objects,
                        lolo_payment=None,
                        st=None,
                        st_payment=None,
                    )
                    gih.save()
                    container_in_out_record = ContainerInOutRecord(
                        container=container_object, in_data=gih, out_data=None
                    )
                    container_in_out_record.save()

                return Response({"successMsg": "Stock Imported"}, status=200)
        except Exception as e:
            return Response({"errorMsg": f"Invalid Data Provided [ {e} ]"}, status=200)


class RejectedStockDataFile(views.APIView):
    """Post Function will download and return the rejected data in a xls file"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data = request.data["rejected_data"]
            faults = request.data["faults"]
            shipping_line = []
            client = []
            c_type = []
            c_size = []
            container_no = []
            gross_wt = []
            tare_wt = []
            manufacturing_date = []
            in_date = []
            in_time = []
            condition = []
            grade = []
            arrived = []
            transporter = []
            vehicle_no = []
            shipper = []
            vessel = []
            voyage = []
            cargo = []
            for each in data:
                shipping_line.append(each["shipping_line"])
                client.append(each["client"])
                c_type.append(each["type"])
                c_size.append(each["size"])
                container_no.append(each["container_no"])
                if each["gross_wt"] == "":
                    each["gross_wt"] = "_"
                    gross_wt.append(each["gross_wt"])
                else:
                    gross_wt.append(each["gross_wt"])
                if each["tare_wt"] == "":
                    each["tare_wt"] = "_"
                    tare_wt.append(each["tare_wt"])
                else:
                    tare_wt.append(each["tare_wt"])
                manufacturing_date.append(each["manufacturing_date"])
                in_date.append(each["in_date"])
                in_time.append(each["in_time"])
                condition.append(each["condition"])
                if each["grade"] == "":
                    each["grade"] = "_"
                    grade.append(each["grade"])
                else:
                    grade.append(each["grade"])
                arrived.append(each["arrived"])
                transporter.append(each["transporter"])
                if each["vehicle_no"] == "":
                    each["vehicle_no"] = "_"
                    vehicle_no.append(each["vehicle_no"])
                else:
                    vehicle_no.append(each["vehicle_no"])
                if each["shipper"] == "":
                    each["shipper"] = "_"
                    shipper.append(each["shipper"])
                else:
                    shipper.append(each["shipper"])
                if each["vessel"] == "":
                    each["vessel"] = "_"
                    vessel.append(each["vessel"])
                else:
                    vessel.append(each["vessel"])
                if each["voyage"] == "":
                    each["voyage"] = "_"
                    voyage.append(each["voyage"])
                else:
                    voyage.append(each["voyage"])
                if each["cargo"] == "":
                    each["cargo"] = "_"
                    cargo.append(each["cargo"])
                else:
                    cargo.append(each["cargo"])

            stock_df_data = {
                "shipping_line": shipping_line,
                "client": client,
                "type": c_type,
                "size": c_size,
                "container_no": container_no,
                "gross_wt": gross_wt,
                "tare_wt": tare_wt,
                "manufacturing_date": manufacturing_date,
                "in_date": in_date,
                "in_time": in_time,
                "condition": condition,
                "grade": grade,
                "arrived": arrived,
                "transporter": transporter,
                "vehicle_no": vehicle_no,
                "shipper": shipper,
                "vessel": vessel,
                "voyage": voyage,
                "cargo": cargo,
            }
            if not os.path.exists(os.path.join(BASE_DIR, "temp/sample_stock/")):
                os.makedirs(os.path.join(BASE_DIR, "temp/sample_stock/"))
            new_temp_file_path = os.path.join(
                BASE_DIR, f"temp/sample_stock/sample_stock_upload.xlsx"
            )
            temp_file_path = os.path.join(
                BASE_DIR, "sample_stock/sample_stock_upload.xlsx"
            )
            old_temp_data = None
            with open(temp_file_path, "rb") as temp:
                old_temp_data = temp.read()
            with open(new_temp_file_path, "wb") as f:
                f.write(old_temp_data)

            # New Version code
            book = load_workbook(new_temp_file_path)
            with pd.ExcelWriter(
                new_temp_file_path,
                engine="openpyxl",
                mode="a",
                if_sheet_exists="replace",
            ) as writer:

                stock_df = pd.DataFrame(stock_df_data)
                stock_df.to_excel(writer, sheet_name="stock", index=False)

                faults_df = pd.DataFrame(faults)
                faults_df.to_excel(writer, sheet_name="faults", index=False)

            # Old version
            # book = load_workbook(new_temp_file_path)
            # writer = pd.ExcelWriter(new_temp_file_path, engine="openpyxl")
            # writer.book = book
            # writer.sheets = dict((ws.title, ws) for ws in book.worksheets)
            # stock_df = pd.DataFrame(stock_df_data)
            # faults_df = pd.DataFrame(faults)
            # stock_df.to_excel(
            #     writer, sheet_name="stock", startrow=0, startcol=0, index=False
            # )
            # faults_df.to_excel(
            #     writer, sheet_name="faults", startrow=0, startcol=0, index=False
            # )
            # fault_sheet = book.get_sheet_by_name("faults")
            # writer.save()

            with open(new_temp_file_path, "rb") as temp:
                file_response = HttpResponse(
                    temp.read(), content_type=f"application/xlsx"
                )
                file_response["Content-Disposition"] = (
                    f'attachment; filename="sample_stock_upload.xlsx"'
                )
                os.remove(new_temp_file_path)
            return file_response
        except Exception as e:
            return Response({"errorMsg": f"Invalid Data Provided [ {e} ]"}, status=200)


class StockSheetDownloadView(views.APIView):
    """
    the Post function will give you all container in stock sheet
    """

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data = request.data
            location = request.data["location"]
            site = request.data["site"]
            ref_code = data.get("ref_code", None)
            queryset = None

            if data["out_history"] == "True":
                if location == "ALL":
                    queryset = (
                        ContainerStock.objects.select_related(
                            "container",
                            "container__client",
                            "container__type",
                            "container__size",
                            "gate_in",
                        )
                        .filter(container_status="OUT")
                        .order_by("-gate_in__in_date")
                    )

                elif site == "ALL":
                    queryset = (
                        ContainerStock.objects.select_related(
                            "container",
                            "container__client",
                            "container__type",
                            "container__size",
                            "container__location",
                            "gate_in",
                        )
                        .filter(
                            container_status="OUT", container__location__name=location
                        )
                        .order_by("-gate_in__in_date")
                    )

                else:
                    queryset = (
                        ContainerStock.objects.select_related(
                            "container",
                            "container__client",
                            "container__type",
                            "container__size",
                            "container__location",
                            "container__site",
                            "gate_in",
                        )
                        .filter(
                            container_status="OUT",
                            container__location__name=location,
                            container__site__name=site,
                        )
                        .order_by("-gate_in__in_date")
                    )
            else:
                if data["do_not_lift_queue"] == "True":
                    if location == "ALL":
                        queryset = (
                            ContainerStock.objects.select_related(
                                "container",
                                "container__client",
                                "container__type",
                                "container__size",
                                "gate_in",
                            )
                            .filter(
                                container_status="IN",
                                container__in_do_not_lift_queue=True,
                            )
                            .order_by("-gate_in__in_date")
                        )
                    elif site == "ALL":
                        queryset = (
                            ContainerStock.objects.select_related(
                                "container",
                                "container__client",
                                "container__type",
                                "container__size",
                                "container__location",
                                "gate_in",
                            )
                            .filter(
                                container_status="IN",
                                container__location__name=location,
                                container__in_do_not_lift_queue=True,
                            )
                            .order_by("-gate_in__in_date")
                        )
                    else:
                        queryset = (
                            ContainerStock.objects.select_related(
                                "container",
                                "container__client",
                                "container__type",
                                "container__size",
                                "container__location",
                                "container__site",
                                "gate_in",
                            )
                            .filter(
                                container_status="IN",
                                container__location__name=location,
                                container__site__name=site,
                                container__in_do_not_lift_queue=True,
                            )
                            .order_by("-gate_in__in_date")
                        )
                else:
                    if location == "ALL":
                        queryset = (
                            ContainerStock.objects.select_related(
                                "container",
                                "container__client",
                                "container__type",
                                "container__size",
                                "gate_in",
                            )
                            .filter(
                                container_status="IN",
                                container__in_do_not_lift_queue=False,
                            )
                            .order_by("-gate_in__in_date")
                        )
                    elif site == "ALL":
                        queryset = (
                            ContainerStock.objects.select_related(
                                "container",
                                "container__client",
                                "container__type",
                                "container__size",
                                "container__location",
                                "gate_in",
                            )
                            .filter(
                                container_status="IN",
                                container__location__name=location,
                                container__in_do_not_lift_queue=False,
                            )
                            .order_by("-gate_in__in_date")
                        )
                    else:
                        queryset = (
                            ContainerStock.objects.select_related(
                                "container",
                                "container__client",
                                "container__type",
                                "container__size",
                                "container__location",
                                "container__site",
                                "gate_in",
                            )
                            .filter(
                                container_status="IN",
                                container__location__name=location,
                                container__site__name=site,
                                container__in_do_not_lift_queue=False,
                            )
                            .order_by("-gate_in__in_date")
                        )

            if ref_code is not None:
                if not len(ref_code) == 0:
                    queryset = queryset.filter(container__client__ref_code=ref_code)

            if not len(data["client"]) == 0:
                queryset = queryset.filter(container__client__name=data["client"])

            if not len(data["container_no"]) == 0:
                queryset = queryset.filter(
                    container__container_no__in=data["container_no"]
                )

            if not len(data["booking_no"]) == 0:
                queryset = queryset.filter(booking_no=data["booking_no"])

            if not len(data["status"]) == 0:
                queryset = queryset.filter(status=data["status"])

            gate_in_date = data["gate_in_date"]
            if not len(gate_in_date["from"]) == 0 and not len(gate_in_date["to"]) == 0:
                queryset = queryset.filter(
                    gate_in__in_date__range=[gate_in_date["from"], gate_in_date["to"]]
                )

            gate_out_date = data["gate_out_date"]
            if (
                not len(gate_out_date["from"]) == 0
                and not len(gate_out_date["to"]) == 0
            ):
                queryset = queryset.filter(
                    gate_out__out_date__range=[
                        gate_out_date["from"],
                        gate_out_date["to"],
                    ]
                )

            available_date = data["available_date"]
            if (
                not len(available_date["from"]) == 0
                and not len(available_date["to"]) == 0
            ):
                queryset = queryset.filter(
                    available_date__range=[available_date["from"], available_date["to"]]
                )

            allotment_date = data["allotment_date"]
            if (
                not len(allotment_date["from"]) == 0
                and not len(allotment_date["to"]) == 0
            ):
                queryset = queryset.filter(
                    allotment_date__range=[allotment_date["from"], allotment_date["to"]]
                )

            location_obj = Location.objects.get(name=location)
            site_obj = Site.objects.get(name=site, location=location_obj)
            user_data_dict = user_data(location=location_obj, site=site_obj)
            stock_df_data = generic_stock_sheet_main_data(stock_data=queryset)
            stock_file_path = create_generic_stock_sheet_wb(
                stock_df_data=stock_df_data, user_data=user_data_dict
            )
            dt = datetime.datetime.now().astimezone(timezone.get_current_timezone())
            date = dt.date().strftime("%Y%m%d")
            time = dt.time().strftime("%H%M%S")
            with open(stock_file_path, "rb") as temp:
                file_response = HttpResponse(
                    temp.read(), content_type=f"application/xlsx"
                )
                file_response["Content-Disposition"] = (
                    f'attachment; filename="stock_sheet_{date}{time}.xlsx"'
                )
                os.remove(stock_file_path)
            return file_response
        except Exception as e:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return Response({"errorMsg": f"Data Not Found [ {e} ]"}, status=200)
