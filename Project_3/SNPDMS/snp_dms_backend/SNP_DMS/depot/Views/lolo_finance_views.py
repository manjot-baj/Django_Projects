from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from SNP_DMS.settings.base import BASE_DIR
from django.http import HttpResponse
import logging, traceback, datetime, os
from django.db import transaction
from master.models_two import Client
from master.models import ContainerType, ContainerSize, Location, Site, TypeSizeCode
import pandas as pd
from openpyxl import load_workbook
from depot.models import (
    Container,
    ContainerStock,
    GateInHistory,
    GateOutHistory,
    SealNo,
    ContainerAllotment,
)
from depot.lolo_finance_models import (
    PreGateIn,
    PreGateOut,
    AdvancedHandlingPayment,
    DoValidityUpgradeHistory,
    CustomerFinAccount,
    FinAccountTransaction,
)
from depot.functions_two import *
from depot.lolo_finance_functions import *
from django.core.paginator import Paginator
from django.db.models import Q
from depot.lolo_finance_do_scan_functions import do_data_extraction
from wkhtmltopdf.views import PDFTemplateResponse
from account.permissions import HasAllowedRoles
from billing_invoice.models import (
    CustomerBill,
    CustomerBillInvoiceLine,
    CustomerBillInvoice,
)
import xlsxwriter
import random
from SNP_DMS.settings.base import BASE_DIR
from django.utils import timezone


class AdvancedHandlingPaymentListAPIView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        try:
            pg_no = request.data["pg_no"]
            on_page_data = request.data["on_page_data"]
            client_name = request.data.get("client")
            entry_type = request.data.get("entry_type")
            bl_no = request.data.get("bl_no")
            bk_no = request.data.get("bk_no")
            location_name = request.data.get("location")
            site_name = request.data.get("site")
            payment_type = request.data.get("payment_type")
            with_gst = request.data.get("with_gst", None)

            is_balance_pymt_adjusted = request.data.get(
                "is_balance_pymt_adjusted", None
            )

            from_date_time = None
            to_date_time = None
            if (
                not len(request.data.get("from_date")) == 0
                and not len(request.data.get("to_date")) == 0
            ):
                from_date = datetime.datetime.strptime(
                    request.data.get("from_date"), "%Y-%m-%d"
                ).date()
                from_time = datetime.datetime.strptime("00:00", "%H:%M").time()
                from_date_time = datetime.datetime.combine(from_date, from_time)
                to_date = datetime.datetime.strptime(
                    request.data.get("to_date"), "%Y-%m-%d"
                ).date()
                to_time = datetime.datetime.strptime("23:59", "%H:%M").time()
                to_date_time = datetime.datetime.combine(to_date, to_time)

            filters = Q()
            if client_name:
                filters &= Q(client__name=client_name)
            if location_name:
                filters &= Q(location__name=location_name)
            if site_name:
                filters &= Q(site__name=site_name)
            if payment_type:
                filters &= Q(payment_type=payment_type)
            if from_date_time and to_date_time:
                filters &= Q(date__range=(from_date_time, to_date_time))
            if entry_type:
                filters &= Q(entry_type=entry_type)
            if bl_no:
                filters &= Q(bl_no=bl_no)
            if bk_no:
                filters &= Q(bk_no=bk_no)
            if is_balance_pymt_adjusted is not None:
                filters &= Q(is_balance_pymt_adjusted=is_balance_pymt_adjusted)
            if with_gst is not None:
                filters &= Q(with_gst=with_gst)

            payments_qs = AdvancedHandlingPayment.objects.filter(filters)

            # pagination
            paginator = Paginator(payments_qs, on_page_data)
            current_page = paginator.page(pg_no)
            no_of_data_count = paginator.count
            no_of_pages = paginator.num_pages
            on_page_data_count = current_page.object_list.count()
            prev_page = (
                current_page.previous_page_number()
                if current_page.has_previous()
                else None
            )
            next_page = (
                current_page.next_page_number() if current_page.has_next() else None
            )

            response_data = [
                {
                    **payment.get_payment_data(),
                    "sr_no": current_page.start_index() + index,
                }
                for index, payment in enumerate(current_page.object_list)
            ]

            return Response(
                {
                    "no_of_data": no_of_data_count,
                    "on_page_data": on_page_data_count,
                    "total_pages": no_of_pages,
                    "prev_page": prev_page,
                    "next_page": next_page,
                    "data": response_data,
                },
                status=200,
            )

        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)


class AdjustBalanceAdvancedHandlingPaymentAPIView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def get(self, request, pk, *args, **kwargs):
        try:
            adv_payment_obj = AdvancedHandlingPayment.objects.get(pk=pk)
            if adv_payment_obj.is_adjusted:

                if (
                    adv_payment_obj.payment_type == "Finance_Account"
                    and CustomerFinAccount.objects.filter(
                        customer=adv_payment_obj.client
                    ).exists()
                ):
                    fin_account_obj = CustomerFinAccount.objects.filter(
                        customer=adv_payment_obj.client
                    ).first()
                    data = {
                        "account": fin_account_obj,
                        "amount": adv_payment_obj.balance_amount,
                        "linked_payment": adv_payment_obj.pk,
                    }
                    FinAccountTransaction.balance_adjust(**data)

                adv_payment_obj.is_balance_pymt_adjusted = True
                adv_payment_obj.save()
                return Response({"successMsg": "Balance Payment Adjusted"}, status=200)
            return Response(
                {"errorMsg": "Balance Payment Cannot be Adjusted at this moment"},
                status=200,
            )
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)


class AdvancedHandlingPaymentAPIView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def _get_client(self, data):
        """
        Helper function to get client based on location, site, and client name.
        """
        location = Location.objects.get(name=data.get("location"))
        site = Site.objects.get(name=data.get("site"), location=location)
        return Client.objects.get(
            name=data.get("client"), type="Party", location=location, site=site
        )

    def revert_fin_account_transaction(self, adv_payment_obj):
        fin_account_obj = CustomerFinAccount.objects.filter(
            customer=adv_payment_obj.client
        ).first()
        if FinAccountTransaction.objects.filter(
            account=fin_account_obj, linked_payment=adv_payment_obj.pk
        ).exists():
            fin_account_trans_obj = FinAccountTransaction.objects.filter(
                account=fin_account_obj, linked_payment=adv_payment_obj.pk
            ).first()
            fin_account_obj.balance = (
                fin_account_obj.balance + fin_account_trans_obj.amount
            )
            fin_account_obj.save()
            fin_account_trans_obj.delete()

    def create_fin_account_debit_transaction(self, adv_payment_obj):
        fin_account_obj = CustomerFinAccount.objects.filter(
            customer=adv_payment_obj.client
        ).first()

        data = {
            "account": fin_account_obj,
            "amount": adv_payment_obj.original_amount,
            "linked_payment": adv_payment_obj.pk,
        }

        FinAccountTransaction.debit_amount(**data)

    def _validate_payment_data(self, data, adv_payment_obj=None):
        """
        Helper function to validate payment data fields and return the necessary values.
        """
        payment_type = data.get("payment_type")
        cheque_no = data.get("cheque_no", None)
        utr_no = data.get("utr_no", None)
        transaction_id = data.get("transaction_id", None)
        original_amount = data.get("original_amount")
        entry_type = data.get("entry_type")
        bl_no = data.get("bl_no")
        bk_no = data.get("bk_no")
        with_gst = True

        tds = data.get("tds", 1)
        # 20
        size_20_quantity = data.get("size_20_quantity", 0)
        size_20_rate = 0
        size_20_tds_amount = 0
        # 40
        size_40_quantity = data.get("size_40_quantity", 0)
        size_40_rate = 0
        size_40_tds_amount = 0

        client = self._get_client(data)

        if entry_type == "IN":
            if AdvancedHandlingPayment.objects.filter(bl_no=bl_no).exists():
                temp_adv_payment = AdvancedHandlingPayment.objects.filter(bl_no=bl_no)
                if not temp_adv_payment[0].client == client:
                    raise ValueError(
                        f"Client should be [ {temp_adv_payment[0].client.name} ] with given bl_no [ {bl_no} ]"
                    )
        else:
            if AdvancedHandlingPayment.objects.filter(bk_no=bk_no).exists():
                temp_adv_payment = AdvancedHandlingPayment.objects.filter(bk_no=bk_no)
                if not temp_adv_payment[0].client == client:
                    raise ValueError(
                        f"Client should be [ {temp_adv_payment[0].client.name} ] with given bk_no [ {bk_no} ]"
                    )

        # Validate that quantity and original_amount are not zero or negative
        if int(size_20_quantity) == 0 and int(size_40_quantity) == 0:
            raise ValueError("Both Size Quantity cannot be zero")

        if int(size_20_quantity) < 0 or int(size_40_quantity) < 0:
            raise ValueError("Size Quantity cannot be negative")

        if original_amount <= 0:
            raise ValueError("Amount cannot be zero or negative")

        def validate_cheque_or_utr():
            nonlocal cheque_no, utr_no, transaction_id
            transaction_id = None
            if not cheque_no and not utr_no:
                raise ValueError("Please fill cheque_no or utr_no")
            if cheque_no:
                utr_no = None
                if AdvancedHandlingPayment.objects.filter(cheque_no=cheque_no).exists():
                    raise ValueError("Payment already exists with the given cheque_no")
            if utr_no:
                cheque_no = None
                if AdvancedHandlingPayment.objects.filter(utr_no=utr_no).exists():
                    raise ValueError("Payment already exists with the given utr_no")

        def validate_upi():
            nonlocal cheque_no, utr_no, transaction_id
            cheque_no = utr_no = None
            if not transaction_id:
                raise ValueError("Please fill transaction_id")
            if AdvancedHandlingPayment.objects.filter(
                transaction_id=transaction_id
            ).exists():
                raise ValueError("Payment already exists with the given transaction_id")

        def validate_fin_account():
            nonlocal cheque_no, utr_no, transaction_id
            cheque_no = utr_no = transaction_id = None
            if not CustomerFinAccount.objects.filter(customer=client).exists():
                raise ValueError("Customer Finance Account Not Exist")

            fin_account_obj = CustomerFinAccount.objects.filter(customer=client).first()
            if float(fin_account_obj.balance) < float(original_amount):
                raise ValueError("Customer Finance Account balance is not sufficient")
            return fin_account_obj.with_gst

        def validate_original_amt():
            nonlocal size_20_rate, size_20_tds_amount, size_40_rate, size_40_tds_amount
            size_20_original_amount = 0
            size_40_original_amount = 0

            if not int(size_20_quantity) == 0:
                size_20_rate = client.site.size_20_rate
                # tds
                tds_amount = float(size_20_rate) * (int(tds) / 100)
                size_20_tds_amount = round(float(tds_amount) * int(size_20_quantity))

                # tax
                tax = 0.18
                if not with_gst:
                    tax = 0
                tax_amount = float(size_20_rate) * tax
                size_20_total_amount_with_gst = round(
                    float(size_20_rate) + float(tax_amount)
                ) * int(size_20_quantity)
                # cal
                size_20_original_amount = float(size_20_total_amount_with_gst) - float(
                    size_20_tds_amount
                )

            if not int(size_40_quantity) == 0:
                size_40_rate = client.site.size_40_rate
                # tds
                tds_amount = float(size_40_rate) * (int(tds) / 100)
                size_40_tds_amount = round(float(tds_amount) * int(size_40_quantity))

                # tax
                tax = 0.18
                if not with_gst:
                    tax = 0
                tax_amount = float(size_40_rate) * tax
                size_40_total_amount_with_gst = round(
                    float(size_40_rate) + float(tax_amount)
                ) * int(size_40_quantity)
                # cal
                size_40_original_amount = float(size_40_total_amount_with_gst) - float(
                    size_40_tds_amount
                )

            my_original_amount = size_20_original_amount + size_40_original_amount
            if float(original_amount) < float(my_original_amount):
                raise ValueError(
                    "Sorry...Your original_amount is insufficient with respect to the quantity and current rate"
                )

        if adv_payment_obj is None:
            if payment_type in ["Cheque", "NEFT", "RTGS"]:
                validate_cheque_or_utr()
            elif payment_type == "UPI":
                validate_upi()
            elif payment_type == "Finance_Account":
                with_gst = validate_fin_account()
            else:
                pass
        else:
            if adv_payment_obj.payment_type != payment_type:
                if payment_type in ["Cheque", "NEFT", "RTGS"]:
                    validate_cheque_or_utr()
                elif payment_type == "UPI":
                    validate_upi()
                elif payment_type == "Finance_Account":
                    with_gst = validate_fin_account()
                else:
                    pass
            else:
                if payment_type in ["Cheque", "NEFT", "RTGS"]:
                    transaction_id = None
                    if not cheque_no and not utr_no:
                        raise ValueError("Please fill cheque_no or utr_no")
                    if cheque_no and adv_payment_obj.cheque_no != cheque_no:
                        utr_no = None
                        if AdvancedHandlingPayment.objects.filter(
                            cheque_no=cheque_no
                        ).exists():
                            raise ValueError(
                                "Payment already exists with the given cheque_no"
                            )
                    if utr_no and adv_payment_obj.utr_no != utr_no:
                        cheque_no = None
                        if AdvancedHandlingPayment.objects.filter(
                            utr_no=utr_no
                        ).exists():
                            raise ValueError(
                                "Payment already exists with the given utr_no"
                            )
                elif payment_type == "UPI":
                    cheque_no = utr_no = None
                    if not transaction_id:
                        raise ValueError("Please fill transaction_id")
                    if (
                        adv_payment_obj.transaction_id != transaction_id
                        and AdvancedHandlingPayment.objects.filter(
                            transaction_id=transaction_id
                        ).exists()
                    ):
                        raise ValueError(
                            "Payment already exists with the given transaction_id"
                        )

        validate_original_amt()

        return (
            payment_type,
            cheque_no,
            utr_no,
            transaction_id,
            size_20_quantity,
            size_40_quantity,
            original_amount,
            entry_type,
            bl_no,
            bk_no,
            with_gst,
            size_20_rate,
            size_20_tds_amount,
            size_40_rate,
            size_40_tds_amount,
            tds,
        )

    def get(self, request, pk, *args, **kwargs):
        try:
            adv_payment_obj = AdvancedHandlingPayment.objects.get(pk=pk)
            adv_payment_data = adv_payment_obj.get_payment_data()
            pregate_in_list = []
            pregate_out_list = []

            # 20 ---------------------------
            size_20_total_amount_without_tax = 0
            size_20_total_tds = 0
            size_20_total_tax = 0
            # tds
            size_20_tds_rate = float(adv_payment_obj.size_20_rate) * (
                int(adv_payment_obj.tds) / 100
            )
            size_20_total_tds = round(
                float(size_20_tds_rate) * int(adv_payment_obj.size_20_quantity)
            )
            # tax
            size_20_tax = float(adv_payment_obj.size_20_rate) * 0.18
            if not adv_payment_obj.with_gst:
                size_20_tax = 0
            size_20_total_tax = round(
                float(size_20_tax) * int(adv_payment_obj.size_20_quantity)
            )
            # amount_without_tax
            size_20_total_amount_without_tax = float(
                adv_payment_obj.size_20_rate
            ) * int(adv_payment_obj.size_20_quantity)

            # 40 ----------------------------------------
            size_40_total_amount_without_tax = 0
            size_40_total_tds = 0
            size_40_total_tax = 0
            # tds
            size_40_tds_rate = float(adv_payment_obj.size_40_rate) * (
                int(adv_payment_obj.tds) / 100
            )
            size_40_total_tds = round(
                float(size_40_tds_rate) * int(adv_payment_obj.size_40_quantity)
            )
            # tax
            size_40_tax = float(adv_payment_obj.size_40_rate) * 0.18
            if not adv_payment_obj.with_gst:
                size_40_tax = 0
            size_40_total_tax = round(
                float(size_40_tax) * int(adv_payment_obj.size_40_quantity)
            )
            # amount_without_tax
            size_40_total_amount_without_tax = float(
                adv_payment_obj.size_40_rate
            ) * int(adv_payment_obj.size_40_quantity)
            # -------------------------------------------

            # total
            total_amount_without_tax = 0
            total_tax = 0
            total_tds = 0

            total_amount_without_tax = float(size_20_total_amount_without_tax) + float(
                size_40_total_amount_without_tax
            )
            total_tax = float(size_20_total_tax) + float(size_40_total_tax)
            total_tds = float(size_20_total_tds) + float(size_40_total_tds)
            total_amount_with_gst_tds = (
                total_amount_without_tax + total_tax
            ) - total_tds

            if adv_payment_obj.entry_type == "IN":
                pregate_in_list = [
                    each.get_pre_gatein_data()
                    for each in PreGateIn.objects.filter(
                        adv_payment_id=adv_payment_obj.pk
                    )
                    if PreGateIn.objects.filter(
                        adv_payment_id=adv_payment_obj.pk
                    ).exists()
                ]

                # if PreGateIn.objects.filter(adv_payment_id=adv_payment_obj.pk).exists():
                #     count_20 = PreGateIn.objects.filter(
                #         adv_payment_id=adv_payment_obj.pk, size__name="20"
                #     ).count()
                #     count_40 = PreGateIn.objects.filter(
                #         adv_payment_id=adv_payment_obj.pk, size__name="40"
                #     ).count()

            else:
                pregate_out_list = [
                    each.get_pre_gateout_data()
                    for each in PreGateOut.objects.filter(
                        adv_payment_id=adv_payment_obj.pk
                    )
                    if PreGateOut.objects.filter(
                        adv_payment_id=adv_payment_obj.pk
                    ).exists()
                ]

                # if PreGateOut.objects.filter(
                #     adv_payment_id=adv_payment_obj.pk
                # ).exists():
                #     count_20 = PreGateOut.objects.filter(
                #         adv_payment_id=adv_payment_obj.pk,
                #         stock__container__size__name="20",
                #     ).count()
                #     count_40 = PreGateOut.objects.filter(
                #         adv_payment_id=adv_payment_obj.pk,
                #         stock__container__size__name="40",
                #     ).count()

            adv_payment_data["total_amount_without_tax"] = total_amount_without_tax
            adv_payment_data["total_tax"] = total_tax
            adv_payment_data["total_tds"] = total_tds
            adv_payment_data["total_amount_with_gst_tds"] = total_amount_with_gst_tds

            data = {
                "adv_payment_data": adv_payment_data,
                "pregate_in_list": pregate_in_list,
                "pregate_out_list": pregate_out_list,
            }
            return Response(data, status=200)
        except AdvancedHandlingPayment.DoesNotExist:
            return Response({"errorMsg": "Payment record not found"}, status=200)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)

    def put(self, request, pk, *args, **kwargs):
        try:
            adv_payment_qs = AdvancedHandlingPayment.objects.filter(
                pk=pk, is_locked=False
            )
            if not adv_payment_qs.exists():
                return Response(
                    {"errorMsg": "Can't Update Data, Advance Payment is Locked!!!"},
                    status=200,
                )

            adv_payment_obj = adv_payment_qs
            data = request.data

            # Validate input data
            (
                payment_type,
                cheque_no,
                utr_no,
                transaction_id,
                size_20_quantity,
                size_40_quantity,
                original_amount,
                entry_type,
                bl_no,
                bk_no,
                with_gst,
                size_20_rate,
                size_20_tds_amount,
                size_40_rate,
                size_40_tds_amount,
                tds,
            ) = self._validate_payment_data(data, adv_payment_obj[0])

            with transaction.atomic():
                client = self._get_client(data)
                tz = timezone.get_current_timezone()
                main_data = {
                    "client": client,
                    "bank_name": data.get("bank_name", ""),
                    "account_name": data.get("account_name", ""),
                    "account_no": data.get("account_no", ""),
                    "cheque_no": cheque_no,
                    "utr_no": utr_no,
                    "transaction_id": transaction_id,
                    "payment_type": payment_type,
                    "quantity": int(size_20_quantity) + int(size_40_quantity),
                    "remaining": int(size_20_quantity) + int(size_40_quantity),
                    "original_amount": original_amount,
                    "remaining_amount": original_amount,
                    "entry_type": entry_type,
                    "with_gst": with_gst,
                    "updated_at": datetime.datetime.now().astimezone(tz),
                    "tds": tds,
                    # 20
                    "size_20_rate": size_20_rate,
                    "size_20_quantity": size_20_quantity,
                    "size_20_tds_amount": size_20_tds_amount,
                    # 40
                    "size_40_rate": size_40_rate,
                    "size_40_quantity": size_40_quantity,
                    "size_40_tds_amount": size_40_tds_amount,
                    "remarks": data.get("remarks"),
                }

                if adv_payment_obj[0].payment_type != payment_type:
                    if adv_payment_obj[0].payment_type == "Finance_Account":
                        self.revert_fin_account_transaction(adv_payment_obj[0])
                    if payment_type == "Finance_Account":
                        self.create_fin_account_debit_transaction(adv_payment_obj[0])

                if entry_type == "IN":
                    main_data["bl_no"] = bl_no
                    adv_payment_obj[0].bk_no = None
                    adv_payment_obj[0].save()
                else:
                    main_data["bk_no"] = bk_no
                    adv_payment_obj[0].bl_no = None
                    adv_payment_obj[0].save()

                adv_payment_obj.update(**main_data)
                return Response({"successMsg": "Data Saved"}, status=200)

        except (Client.DoesNotExist, Location.DoesNotExist, Site.DoesNotExist):
            return Response(
                {"errorMsg": "Invalid client, location, or site details"}, status=200
            )
        except ValueError as e:
            return Response({"errorMsg": str(e)}, status=200)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)

    def post(self, request, *args, **kwargs):
        try:
            data = request.data

            # Validate input data
            (
                payment_type,
                cheque_no,
                utr_no,
                transaction_id,
                size_20_quantity,
                size_40_quantity,
                original_amount,
                entry_type,
                bl_no,
                bk_no,
                with_gst,
                size_20_rate,
                size_20_tds_amount,
                size_40_rate,
                size_40_tds_amount,
                tds,
            ) = self._validate_payment_data(data)

            with transaction.atomic():
                client = self._get_client(data)
                main_data = {
                    "client": client,
                    "bank_name": data.get("bank_name", ""),
                    "account_name": data.get("account_name", ""),
                    "account_no": data.get("account_no", ""),
                    "cheque_no": cheque_no,
                    "utr_no": utr_no,
                    "transaction_id": transaction_id,
                    "payment_type": payment_type,
                    "quantity": int(size_20_quantity) + int(size_40_quantity),
                    "original_amount": original_amount,
                    "location": client.location,
                    "site": client.site,
                    "entry_type": entry_type,
                    "tds": tds,
                    "with_gst": with_gst,
                    # 20
                    "size_20_rate": size_20_rate,
                    "size_20_quantity": size_20_quantity,
                    "size_20_tds_amount": size_20_tds_amount,
                    # 40
                    "size_40_rate": size_40_rate,
                    "size_40_quantity": size_40_quantity,
                    "size_40_tds_amount": size_40_tds_amount,
                    "remarks": data.get("remarks"),
                }

                adv_payment_obj = None
                if entry_type == "IN":
                    main_data["bl_no"] = bl_no
                    adv_payment_obj = AdvancedHandlingPayment.create_in_adv_payment(
                        **main_data
                    )
                else:
                    main_data["bk_no"] = bk_no
                    adv_payment_obj = AdvancedHandlingPayment.create_out_adv_payment(
                        **main_data
                    )

                if payment_type == "Finance_Account":
                    self.create_fin_account_debit_transaction(adv_payment_obj)

                return Response({"successMsg": "Data Saved"}, status=200)

        except (Client.DoesNotExist, Location.DoesNotExist, Site.DoesNotExist):
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response(
                {"errorMsg": "Invalid client, location, or site details"}, status=200
            )
        except ValueError as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)


class DeleteAdvanceHandlingPaymentAPIView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def revert_fin_account_transaction(self, adv_payment_obj):
        fin_account_obj = CustomerFinAccount.objects.filter(
            customer=adv_payment_obj.client
        ).first()
        if FinAccountTransaction.objects.filter(
            account=fin_account_obj, linked_payment=adv_payment_obj.pk
        ).exists():
            fin_account_trans_obj = FinAccountTransaction.objects.filter(
                account=fin_account_obj, linked_payment=adv_payment_obj.pk
            ).first()
            fin_account_obj.balance = (
                fin_account_obj.balance + fin_account_trans_obj.amount
            )
            fin_account_obj.save()
            fin_account_trans_obj.delete()

    def post(self, request, *args, **kwargs):

        try:
            data = request.data
            adv_payment_pk_list = data["pk_list"]
            adv_payment_query = AdvancedHandlingPayment.objects.filter(
                pk__in=adv_payment_pk_list, is_locked=False
            )
            if len(adv_payment_query) == 0:
                return Response({"errorMsg": f"Advance Payment Not Found"}, status=200)
            for adv_payment in adv_payment_query:
                if adv_payment.payment_type == "Finance_Account":
                    self.revert_fin_account_transaction(adv_payment_obj=adv_payment)
            adv_payment_query.delete()
            return Response({"successMsg": "Advanced Payment Delete"}, status=200)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)


class UpdateRemarksAdvLoloPymtAPIView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def put(self, request, pk, *args, **kwargs):
        try:
            adv_payment_obj = AdvancedHandlingPayment.objects.get(pk=pk)
            adv_payment_obj.remarks = request.data.get("remarks")
            adv_payment_obj.save()
            return Response({"successMsg": "Data Saved"}, status=200)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)


class PregateInListAPIView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        try:
            pg_no = request.data["pg_no"]
            on_page_data = request.data["on_page_data"]
            client_name = request.data.get("client", None)
            shipping_line = request.data.get("shipping_line", None)
            location_name = request.data.get("location", None)
            site_name = request.data.get("site", None)
            container_no = request.data.get("container_no", None)
            do_validity = request.data.get("do_validity_date", None)
            do_validity_date_from_date = do_validity.get("from_date", None)
            do_validity_date_to_date = do_validity.get("to_date", None)
            # on_hold = request.data.get("on_hold", None)
            validity_expired = request.data.get("validity_expired", None)
            is_gatein_done = request.data.get("is_gatein_done", None)
            is_survey_done = request.data.get("is_survey_done", None)

            filters = Q()
            # if on_hold:
            #     filters &= Q(on_hold=on_hold)
            if validity_expired:
                filters &= Q(validity_expired=validity_expired)
            if is_gatein_done:
                filters &= Q(is_gatein_done=is_gatein_done)
            else:
                filters &= Q(is_gatein_done=False)
            if is_survey_done:
                filters &= Q(is_survey_done=is_survey_done)
            else:
                filters &= Q(is_survey_done=False)
            if client_name:
                filters &= Q(client__name__icontains=client_name)
            if shipping_line:
                filters &= Q(shipping_line__icontains=shipping_line)
            if location_name:
                filters &= Q(location__name=location_name)
            if site_name:
                filters &= Q(site__name=site_name)
            if container_no:
                filters &= Q(container_no__in=container_no)
            if do_validity_date_from_date and do_validity_date_to_date:
                filters &= Q(
                    do_validity_in_date__range=[
                        do_validity_date_from_date,
                        do_validity_date_to_date,
                    ]
                )

            pregatein_qs = PreGateIn.objects.filter(filters)

            # pagination
            paginator = Paginator(pregatein_qs, on_page_data)
            current_page = paginator.page(pg_no)
            no_of_data_count = paginator.count
            no_of_pages = paginator.num_pages
            on_page_data_count = current_page.object_list.count()
            prev_page = (
                current_page.previous_page_number()
                if current_page.has_previous()
                else None
            )
            next_page = (
                current_page.next_page_number() if current_page.has_next() else None
            )

            response_data = [
                {
                    **pregatein.get_pre_gatein_data(),
                    "sr_no": current_page.start_index() + index,
                }
                for index, pregatein in enumerate(current_page.object_list)
            ]

            return Response(
                {
                    "no_of_data": no_of_data_count,
                    "on_page_data": on_page_data_count,
                    "total_pages": no_of_pages,
                    "prev_page": prev_page,
                    "next_page": next_page,
                    "data": response_data,
                },
                status=200,
            )

        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)


class PreGateInBulkUploadAPIView(APIView):
    """Get Function will provide the sample file"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def get(self, request, *args, **kwargs):
        try:
            file_path = os.path.join(
                BASE_DIR, "sample_stock/sample_pre_gatein_upload.xlsx"
            )
            with open(file_path, "rb") as file:
                response = HttpResponse(file.read(), content_type="application/xlsx")
                response["Content-Disposition"] = (
                    'attachment; filename="sample_pre_gatein_upload.xlsx"'
                )
            return response
        except:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": "Data Not Found"}, status=200)

    def post(self, request, *args, **kwargs):
        file = request.data.get("file")
        adv_payment_id = request.data.get("adv_payment_id")
        adv_pymt_obj = AdvancedHandlingPayment.objects.get(pk=adv_payment_id)

        try:
            location = adv_pymt_obj.location
            site = adv_pymt_obj.site
            result = extract_pre_gate_in_excel_data(
                input_excel=file,
                location=location,
                site=site,
                adv_payment_id=adv_payment_id,
            )
            if result in ["Header Not Found", False]:
                return Response(
                    {"errorMsg": "Data file is corrupted, unable to import data"},
                    status=200,
                )
            return Response(result, status=200)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": f"Invalid Data Provided [{e}]"}, status=200)


class PreGateInDOUploadAPIView(APIView):
    """Upload DO for Scaning"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        try:
            file = request.data.get("file")
            adv_payment_id = request.data.get("adv_payment_id")
            adv_pymt_obj = AdvancedHandlingPayment.objects.get(pk=adv_payment_id)
            extracted_data = do_data_extraction(adv_pymt_obj=adv_pymt_obj, file=file)
            return Response(extracted_data, status=200)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": f"Invalid Data Provided [{e}]"}, status=200)


class PreGateInBulkImportAPIView(APIView):
    """Post Function will Import the data"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def get_comman_object_data(self, data, location, site):
        client = Client.objects.get(
            name=data.get("client"), type="Line", location=location, site=site
        )
        return {
            "location": location,
            "site": site,
            "type_obj": ContainerType.objects.get(name=data.get("type")),
            "size_obj": ContainerSize.objects.get(name=data.get("size")),
            "client": client,
            "shipping_line": client.ref_code,
        }

    def handle_advanced_payment(self, adv_payment_id, size_name, site):
        adv_pymt = give_ideal_advance_payment_obj(
            adv_payment_id=adv_payment_id, size=size_name
        )
        if isinstance(adv_pymt, str):
            return Response({"errorMsg": adv_pymt}, status=200)
        adv_pymt.manage_qty_n_payment(size=size_name)
        return adv_pymt

    def do_validity_check(self, data_object):
        tz = timezone.get_current_timezone()
        now = datetime.datetime.now().astimezone(tz)
        do_date = datetime.datetime.combine(
            data_object.do_validity_in_date, data_object.do_validity_in_time
        ).astimezone(tz)
        if now > do_date:
            data_object.validity_expired = True
            data_object.on_hold = True
            data_object.save()
        return True

    def adjust_quantity_n_balance(self, payment_record):
        if float(payment_record.remaining_amount) > 0:
            payment_record.quantity = payment_record.quantity - payment_record.remaining
            payment_record.remaining = 0
            payment_record.balance_amount = float(payment_record.remaining_amount)
            payment_record.remaining_amount = 0
            payment_record.is_adjusted = True
            payment_record.save()
        else:
            payment_record.quantity = payment_record.quantity - payment_record.remaining
            payment_record.remaining = 0
            payment_record.is_adjusted = True
            payment_record.save()

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data_list = request.data["importable_data"]
            location = Location.objects.get(name=request.data.get("location"))
            site = Site.objects.get(name=request.data.get("site"), location=location)
            adv_payment_id = request.data.get("adv_payment_id")
            # adv_payment_obj = AdvancedHandlingPayment.objects.filter(pk=adv_payment_id)

            for data in data_list:
                with transaction.atomic():
                    container_no = data.get("container_no")
                    gatein_exists = Container.objects.filter(
                        container_no=container_no, status="IN", location=location
                    ).exists()
                    pregatein_exists = PreGateIn.objects.filter(
                        location=location,
                        container_no=container_no,
                        is_gatein_done=False,
                    ).exists()
                    if not gatein_exists and not pregatein_exists:
                        comman_objects = self.get_comman_object_data(
                            data, location, site
                        )
                        adv_pymt = self.handle_advanced_payment(
                            adv_payment_id=adv_payment_id,
                            size_name=comman_objects["size_obj"].name,
                            site=comman_objects["site"],
                        )
                        lolo_amount = get_total_site_lolo_amount(
                            comman_objects["size_obj"].name,
                            count=1,
                            payment_data=adv_pymt,
                            with_gst=adv_pymt.with_gst,
                        )

                        main_data = {
                            "client": comman_objects["client"],
                            "type": comman_objects["type_obj"],
                            "size": comman_objects["size_obj"],
                            "container_no": container_no,
                            "shipping_line": comman_objects["shipping_line"],
                            "do_validity_in_date": datetime.datetime.strptime(
                                data.get("do_validity_in_date"), "%d_%m_%Y"
                            ).date(),
                            "do_validity_in_time": datetime.datetime.strptime(
                                data.get("do_validity_in_time"), "%H_%M"
                            ).time(),
                            "bl_no": adv_pymt.bl_no,
                            "consignee": (
                                None
                                if len(data.get("consignee")) == 0
                                else data.get("consignee")
                            ),
                            "shipper": (
                                None
                                if len(data.get("shipper")) == 0
                                else data.get("shipper")
                            ),
                            "cargo": (
                                None
                                if len(data.get("cargo")) == 0
                                else data.get("cargo")
                            ),
                            "remarks": (
                                None
                                if len(data.get("remarks")) == 0
                                else data.get("remarks")
                            ),
                            "arrived": (
                                None
                                if len(data.get("arrived")) == 0
                                else data.get("arrived")
                            ),
                            "location": comman_objects["location"],
                            "site": comman_objects["site"],
                            "lolo_amount": lolo_amount,
                        }

                        pregatein = PreGateIn.create_pre_gatein(**main_data)
                        adv_pymt.is_locked = True
                        adv_pymt.save()
                        pregatein.adv_payment_id = adv_pymt.pk
                        pregatein.save()
                        self.do_validity_check(pregatein)
                        if adv_pymt.remaining <= 0:
                            self.adjust_quantity_n_balance(adv_pymt)
            return Response({"successMsg": "Data Saved"}, status=200)
        except Exception as e:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return Response({"errorMsg": f"{e}"}, status=200)


class RejectedPreGateInDataApiView(APIView):
    """Post Function will get excel with rejected pregateIn data"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data = request.data["rejected_data"]
            faults = request.data["faults"]

            stock_df_data = {
                "client": [each["client"] for each in data],
                "container_no": [each["container_no"] for each in data],
                "type": [each["type"] for each in data],
                "size": [each["size"] for each in data],
                "arrived": [each["arrived"] for each in data],
                "consignee": [each["consignee"] for each in data],
                "shipper": [each["shipper"] for each in data],
                "cargo": [each["cargo"] for each in data],
                "remarks": [each["remarks"] for each in data],
            }

            faults_list = []
            for row, messages in faults.items():
                for msg in messages:
                    faults_list.append({"row": row, "message": msg})

            temp_dir = os.path.join(BASE_DIR, "temp/sample_stock/")
            os.makedirs(temp_dir, exist_ok=True)

            new_temp_file_path = os.path.join(temp_dir, "sample_pre_gatein_upload.xlsx")
            temp_file_path = os.path.join(
                BASE_DIR, "sample_stock/sample_pre_gatein_upload.xlsx"
            )

            with open(temp_file_path, "rb") as temp:
                old_temp_data = temp.read()

            with open(new_temp_file_path, "wb") as f:
                f.write(old_temp_data)

            # New Version code
            book = load_workbook(new_temp_file_path)
            with pd.ExcelWriter(
                new_temp_file_path,
                engine="openpyxl",
                mode="a",
                if_sheet_exists="replace",
            ) as writer:

                stock_df = pd.DataFrame(stock_df_data)
                stock_df.to_excel(writer, sheet_name="pre_gatein", index=False)

                faults_df = pd.DataFrame(faults_list)
                faults_df.to_excel(writer, sheet_name="faults", index=False)

            # Old Version code
            # book = load_workbook(new_temp_file_path)
            # writer = pd.ExcelWriter(new_temp_file_path, engine="openpyxl")
            # writer.book = book
            # writer.sheets = dict((ws.title, ws) for ws in book.worksheets)
            # stock_df = pd.DataFrame(stock_df_data)
            # stock_df.to_excel(
            #     writer, sheet_name="pre_gatein", startrow=0, startcol=0, index=False
            # )
            # faults_df = pd.DataFrame(faults)
            # faults_df.to_excel(
            #     writer, sheet_name="faults", startrow=0, startcol=0, index=False
            # )
            # fault_sheet = book.get_sheet_by_name("faults")
            # writer.save()

            with open(new_temp_file_path, "rb") as temp:
                file_response = HttpResponse(
                    temp.read(), content_type=f"application/xlsx"
                )
                file_response["Content-Disposition"] = (
                    f'attachment; filename="sample_pre_gatein_upload.xlsx"'
                )
                os.remove(new_temp_file_path)
            return file_response

        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response(
                {"errorMsg": f"Invalid Data Provided [ {str(e)} ]"}, status=200
            )


class PreGateInAPIView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def validate_container_no(self, container_no, location):
        if not container_no:
            return Response({"errorMsg": "Please Enter Container_no"}, status=200)

        if len(container_no) != 11 or not check_char_digit(container_no):
            return Response(
                {
                    "errorMsg": "Invalid Entry, "
                    "Please Enter Container_no having first 4 uppercase alphabets and "
                    "rest 7 digits"
                },
                status=200,
            )

        # validation = check_digit(container_no=container_no)
        gatein_exists = Container.objects.filter(
            container_no=container_no, status="IN", location=location
        ).exists()
        pregatein_exists = PreGateIn.objects.filter(
            location=location,
            container_no=container_no,
            is_gatein_done=False,
        ).exists()
        if gatein_exists or pregatein_exists:
            return Response(
                {
                    "errorMsg": "Container_no already exists in system with status 'IN' or in PreGateIn !!!"
                },
                status=200,
            )
        return True

    def get_comman_object_data(self, data):
        location = Location.objects.get(name=data.get("location"))
        site = Site.objects.get(name=data.get("site"), location=location)
        client = Client.objects.get(
            name=data.get("client"), type="Line", location=location, site=site
        )
        return {
            "location": location,
            "site": site,
            "type_obj": ContainerType.objects.get(name=data.get("type")),
            "size_obj": ContainerSize.objects.get(name=data.get("size")),
            "client": client,
            "shipping_line": client.ref_code,
        }

    def do_validity_check(self, data_object):
        tz = timezone.get_current_timezone()
        now = datetime.datetime.now().astimezone(tz)
        do_date = datetime.datetime.combine(
            data_object.do_validity_in_date, data_object.do_validity_in_time
        ).astimezone(tz)
        if now > do_date:
            data_object.validity_expired = True
            data_object.on_hold = True
            data_object.save()
        return True

    def handle_advanced_payment(self, adv_payment_id, size_name, site, pregatein):
        adv_pymt = AdvancedHandlingPayment.objects.get(pk=adv_payment_id)

        if pregatein and pregatein[0].size.name != size_name:
            adv_pymt.manage_qty_n_payment(
                size=pregatein[0].size.name,
                increment=True,
                total_amount=pregatein[0].lolo_amount,
            )
            validated_adv_pymt = give_ideal_advance_payment_obj(
                adv_payment_id=adv_payment_id, size=size_name
            )
            if isinstance(validated_adv_pymt, str):
                return Response({"errorMsg": validated_adv_pymt}, status=200)
            validated_adv_pymt.manage_qty_n_payment(size=size_name)
            return validated_adv_pymt

        if not pregatein:
            validated_adv_pymt = give_ideal_advance_payment_obj(
                adv_payment_id=adv_payment_id, size=size_name
            )
            if isinstance(validated_adv_pymt, str):
                return Response({"errorMsg": validated_adv_pymt}, status=200)
            validated_adv_pymt.manage_qty_n_payment(size=size_name)
            return validated_adv_pymt

        return adv_pymt

    def adjust_quantity_n_balance(self, payment_record):
        if float(payment_record.remaining_amount) > 0:
            payment_record.quantity = payment_record.quantity - payment_record.remaining
            payment_record.remaining = 0
            payment_record.balance_amount = float(payment_record.remaining_amount)
            payment_record.remaining_amount = 0
            payment_record.is_adjusted = True
            payment_record.save()
        else:
            payment_record.quantity = payment_record.quantity - payment_record.remaining
            payment_record.remaining = 0
            payment_record.is_adjusted = True
            payment_record.save()

    def handle_pre_gatein(self, request, pregatein=None):
        try:
            with transaction.atomic():
                data = request.data
                adv_payment_id = data.get("adv_payment_id")
                container_no = data.get("container_no")
                comman_objects = self.get_comman_object_data(data)
                if pregatein:
                    if not pregatein[0].container_no == container_no:
                        validation = self.validate_container_no(
                            container_no=container_no,
                            location=comman_objects["location"],
                        )

                        if isinstance(validation, Response):
                            return validation
                else:
                    validation = self.validate_container_no(
                        container_no=container_no, location=comman_objects["location"]
                    )

                    if isinstance(validation, Response):
                        return validation

                adv_pymt = self.handle_advanced_payment(
                    adv_payment_id=adv_payment_id,
                    size_name=comman_objects["size_obj"].name,
                    site=comman_objects["site"],
                    pregatein=pregatein,
                )

                if isinstance(adv_pymt, Response):
                    return adv_pymt
                if not adv_pymt:
                    return Response({"errorMsg": "Something Went Wrong!"}, status=200)

                lolo_amount = get_total_site_lolo_amount(
                    comman_objects["size_obj"].name,
                    count=1,
                    payment_data=adv_pymt,
                    with_gst=adv_pymt.with_gst,
                )

                main_data = {
                    "client": comman_objects["client"],
                    "type": comman_objects["type_obj"],
                    "size": comman_objects["size_obj"],
                    "container_no": container_no,
                    "shipping_line": comman_objects["shipping_line"],
                    "do_validity_in_date": datetime.datetime.strptime(
                        data.get("do_validity_in_date"), "%Y-%m-%d"
                    ).date(),
                    "do_validity_in_time": datetime.datetime.strptime(
                        data.get("do_validity_in_time"), "%H:%M"
                    ).time(),
                    "consignee": data.get("consignee"),
                    "shipper": data.get("shipper"),
                    "bl_no": data.get("bl_no"),
                    "cargo": data.get("cargo"),
                    "remarks": data.get("remarks"),
                    "arrived": data.get("arrived"),
                    "lolo_amount": lolo_amount,
                    "location": comman_objects["location"],
                    "site": comman_objects["site"],
                }

                if pregatein:
                    if (
                        not pregatein[0].on_hold
                        and not pregatein[0].validity_expired
                        and not pregatein[0].is_gatein_done
                    ):
                        pregatein.update(**main_data)
                        adv_pymt.is_locked = True
                        adv_pymt.save()
                        pregatein[0].adv_payment_id = adv_pymt.pk
                        pregatein[0].save()
                        self.do_validity_check(pregatein[0])
                    else:
                        return Response(
                            {"errorMsg": "Sorry Can't Update, PregateIn is Locked !!!"},
                            status=200,
                        )
                else:
                    pregatein = PreGateIn.create_pre_gatein(**main_data)
                    adv_pymt.is_locked = True
                    adv_pymt.save()
                    pregatein.adv_payment_id = adv_pymt.pk
                    pregatein.save()
                    self.do_validity_check(pregatein)

                if adv_pymt.remaining <= 0:
                    self.adjust_quantity_n_balance(adv_pymt)

                return Response({"successMsg": "Data Saved"}, status=200)

        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)

    def get(self, request, pk, *args, **kwargs):
        try:
            pregatein = PreGateIn.objects.get(pk=pk)
            return Response(pregatein.get_pre_gatein_data(), status=200)
        except PreGateIn.DoesNotExist:
            return Response({"errorMsg": "PreGateIn record not found"}, status=200)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)

    def put(self, request, pk, *args, **kwargs):
        try:
            pregatein = PreGateIn.objects.filter(pk=pk)
            return self.handle_pre_gatein(request, pregatein=pregatein)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)

    def post(self, request, *args, **kwargs):
        try:
            return self.handle_pre_gatein(request)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)


class UpgradePregateInValidityApiView(APIView):
    """Post Function will upgrade the validity"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        try:
            data = request.data
            pregatein_pk_list = data["pk_list"]
            date_str = request.data["date"]
            time_str = request.data["time"]
            date = datetime.datetime.strptime(date_str, "%Y-%m-%d").date()
            time = datetime.datetime.strptime(time_str, "%H:%M").time()
            tz = timezone.get_current_timezone()
            date_time = datetime.datetime.combine(date, time).astimezone(tz)

            pregatein_query = PreGateIn.objects.filter(
                pk__in=pregatein_pk_list,
                is_gatein_done=False,
            )

            if len(pregatein_query) == 0:
                return Response({"errorMsg": f"PregateIN Not Found"}, status=200)

            for each in pregatein_query:
                # tz = timezone.get_current_timezone()
                now = datetime.datetime.now().astimezone(tz)
                if date_time > now:
                    _ = DoValidityUpgradeHistory.objects.get_or_create(
                        pregate_in_id=each.pk,
                        container_no=each.container_no,
                        old_do_validity_date=each.do_validity_in_date,
                        old_do_validity_time=each.do_validity_in_time,
                        new_do_validity_date=date,
                        new_do_validity_time=time,
                        location=each.location,
                        site=each.site,
                        entry_type="IN",
                    )
                    each.upgrade_do_validity(date=date, time=time)
            return Response({"sucessMsg": f"PregateIN Validity upgraded"}, status=200)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)


class DeletePregateInApiView(APIView):
    """Post Function will delete pregateIN"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        try:
            data = request.data
            pregatein_pk_list = data["pk_list"]
            pregatein_query = PreGateIn.objects.filter(
                pk__in=pregatein_pk_list,
                is_gatein_done=False,
            )

            if len(pregatein_query) == 0:
                return Response({"errorMsg": f"PregateIN Not Found"}, status=200)

            adv_pymt_list = list(
                set(pregatein_query.values_list("adv_payment_id", flat=True))
            )

            for each in pregatein_query:
                adv_pymt_obj = AdvancedHandlingPayment.objects.get(
                    pk=each.adv_payment_id
                )
                adv_pymt_obj.manage_qty_n_payment(
                    size=each.size.name,
                    increment=True,
                    total_amount=each.lolo_amount,
                )
                each.delete()

            for pk in adv_pymt_list:
                if not PreGateIn.objects.filter(adv_payment_id=pk).exists():
                    adv_pymt_obj = AdvancedHandlingPayment.objects.get(pk=pk)
                    adv_pymt_obj.is_locked = False
                    adv_pymt_obj.save()

            return Response({"sucessMsg": f"PregateIN Deleted"}, status=200)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)


class BulkUploadExcelPregateInToInApiView(APIView):
    """Post Function will get excel with pregateIn data to bulk IN"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data = request.data
            pregatein_pk_list = data["pk_list"]
            pregatein_query = PreGateIn.objects.filter(
                pk__in=pregatein_pk_list,
                validity_expired=False,
                is_gatein_done=False,
                on_hold=False,
            )

            if len(pregatein_query) == 0:
                return Response({"errorMsg": f"PregateIN Not Found"}, status=200)

            shipping_line = []
            client = []
            c_type = []
            c_size = []
            container_no = []
            gross_wt = []
            tare_wt = []
            manufacturing_date = []
            in_date = []
            in_time = []
            condition = []
            grade = []
            arrived = []
            transporter = []
            vehicle_no = []
            shipper = []
            vessel = []
            voyage = []
            cargo = []

            for each in pregatein_query:
                tz = timezone.get_current_timezone()
                now = datetime.datetime.now().astimezone(tz)
                shipping_line.append(
                    each.shipping_line if each.shipping_line is not None else "_"
                )
                client.append(each.client.name if each.client is not None else "_")
                c_type.append(each.type.name if each.type is not None else "_")
                c_size.append(each.size.name if each.size is not None else "_")
                container_no.append(
                    each.container_no if each.container_no is not None else "_"
                )
                gross_wt.append("_")
                tare_wt.append("_")
                manufacturing_date.append("_")
                in_date.append(now.date().strftime("%d_%m_%Y"))
                in_time.append(now.time().strftime("%H_%M"))
                condition.append("_")
                grade.append("_")
                arrived.append(each.arrived)
                transporter.append("_")
                vehicle_no.append("_")
                shipper.append(each.shipper if each.shipper is not None else "_")
                vessel.append("_")
                voyage.append("_")
                cargo.append(each.cargo if each.cargo is not None else "_")

            stock_df_data = {
                "shipping_line": shipping_line,
                "client": client,
                "type": c_type,
                "size": c_size,
                "container_no": container_no,
                "gross_wt": gross_wt,
                "tare_wt": tare_wt,
                "manufacturing_date": manufacturing_date,
                "in_date": in_date,
                "in_time": in_time,
                "condition": condition,
                "grade": grade,
                "arrived": arrived,
                "transporter": transporter,
                "vehicle_no": vehicle_no,
                "shipper": shipper,
                "vessel": vessel,
                "voyage": voyage,
                "cargo": cargo,
            }
            if not os.path.exists(os.path.join(BASE_DIR, "temp/sample_stock/")):
                os.makedirs(os.path.join(BASE_DIR, "temp/sample_stock/"))
            new_temp_file_path = os.path.join(
                BASE_DIR, f"temp/sample_stock/sample_stock_upload.xlsx"
            )
            temp_file_path = os.path.join(
                BASE_DIR, "sample_stock/sample_stock_upload.xlsx"
            )
            old_temp_data = None
            with open(temp_file_path, "rb") as temp:
                old_temp_data = temp.read()
            with open(new_temp_file_path, "wb") as f:
                f.write(old_temp_data)

            # New Version
            book = load_workbook(new_temp_file_path)
            with pd.ExcelWriter(
                new_temp_file_path,
                engine="openpyxl",
                mode="a",
                if_sheet_exists="replace",
            ) as writer:
                stock_df = pd.DataFrame(stock_df_data)
                stock_df.to_excel(writer, sheet_name="stock", index=False)

            # Old Version
            # book = load_workbook(new_temp_file_path)
            # writer = pd.ExcelWriter(new_temp_file_path, engine="openpyxl")
            # writer.book = book
            # writer.sheets = dict((ws.title, ws) for ws in book.worksheets)
            # stock_df = pd.DataFrame(stock_df_data)
            # stock_df.to_excel(
            #     writer, sheet_name="stock", startrow=0, startcol=0, index=False
            # )
            # writer.save()

            with open(new_temp_file_path, "rb") as temp:
                file_response = HttpResponse(
                    temp.read(), content_type=f"application/xlsx"
                )
                file_response["Content-Disposition"] = (
                    f'attachment; filename="sample_stock_upload.xlsx"'
                )
                os.remove(new_temp_file_path)
            return file_response
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": f"Invalid Data Provided [ {e} ]"}, status=200)


class PreGateInContainerInfoAPIView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User", "Surveyor"]

    def post(self, request, *args, **kwargs):

        try:
            data = request.data
            location = Location.objects.get(name=data.get("location"))
            site = Site.objects.get(name=data.get("site"), location=location)
            container_no = request.data.get("container_no", None)
            pregatein_obj = PreGateIn.objects.filter(
                container_no=container_no,
                location=location,
                site=site,
                on_hold=False,
                validity_expired=False,
                is_gatein_done=False,
            ).latest("pk")
            main_data = pregatein_obj.get_pre_gatein_data()
            adv_payment = AdvancedHandlingPayment.objects.get(
                pk=main_data["adv_payment_id"]
            )
            main_data["apply_charges"] = "Party"
            main_data["customer_name"] = adv_payment.client.name
            main_data["payment_type"] = "Advance"
            return Response(main_data, status=200)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)


class PreGateOutAPIView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def get_comman_object_data(self, data):
        location = Location.objects.get(name=data.get("location"))
        site = Site.objects.get(name=data.get("site"), location=location)

        if not ContainerStock.objects.filter(
            container__container_no=data.get("container_no"),
            container__location=location,
            container__site=site,
            container__status="IN",
            container_status="IN",
            # container__is_available=True,
        ).exists():
            return Response(
                {"errorMsg": "Please Enter Available Container_no"}, status=200
            )

        stock = ContainerStock.objects.get(
            container__container_no=data.get("container_no"),
            container__location=location,
            container__site=site,
            container__status="IN",
            container_status="IN",
            # container__is_available=True,
        )

        return {"location": location, "site": site, "stock": stock}

    def handle_advanced_payment(self, adv_payment_id, size_name, site, pregateout):
        adv_pymt = AdvancedHandlingPayment.objects.get(pk=adv_payment_id)
        if not pregateout:
            validated_adv_pymt = give_ideal_advance_payment_obj(
                adv_payment_id=adv_payment_id, size=size_name
            )
            if isinstance(validated_adv_pymt, str):
                return Response({"errorMsg": validated_adv_pymt}, status=200)
            validated_adv_pymt.manage_qty_n_payment(size=size_name)
            return validated_adv_pymt
        return adv_pymt

    def do_validity_check(self, data_object):
        tz = timezone.get_current_timezone()
        now = datetime.datetime.now().astimezone(tz)
        do_date = datetime.datetime.combine(
            data_object.do_validity_out_date, data_object.do_validity_out_time
        ).astimezone(tz)
        if now > do_date:
            data_object.validity_expired = True
            data_object.on_hold = True
            data_object.save()
        return True

    def adjust_quantity_n_balance(self, payment_record):
        if float(payment_record.remaining_amount) > 0:
            payment_record.quantity = payment_record.quantity - payment_record.remaining
            payment_record.remaining = 0
            payment_record.balance_amount = float(payment_record.remaining_amount)
            payment_record.remaining_amount = 0
            payment_record.is_adjusted = True
            payment_record.save()
        else:
            payment_record.quantity = payment_record.quantity - payment_record.remaining
            payment_record.remaining = 0
            payment_record.is_adjusted = True
            payment_record.save()

    def handle_pre_gateout(self, request, pregateout=None):
        try:
            with transaction.atomic():
                data = request.data
                adv_payment_id = data["adv_payment_id"]
                if pregateout is None:
                    comman_objects = self.get_comman_object_data(data)
                    pregateout_exists = PreGateOut.objects.filter(
                        stock=comman_objects["stock"],
                        is_gateout_done=False,
                    ).exists()
                    if pregateout_exists:
                        return Response(
                            {
                                "errorMsg": "Container_no already exists in system in PreGateOut !!!"
                            },
                            status=200,
                        )

                lolo_amount = 0
                if pregateout is None:
                    comman_objects = self.get_comman_object_data(data)
                    size_name = comman_objects["stock"].container.size.name
                    site_obj = comman_objects["site"]
                    adv_pymt = self.handle_advanced_payment(
                        adv_payment_id=adv_payment_id,
                        size_name=size_name,
                        site=site_obj,
                        pregateout=pregateout,
                    )

                    if isinstance(adv_pymt, Response):
                        return adv_pymt
                    if not adv_pymt:
                        return Response(
                            {"errorMsg": "Something Went Wrong!"}, status=200
                        )

                    lolo_amount = get_total_site_lolo_amount(
                        size_name,
                        count=1,
                        payment_data=adv_pymt,
                        with_gst=adv_pymt.with_gst,
                    )
                else:
                    size_name = pregateout[0].stock.container.size.name
                    site_obj = pregateout[0].site
                    adv_pymt = self.handle_advanced_payment(
                        adv_payment_id=adv_payment_id,
                        size_name=size_name,
                        site=site_obj,
                        pregateout=pregateout,
                    )

                    if isinstance(adv_pymt, Response):
                        return adv_pymt
                    if not adv_pymt:
                        return Response(
                            {"errorMsg": "Something Went Wrong!"}, status=200
                        )

                    lolo_amount = get_total_site_lolo_amount(
                        size_name,
                        count=1,
                        payment_data=adv_pymt,
                        with_gst=adv_pymt.with_gst,
                    )

                main_data = {
                    "do_validity_out_date": datetime.datetime.strptime(
                        data.get("do_validity_out_date"), "%Y-%m-%d"
                    ).date(),
                    "do_validity_out_time": datetime.datetime.strptime(
                        data.get("do_validity_out_time"), "%H:%M"
                    ).time(),
                    "consignee": data.get("consignee"),
                    "shipper": data.get("shipper"),
                    "bk_no": data.get("bk_no"),
                    "cargo": data.get("cargo"),
                    "remarks": data.get("remarks"),
                    "departed": data.get("departed"),
                    "lolo_amount": lolo_amount,
                }

                if pregateout is None:
                    main_data["stock"] = comman_objects["stock"]
                    main_data["location"] = comman_objects["location"]
                    main_data["site"] = comman_objects["site"]

                if pregateout:
                    if (
                        not pregateout[0].on_hold
                        and not pregateout[0].validity_expired
                        and not pregateout[0].is_gateout_done
                    ):
                        pregateout.update(**main_data)
                        adv_pymt.is_locked = True
                        adv_pymt.save()
                        pregateout[0].adv_payment_id = adv_pymt.pk
                        pregateout[0].save()
                        self.do_validity_check(pregateout[0])
                    else:
                        return Response(
                            {
                                "errorMsg": "Sorry Can't Update, PregateOut is Locked !!!"
                            },
                            status=200,
                        )
                else:
                    pregateout = PreGateOut.create_pre_gateout(**main_data)
                    adv_pymt.is_locked = True
                    adv_pymt.save()
                    pregateout.adv_payment_id = adv_pymt.pk
                    pregateout.save()
                    self.do_validity_check(pregateout)

                if adv_pymt.remaining <= 0:
                    self.adjust_quantity_n_balance(adv_pymt)
                return Response({"successMsg": "Data Saved"}, status=200)

        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)

    def get(self, request, pk, *args, **kwargs):
        try:
            pregateout = PreGateOut.objects.get(pk=pk)
            return Response(pregateout.get_pre_gateout_data(), status=200)
        except PreGateOut.DoesNotExist:
            return Response({"errorMsg": "PreGateOut record not found"}, status=200)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)

    def put(self, request, pk, *args, **kwargs):
        try:
            pregateout = PreGateOut.objects.filter(pk=pk)
            return self.handle_pre_gateout(request, pregateout=pregateout)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)

    def post(self, request, *args, **kwargs):
        try:
            return self.handle_pre_gateout(request)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)


class UpgradePregateOutValidityApiView(APIView):
    """Post Function will upgrade the validity"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        try:
            data = request.data
            pregateout_pk_list = data["pk_list"]
            date_str = request.data["date"]
            time_str = request.data["time"]
            date = datetime.datetime.strptime(date_str, "%Y-%m-%d").date()
            time = datetime.datetime.strptime(time_str, "%H:%M").time()

            tz = timezone.get_current_timezone()
            date_time = datetime.datetime.combine(date, time).astimezone(tz)

            pregateout_query = PreGateOut.objects.filter(
                pk__in=pregateout_pk_list,
                is_gateout_done=False,
            )

            if len(pregateout_query) == 0:
                return Response({"errorMsg": f"PregateOUT Not Found"}, status=200)

            for each in pregateout_query:
                # tz = timezone.get_current_timezone()
                now = datetime.datetime.now().astimezone(tz)
                if date_time > now:
                    _ = DoValidityUpgradeHistory.objects.get_or_create(
                        pregate_out_id=each.pk,
                        container_no=each.stock.container.container_no,
                        old_do_validity_date=each.do_validity_out_date,
                        old_do_validity_time=each.do_validity_out_time,
                        new_do_validity_date=date,
                        new_do_validity_time=time,
                        location=each.location,
                        site=each.site,
                        entry_type="OUT",
                    )
                    each.upgrade_do_validity(date=date, time=time)
            return Response({"sucessMsg": f"PregateOUT Validity upgraded"}, status=200)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)


class DeletePregateOutApiView(APIView):
    """Post Function will delete pregateOUT"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        try:
            data = request.data
            pregateout_pk_list = data["pk_list"]
            pregateout_query = PreGateOut.objects.filter(
                pk__in=pregateout_pk_list,
                is_gateout_done=False,
            )

            if len(pregateout_query) == 0:
                return Response({"errorMsg": f"PregateOUT Not Found"}, status=200)

            adv_pymt_list = list(
                set(pregateout_query.values_list("adv_payment_id", flat=True))
            )

            for each in pregateout_query:
                adv_pymt_obj = AdvancedHandlingPayment.objects.get(
                    pk=each.adv_payment_id
                )
                adv_pymt_obj.manage_qty_n_payment(
                    size=each.stock.container.size.name,
                    increment=True,
                    total_amount=each.lolo_amount,
                )
                each.delete()

            for pk in adv_pymt_list:
                if not PreGateOut.objects.filter(adv_payment_id=pk).exists():
                    adv_pymt_obj = AdvancedHandlingPayment.objects.get(pk=pk)
                    adv_pymt_obj.is_locked = False
                    adv_pymt_obj.save()
            return Response({"sucessMsg": f"PregateOut Deleted"}, status=200)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)


class PreGateOutContainerInfoAPIView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):

        try:
            data = request.data
            location = Location.objects.get(name=data.get("location"))
            site = Site.objects.get(name=data.get("site"), location=location)
            container_no = request.data.get("container_no", None)

            pregateout_obj = PreGateOut.objects.filter(
                stock__container__container_no=container_no,
                location=location,
                site=site,
                on_hold=False,
                validity_expired=False,
                is_gateout_done=False,
            ).latest("pk")
            adv_payment = AdvancedHandlingPayment.objects.get(
                pk=pregateout_obj.adv_payment_id
            )
            pregateout_obj_data = pregateout_obj.get_pre_gateout_data()
            total_amount = get_total_site_lolo_amount(
                size=pregateout_obj.stock.container.size.name,
                count=1,
                payment_data=adv_payment,
                with_gst=adv_payment.with_gst,
            )
            try:
                stock_object = ContainerStock.objects.get(
                    pk=pregateout_obj.stock.pk, container_status="IN"
                )
            except:
                stock_object = None
            container_object = pregateout_obj.stock.container

            line = container_object.client.ref_code
            seal_no_object_list = SealNo.objects.select_related(
                "location", "site"
            ).filter(is_available=True, location=location, site=site, line=line)
            seal_no_list = [each.number for each in seal_no_object_list]
            if (
                not stock_object is None
                and container_object.status == "IN"
                and container_object.in_do_not_lift_queue is False
            ):
                if (
                    stock_object.status == "Available"
                    or stock_object.status == "Alloted"
                ):
                    container_object_data = container_object.get_container()
                    try:
                        manufacturing_date = (
                            datetime.datetime.strptime(
                                container_object_data["manufacturing_date"],
                                "%Y-%m-%d",
                            )
                            .date()
                            .strftime("%d/%m/%Y")
                        )
                    except:
                        manufacturing_date = ""
                    container_object_data["manufacturing_date"] = manufacturing_date
                    gate_in_object_data = stock_object.gate_in.get_gate_in()
                    condition = gate_in_object_data["condition"]
                    try:
                        allotment_object = ContainerAllotment.objects.get(
                            booking_no=stock_object.booking_no
                        )
                        allotment_object_data = allotment_object.get_allotment()
                        booking_no = allotment_object_data["booking_no"]
                        booking_date = allotment_object_data["booking_date"]
                        booking_party = allotment_object_data["booking_party"]
                    except:
                        booking_no = ""
                        booking_date = ""
                        booking_party = ""
                    stock_object_data = stock_object.get_stock()
                    grade = stock_object_data["grade"]
                    gate_out_data = {
                        "booking_no": booking_no,
                        "booking_date": booking_date,
                        "booking_party": booking_party,
                        "seal_no": stock_object_data["seal_no"],
                        "condition": condition,
                        "grade": grade,
                        "consignee": pregateout_obj_data["consignee"],
                        "shipper": pregateout_obj_data["shipper"],
                        "cargo": pregateout_obj_data["cargo"],
                        "remarks": pregateout_obj_data["remarks"],
                        "departed": pregateout_obj_data["departed"],
                    }

                    if not len(gate_out_data["booking_date"]) == 0:
                        booking_date = (
                            datetime.datetime.strptime(
                                gate_out_data["booking_date"], "%Y-%m-%d"
                            )
                            .date()
                            .strftime("%d/%m/%Y")
                        )
                        gate_out_data["booking_date"] = booking_date
                    gih_object = GateInHistory.objects.get(
                        container=container_object, gate_in=stock_object.gate_in
                    )
                    lolo_data = {}
                    lolo_data["apply_charges"] = "Party"
                    lolo_data["customer_name"] = adv_payment.client.name
                    lolo_data["payment_type"] = "Advance"
                    lolo_data["lolo_amount"] = pregateout_obj_data["lolo_amount"]
                    return Response(
                        {
                            "container_data": container_object_data,
                            "gate_out_data": gate_out_data,
                            "lolo_data": lolo_data,
                            "gih_pk": gih_object.pk,
                            "flag": "OUT",
                            "seal_no_list": seal_no_list,
                        },
                        status=200,
                    )
            return Response(
                {"errorMsg": f"Container is not available yet for Out process."},
                status=200,
            )
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)


class PreGateOutPrefillInfoAPIView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):

        try:
            data = request.data
            location = Location.objects.get(name=data.get("location"))
            site = Site.objects.get(name=data.get("site"), location=location)
            container_no = request.data.get("container_no", None)
            if not ContainerStock.objects.filter(
                container__container_no=container_no,
                container__location=location,
                container__site=site,
                container_status="IN",
                # container__is_available=True,
                container__status="IN",
            ).exists():
                return Response(
                    {"errorMsg": "Container is not Available or exists in system"},
                    status=200,
                )
            stock_obj = ContainerStock.objects.filter(
                container__container_no=container_no,
                container__location=location,
                container__site=site,
                container_status="IN",
                # container__is_available=True,
                container__status="IN",
            ).latest("pk")
            main_data = stock_obj.get_pregateout_prefill_data()
            return Response(main_data, status=200)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)


class PregateOutListAPIView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        try:
            pg_no = request.data["pg_no"]
            on_page_data = request.data["on_page_data"]
            client_name = request.data.get("client", None)
            shipping_line = request.data.get("shipping_line", None)
            location_name = request.data.get("location", None)
            site_name = request.data.get("site", None)
            container_no = request.data.get("container_no", None)
            do_validity = request.data.get("do_validity_date", None)
            do_validity_date_from_date = do_validity.get("from_date", None)
            do_validity_date_to_date = do_validity.get("to_date", None)
            # on_hold = request.data.get("on_hold", None)
            validity_expired = request.data.get("validity_expired", None)
            is_gateout_done = request.data.get("is_gateout_done", None)

            filters = Q()
            # if on_hold:
            #     filters &= Q(on_hold=on_hold)
            if validity_expired:
                filters &= Q(validity_expired=validity_expired)
            if is_gateout_done:
                filters &= Q(is_gateout_done=is_gateout_done)
            else:
                filters &= Q(is_gateout_done=False)
            if client_name:
                filters &= Q(stock__container__client__name__icontains=client_name)
            if shipping_line:
                filters &= Q(
                    stock__container__client__ref_code__icontains=shipping_line
                )
            if location_name:
                filters &= Q(location__name=location_name)
            if site_name:
                filters &= Q(site__name=site_name)
            if container_no:
                filters &= Q(stock__container__container_no__in=container_no)
            if do_validity_date_from_date and do_validity_date_to_date:
                filters &= Q(
                    do_validity_out_date__range=[
                        do_validity_date_from_date,
                        do_validity_date_to_date,
                    ]
                )

            pregateout_qs = PreGateOut.objects.filter(filters)

            # pagination
            paginator = Paginator(pregateout_qs, on_page_data)
            current_page = paginator.page(pg_no)
            no_of_data_count = paginator.count
            no_of_pages = paginator.num_pages
            on_page_data_count = current_page.object_list.count()
            prev_page = (
                current_page.previous_page_number()
                if current_page.has_previous()
                else None
            )
            next_page = (
                current_page.next_page_number() if current_page.has_next() else None
            )

            response_data = [
                {
                    **pregateout.get_pre_gateout_data(),
                    "sr_no": current_page.start_index() + index,
                }
                for index, pregateout in enumerate(current_page.object_list)
            ]

            return Response(
                {
                    "no_of_data": no_of_data_count,
                    "on_page_data": on_page_data_count,
                    "total_pages": no_of_pages,
                    "prev_page": prev_page,
                    "next_page": next_page,
                    "data": response_data,
                },
                status=200,
            )

        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)


class PreGateOutAvailableContainersAPIView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        try:
            data = request.data
            location = Location.objects.get(name=data.get("location"))
            site = Site.objects.get(name=data.get("site"), location=location)
            booking_no = request.data.get("booking_no", None)
            pregate_out_available_container_raw = ContainerStock.objects.filter(
                booking_no=booking_no,
                container__location=location,
                container__site=site,
                container_status="IN",
                # container__is_available=True,
                container__status="IN",
            )
            current_pregate_out_available_container = [
                each.container.container_no
                for each in pregate_out_available_container_raw
                if not PreGateOut.objects.filter(stock=each).exists()
            ]
            return Response(current_pregate_out_available_container, status=200)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": str(e)}, status=200)


class DownloadPreGateOutFormSix(APIView):
    """the get function will download pregate out form 6"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def get(self, request, pk, *args, **kwargs):
        try:
            context = {}
            pgo_obj = PreGateOut.objects.get(pk=pk)
            adv_payment = AdvancedHandlingPayment.objects.get(pk=pgo_obj.adv_payment_id)
            container = pgo_obj.stock.container
            site = container.site
            iso_code = ""
            try:
                iso_code_obj = TypeSizeCode.objects.get(
                    type=pgo_obj.stock.container.type, size=pgo_obj.stock.container.size
                )
                iso_code = iso_code_obj.code
            except:
                pass

            context["company_name"] = (
                "" if site.organization is None else site.organization
            )
            context["site_address"] = "" if site.address is None else site.address
            context["site_contact"] = "" if site.contact is None else site.contact
            context["export_container_booking_no"] = (
                "" if pgo_obj.stock.booking_no is None else pgo_obj.stock.booking_no
            )
            context["validity"] = pgo_obj.do_validity_out_date.strftime("%d/%m/%Y")
            context["line_name_code"] = container.client.ref_code
            context["container_no"] = container.container_no
            context["iso_code"] = iso_code
            context["cha_name"] = adv_payment.client.name
            context["vehicle"] = ""
            context["driver"] = ""
            context["seal_no"] = pgo_obj.stock.seal_no

            response = PDFTemplateResponse(
                request=request,
                template="depot/pre_gateout_form_6.html",
                filename="foo.pdf",
                context=context,
                show_content_in_browser=True,
                cmd_options={
                    "margin-top": 50,
                },
            )
            return response
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": f"Data Not Found [ {e} ]"}, status=200)


class AdvancedPaymentBulkUploadAPIView(APIView):
    """
    Get Function will provide the sample file
    Post Function will upload the file
    """

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def get(self, request, *args, **kwargs):
        try:
            file_path = os.path.join(
                BASE_DIR, "sample_stock/sample_advanced_payment_upload.xlsx"
            )
            with open(file_path, "rb") as file:
                response = HttpResponse(file.read(), content_type="application/xlsx")
                response["Content-Disposition"] = (
                    'attachment; filename="sample_advanced_payment_upload.xlsx"'
                )
            return response
        except:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": "Data Not Found"}, status=200)

    def post(self, request, *args, **kwargs):
        file = request.data.get("file")
        location = Location.objects.get(name=request.data.get("location"))
        site = Site.objects.get(name=request.data.get("site"), location=location)
        try:
            result = extract_advanced_Payment_excel_data(
                input_excel=file,
                location=location,
                site=site,
            )
            if result in ["Header Not Found", False]:
                return Response(
                    {"errorMsg": "Data file is corrupted, unable to import data"},
                    status=200,
                )
            return Response(result, status=200)
        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response({"errorMsg": f"Invalid Data Provided [{e}]"}, status=200)


class AdvancedPaymentBulkImportAPIView(APIView):
    """Post Function will Import the data"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def get_comman_object_data(self, data, location_name, site_name):
        location = Location.objects.get(name=location_name)
        site = Site.objects.get(name=site_name, location=location)
        client = Client.objects.get(
            name=data.get("client"), type="Party", location=location, site=site
        )
        return {
            "location": location,
            "site": site,
            "client": client,
        }

    def create_fin_account_debit_transaction(self, adv_payment_obj):
        fin_account_obj = CustomerFinAccount.objects.filter(
            customer=adv_payment_obj.client
        ).first()

        data = {
            "account": fin_account_obj,
            "amount": adv_payment_obj.original_amount,
            "linked_payment": adv_payment_obj.pk,
        }

        FinAccountTransaction.debit_amount(**data)

    def _validate_payment_data(self, client, data):
        """
        Helper function to validate payment data fields and return the necessary values.
        """
        payment_type = data.get("payment_type")
        cheque_no = data.get("cheque_no", None)
        utr_no = data.get("utr_no", None)
        original_amount = data.get("original_amount")
        entry_type = data.get("entry_type")
        bl_no = data.get("bl_no")
        bk_no = data.get("bk_no")
        with_gst = True

        tds = data.get("tds", 1)
        # 20
        size_20_quantity = data.get("size_20_quantity", 0)
        size_20_rate = 0
        size_20_tds_amount = 0
        # 40
        size_40_quantity = data.get("size_40_quantity", 0)
        size_40_rate = 0
        size_40_tds_amount = 0

        if entry_type == "IN":
            if AdvancedHandlingPayment.objects.filter(bl_no=bl_no).exists():
                temp_adv_payment = AdvancedHandlingPayment.objects.filter(bl_no=bl_no)
                if not temp_adv_payment[0].client == client:
                    raise ValueError(
                        f"Client should be [ {temp_adv_payment[0].client.name} ] with given bl_no [ {bl_no} ]"
                    )
        else:
            if AdvancedHandlingPayment.objects.filter(bk_no=bk_no).exists():
                temp_adv_payment = AdvancedHandlingPayment.objects.filter(bk_no=bk_no)
                if not temp_adv_payment[0].client == client:
                    raise ValueError(
                        f"Client should be [ {temp_adv_payment[0].client.name} ] with given bk_no [ {bk_no} ]"
                    )

        # Validate that quantity and original_amount are not zero or negative
        if int(size_20_quantity) == 0 and int(size_40_quantity) == 0:
            raise ValueError("Both Size Quantity cannot be zero")

        if int(size_20_quantity) < 0 or int(size_40_quantity) < 0:
            raise ValueError("Size Quantity cannot be negative")

        if original_amount <= 0:
            raise ValueError("Amount cannot be zero or negative")

        def validate_cheque_or_utr():
            nonlocal cheque_no, utr_no
            if not cheque_no and not utr_no:
                cheque_no = utr_no = None
            if cheque_no:
                utr_no = None
                if AdvancedHandlingPayment.objects.filter(cheque_no=cheque_no).exists():
                    raise ValueError("Payment already exists with the given cheque_no")
            if utr_no:
                cheque_no = None
                if AdvancedHandlingPayment.objects.filter(utr_no=utr_no).exists():
                    raise ValueError("Payment already exists with the given utr_no")

        if payment_type in ["Cheque", "NEFT", "RTGS", "Finance_Account"]:
            validate_cheque_or_utr()

        if payment_type == "Finance_Account":
            fin_account_obj = CustomerFinAccount.objects.filter(customer=client).first()
            with_gst = fin_account_obj.with_gst

        def validate_original_amt():
            nonlocal size_20_rate, size_20_tds_amount, size_40_rate, size_40_tds_amount
            size_20_original_amount = 0
            size_40_original_amount = 0

            if not int(size_20_quantity) == 0:
                size_20_rate = client.site.size_20_rate
                # tds
                tds_amount = float(size_20_rate) * (int(tds) / 100)
                size_20_tds_amount = round(float(tds_amount) * int(size_20_quantity))

                # tax
                tax = 0.18
                if not with_gst:
                    tax = 0
                tax_amount = float(size_20_rate) * tax
                size_20_total_amount_with_gst = round(
                    float(size_20_rate) + float(tax_amount)
                ) * int(size_20_quantity)
                # cal
                size_20_original_amount = float(size_20_total_amount_with_gst) - float(
                    size_20_tds_amount
                )

            if not int(size_40_quantity) == 0:
                size_40_rate = client.site.size_40_rate
                # tds
                tds_amount = float(size_40_rate) * (int(tds) / 100)
                size_40_tds_amount = round(float(tds_amount) * int(size_40_quantity))

                # tax
                tax = 0.18
                if not with_gst:
                    tax = 0
                tax_amount = float(size_40_rate) * tax
                size_40_total_amount_with_gst = round(
                    float(size_40_rate) + float(tax_amount)
                ) * int(size_40_quantity)
                # cal
                size_40_original_amount = float(size_40_total_amount_with_gst) - float(
                    size_40_tds_amount
                )

            my_original_amount = size_20_original_amount + size_40_original_amount
            if float(original_amount) < float(my_original_amount):
                raise ValueError(
                    "Sorry...Your original_amount is insufficient with respect to the quantity and current rate"
                )

        validate_original_amt()

        return (
            payment_type,
            cheque_no,
            utr_no,
            original_amount,
            entry_type,
            bl_no,
            bk_no,
            with_gst,
            size_20_quantity,
            size_20_rate,
            size_20_tds_amount,
            size_40_quantity,
            size_40_rate,
            size_40_tds_amount,
            tds,
        )

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data_list = request.data["importable_data"]
            for data in data_list:
                with transaction.atomic():
                    comman_objects = self.get_comman_object_data(
                        data=data,
                        location_name=request.data.get("location"),
                        site_name=request.data.get("site"),
                    )
                    (
                        payment_type,
                        cheque_no,
                        utr_no,
                        original_amount,
                        entry_type,
                        bl_no,
                        bk_no,
                        with_gst,
                        size_20_quantity,
                        size_20_rate,
                        size_20_tds_amount,
                        size_40_quantity,
                        size_40_rate,
                        size_40_tds_amount,
                        tds,
                    ) = self._validate_payment_data(
                        client=comman_objects["client"], data=data
                    )
                    main_data = {
                        "client": comman_objects["client"],
                        "location": comman_objects["location"],
                        "site": comman_objects["site"],
                        "bank_name": data.get("bank_name"),
                        "account_name": data.get("account_name"),
                        "account_no": data.get("account_no"),
                        "cheque_no": cheque_no,
                        "utr_no": utr_no,
                        "transaction_id": None,
                        "payment_type": payment_type,
                        "quantity": int(size_20_quantity) + int(size_40_quantity),
                        "original_amount": original_amount,
                        "entry_type": entry_type,
                        "tds": tds,
                        "with_gst": with_gst,
                        # 20
                        "size_20_rate": size_20_rate,
                        "size_20_quantity": size_20_quantity,
                        "size_20_tds_amount": size_20_tds_amount,
                        # 40
                        "size_40_rate": size_40_rate,
                        "size_40_quantity": size_40_quantity,
                        "size_40_tds_amount": size_40_tds_amount,
                    }
                    adv_payment_obj = None
                    if data.get("entry_type") == "IN":
                        main_data["bl_no"] = bl_no
                        adv_payment_obj = AdvancedHandlingPayment.create_in_adv_payment(
                            **main_data
                        )
                    else:
                        main_data["bk_no"] = bk_no
                        adv_payment_obj = (
                            AdvancedHandlingPayment.create_out_adv_payment(**main_data)
                        )

                    if payment_type == "Finance_Account":
                        self.create_fin_account_debit_transaction(adv_payment_obj)

            return Response({"successMsg": "Data Saved"}, status=200)
        except Exception as e:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return Response({"errorMsg": f"{e}"}, status=200)


class RejectedAdvancedPaymentDataApiView(APIView):
    """Post Function will get excel with rejected advanced payment data"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data = request.data["rejected_data"]
            faults = request.data["faults"]
            df_data = {
                "client": [each["client"] for each in data],
                "entry_type": [each["entry_type"] for each in data],
                "bl_no": [each["bl_no"] for each in data],
                "bk_no": [each["bk_no"] for each in data],
                "payment_type": [each["payment_type"] for each in data],
                "cheque_no": [each["cheque_no"] for each in data],
                "utr_no": [each["utr_no"] for each in data],
                "bank_name": [each["bank_name"] for each in data],
                "account_name": [each["account_name"] for each in data],
                "account_no": [each["account_no"] for each in data],
                "size_20_quantity": [each["size_20_quantity"] for each in data],
                "size_40_quantity": [each["size_40_quantity"] for each in data],
                "original_amount": [each["original_amount"] for each in data],
                "tds": [each["tds"] for each in data],
            }

            faults_list = []
            for row, messages in faults.items():
                for msg in messages:
                    faults_list.append({"row": row, "message": msg})

            temp_dir = os.path.join(BASE_DIR, "temp/sample_stock/")
            os.makedirs(temp_dir, exist_ok=True)

            new_temp_file_path = os.path.join(
                temp_dir, "sample_advanced_payment_upload.xlsx"
            )
            temp_file_path = os.path.join(
                BASE_DIR, "sample_stock/sample_advanced_payment_upload.xlsx"
            )

            with open(temp_file_path, "rb") as temp:
                old_temp_data = temp.read()

            with open(new_temp_file_path, "wb") as f:
                f.write(old_temp_data)

            book = load_workbook(new_temp_file_path)
            with pd.ExcelWriter(
                new_temp_file_path,
                engine="openpyxl",
                mode="a",
                if_sheet_exists="replace",
            ) as writer:

                stock_df = pd.DataFrame(df_data)
                stock_df.to_excel(writer, sheet_name="advanced_payment", index=False)

                faults_df = pd.DataFrame(faults_list)
                faults_df.to_excel(writer, sheet_name="faults", index=False)

            with open(new_temp_file_path, "rb") as temp:
                file_response = HttpResponse(
                    temp.read(), content_type=f"application/xlsx"
                )
                file_response["Content-Disposition"] = (
                    f'attachment; filename="sample_advanced_payment_upload.xlsx"'
                )
                os.remove(new_temp_file_path)
            return file_response

        except Exception as e:
            logging.getLogger("error_log").error(traceback.format_exc())
            return Response(
                {"errorMsg": f"Invalid Data Provided [ {str(e)} ]"}, status=200
            )


class AdvancedHandlingPaymentReportAPIView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User"]

    def get_invoice_no(self, payment_obj):
        try:
            lolo_pk_list = []
            invoice_no = ""
            if payment_obj.bl_no:
                lolo_pk_list = list(
                    GateInHistory.objects.select_related(
                        "container",
                        "container__location",
                        "container__site",
                        "lolo",
                        "gate_in",
                    )
                    .filter(
                        container__location=payment_obj.location,
                        container__site=payment_obj.site,
                        gate_in__bl_no=payment_obj.bl_no,
                    )
                    .values_list("lolo__pk", flat=True)
                )
            else:
                lolo_pk_list = (
                    GateOutHistory.objects.select_related(
                        "container",
                        "container__location",
                        "container__site",
                        "lolo",
                        "gate_out",
                    )
                    .filter(
                        container__location=payment_obj.location,
                        container__site=payment_obj.site,
                        gate_out__booking_no=payment_obj.bk_no,
                    )
                    .values_list("lolo__pk", flat=True)
                )

            bill_pk_list = list(
                CustomerBill.objects.filter(lolo_id__in=lolo_pk_list).values_list(
                    "pk", flat=True
                )
            )

            if not len(bill_pk_list) == 0:
                invoice_no_list = list(
                    CustomerBillInvoiceLine.objects.select_related("parent")
                    .filter(bill_id__in=bill_pk_list)
                    .values_list("parent__invoice_no", flat=True)
                    .distinct()
                )

                if invoice_no_list:
                    invoice_no = invoice_no_list[0]
            return invoice_no
        except:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return ""

    def get_payment_report_row_data(self, payment_obj):
        try:
            ref_no = ""
            if payment_obj.payment_type == "UPI":
                ref_no = payment_obj.transaction_id
            elif payment_obj.payment_type == "Cheque":
                ref_no = payment_obj.cheque_no
            elif payment_obj.payment_type in ["NEFT", "RTGS"]:
                ref_no = payment_obj.utr_no
            else:
                ref_no = payment_obj.remarks if payment_obj.remarks else ""

            size_20_data = {
                "bl_or_bk": (
                    payment_obj.bl_no
                    if payment_obj.entry_type == "IN"
                    else payment_obj.bk_no
                ),
                "party": payment_obj.client.name if payment_obj.client else None,
                # 20
                "quantity": f"{str(payment_obj.size_20_quantity)}X20",
                "full_amount": payment_obj.get_full_amount(size="20"),
                "recieved_amount": float(payment_obj.original_amount),
                "tds": payment_obj.tds,
                # "invoice_no": self.get_invoice_no(payment_obj),
                "invoice_no": (
                    payment_obj.advance_lolo_payment_rel.values_list(
                        "invoice_no", flat=True
                    ).first()
                    or ""
                ),
                "ref_no": ref_no,
                "date": payment_obj.date.astimezone(timezone.get_current_timezone())
                .date()
                .strftime("%d/%m/%Y"),
                "payment_type": payment_obj.payment_type,
            }

            size_40_data = {
                "bl_or_bk": (
                    payment_obj.bl_no
                    if payment_obj.entry_type == "IN"
                    else payment_obj.bk_no
                ),
                "party": payment_obj.client.name if payment_obj.client else None,
                # 40
                "quantity": f"{str(payment_obj.size_40_quantity)}X40",
                "full_amount": payment_obj.get_full_amount(size="40"),
                "recieved_amount": float(payment_obj.original_amount),
                "tds": payment_obj.tds,
                # "invoice_no": self.get_invoice_no(payment_obj),
                "invoice_no": (
                    payment_obj.advance_lolo_payment_rel.values_list(
                        "invoice_no", flat=True
                    ).first()
                    or ""
                ),
                "ref_no": ref_no,
                "date": payment_obj.date.astimezone(timezone.get_current_timezone())
                .date()
                .strftime("%d/%m/%Y"),
                "payment_type": payment_obj.payment_type,
            }
            return size_20_data, size_40_data
        except:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return None

    def get_df_data(self, request):
        try:
            location_name = request.data.get("location")
            site_name = request.data.get("site")

            from_date = datetime.datetime.strptime(
                request.data.get("from_date"), "%Y-%m-%d"
            ).date()
            from_time = datetime.datetime.strptime("00:00", "%H:%M").time()

            to_date = datetime.datetime.strptime(
                request.data.get("to_date"), "%Y-%m-%d"
            ).date()
            to_time = datetime.datetime.strptime("23:59", "%H:%M").time()

            from_date_time = datetime.datetime.combine(from_date, from_time)
            to_date_time = datetime.datetime.combine(to_date, to_time)

            payments_qs = AdvancedHandlingPayment.objects.filter(
                date__range=(from_date_time, to_date_time),
                location__name=location_name,
                site__name=site_name,
            ).order_by("date")

            data = []
            count = 0
            for each in payments_qs:
                size_20_data, size_40_data = self.get_payment_report_row_data(each)
                # 20
                count += 1
                size_20_data["sl_no"] = count
                # 40
                count += 1
                size_40_data["sl_no"] = count

                data.append(size_20_data)
                data.append(size_40_data)

            sl_no = [each.get("sl_no") for each in data]
            bl_or_bk = [each.get("bl_or_bk") for each in data]
            party = [each.get("party") for each in data]
            quantity = [each.get("quantity") for each in data]
            full_amount = [each.get("full_amount") for each in data]
            recieved_amount = [each.get("recieved_amount") for each in data]
            tds = [each.get("tds") for each in data]
            invoice_no = [each.get("invoice_no") for each in data]
            ref_no = [each.get("ref_no") for each in data]
            date = [each.get("date") for each in data]
            payment_type = [each.get("payment_type") for each in data]

            df_data = [
                [
                    sl_no[i],
                    bl_or_bk[i],
                    party[i],
                    quantity[i],
                    full_amount[i],
                    recieved_amount[i],
                    tds[i],
                    invoice_no[i],
                    ref_no[i],
                    date[i],
                    payment_type[i],
                ]
                for i in range(len(sl_no))
            ]
            if len(df_data) == 0:
                df_data.append(["" for i in range(0, 11)])
            return df_data
        except:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            df_data = [["" for i in range(0, 11)]]

    def create_report_wb(self, df_data):
        try:
            if not os.path.exists(os.path.join(BASE_DIR, "temp/")):
                os.makedirs(os.path.join(BASE_DIR, "temp/"))
            dt = datetime.datetime.now().astimezone(timezone.get_current_timezone())
            date = dt.date().strftime("%Y%m%d")
            time = dt.time().strftime("%H%M")
            temp_file_path = os.path.join(
                BASE_DIR, f"temp/advance_lolo_payment_report_{date}{time}.xlsx"
            )
            generic_workbook = xlsxwriter.Workbook(temp_file_path)
            generic_sheet = generic_workbook.add_worksheet("Advance_Lolo_Payment")

            generic_sheet.add_table(
                f"A1:K{1 + len(df_data)}",
                {
                    "data": df_data,
                    "columns": [
                        {"header": "Sr No"},
                        {"header": "BL /BOOKING NO"},
                        {"header": "PARTY NAME"},
                        {"header": "QUANTITY"},
                        {"header": "FULL AMOUNT"},
                        {"header": "PAYMENT RECEIVED AMT"},
                        {"header": "TDS"},
                        {"header": "INVOICE NO"},
                        {"header": "UTR /REFERENCE NO"},
                        {"header": "PAYMENT DATE"},
                        {"header": "PAYMENT TYPE"},
                    ],
                },
            )
            generic_workbook.close()
            return temp_file_path
        except Exception as e:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return None

    def post(self, request, *args, **kwargs):
        try:
            site_name = request.data.get("site")
            df_data = self.get_df_data(request=request)
            temp_file_path = self.create_report_wb(df_data=df_data)
            with open(temp_file_path, "rb") as temp:
                file_response = HttpResponse(
                    temp.read(), content_type=f"application/xlsx"
                )
                file_response["Content-Disposition"] = (
                    f'attachment; filename="advance_lolo_payment_report_{site_name}_{random.randint(100000, 999999)}.xlsx"'
                )
                os.remove(temp_file_path)
            return file_response
        except Exception as e:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return Response({"errorMsg": e}, status=200)
