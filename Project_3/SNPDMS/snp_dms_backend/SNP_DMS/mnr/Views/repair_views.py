from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from django.http import HttpResponse
import os
from SNP_DMS.settings.base import BASE_DIR
from mnr.functions import repair_extract_excel_data, mnr_img_validation
import pandas as pd
from openpyxl import load_workbook
from django.db import transaction

from account.models import AccountUser
from ..models import Estimate, Repair, MnrStaff, Survey
from master.models import Site, Location
from ..functions import (
    upload_mnr_file_to_s3,
    download_mnr_file_from_s3,
    upload_mnr_repair_img_to_s3,
    get_only_images,
    download_mnr_repair_img_to_s3,
    get_bill_type_for_mnr,
    delete_mnr_repair_img_from_s3,
)
import datetime
from django.utils import timezone
import logging, traceback
from billing_invoice.functions import create_billing_mnr

from account.permissions import HasAllowedRoles


class RepairImageDownloadView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = [
        "Admin",
        "Location Admin",
        "Site Admin",
        "Depot User",
        "MNR Team",
        "Repair",
    ]

    def post(self, request, pk, *args, **kwargs):
        try:
            repair_data = Repair.getById(pk)
            data = request.data
            image_id_list = data["image_id_list"]
            response = download_mnr_repair_img_to_s3(
                process="Repair",
                process_id=repair_data.parent.parent.pk,
                image_id_list=image_id_list,
            )
            return response
        except Exception as e:
            return Response({"errorMsg": f"Data Not Found [ {e} ]"}, status=200)


class RepairImageUploadView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = [
        "Admin",
        "Location Admin",
        "Site Admin",
        "Depot User",
        "MNR Team",
        "Repair",
    ]

    def post(self, request, pk, *args, **kwargs):
        try:
            data = request.data
            file_list = get_only_images(request.FILES.getlist("file_list"))
            if file_list is not None:
                # Condition
                repair_data = None
                try:
                    repair_data = Repair.getById(pk)
                    if not repair_data.is_proceed:
                        return Response(
                            {
                                "errorMsg": "Please Complete Repair stage before Image Upload"
                            },
                            status=200,
                        )
                except:
                    error_log = logging.getLogger("error_log")
                    error_log.error(traceback.format_exc())
                    return Response(
                        {"errorMsg": "Data Not Exist with the given Repair Id"},
                        status=200,
                    )

                # getting data
                container_no = None
                site = Site.objects.get(name=data["site"])
                if site.type == "DEPOT":
                    container_no = (
                        repair_data.parent.parent.depot.container.container_no
                    )
                else:
                    container_no = (
                        repair_data.parent.parent.non_depot.container.container_no
                    )

                response = upload_mnr_repair_img_to_s3(
                    location=data["location"],
                    site=data["site"],
                    site_type=site.type,
                    container_no=container_no,
                    process="Repair",
                    processId=repair_data.parent.parent.pk,
                    file_list=file_list,
                    add=data["add"],
                )

                if response is True:
                    return Response({"successMsg": "File Uploaded"}, status=200)
                else:
                    return Response({"errorMsg": "File Not Uploaded"}, status=200)
            else:
                return Response(
                    {"errorMsg": "Please Provide Images To Upload"}, status=200
                )
        except Exception as e:
            return Response({"errorMsg": f"Data Not Found [ {e} ]"}, status=200)


class RepairUploadView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = [
        "Admin",
        "Location Admin",
        "Site Admin",
        "Depot User",
        "MNR Team",
        "Repair",
    ]

    def post(self, request, pk, *args, **kwargs):
        try:
            data = request.data
            if data["file"] is not None:
                site = Site.objects.get(name=data["site"])
                # Getting Repair Data
                repair_data = None
                if Repair.checkExistById(pk):
                    repair_data = Repair.getById(pk)
                else:
                    return Response(
                        {"errorMsg": "Data Not Exist with the given Repair Id"},
                        status=200,
                    )
                container_no = None
                if site.type == "DEPOT":
                    container_no = (
                        repair_data.parent.parent.depot.container.container_no
                    )
                else:
                    container_no = (
                        repair_data.parent.parent.non_depot.container.container_no
                    )
                response = upload_mnr_file_to_s3(
                    location=data["location"],
                    site=data["site"],
                    site_type=site.type,
                    container_no=container_no,
                    process="Repair",
                    processId=repair_data.parent.parent.pk,
                    file=data["file"],
                )
                if response is True:
                    return Response({"successMsg": "File Uploaded"}, status=200)
                else:
                    return Response({"errorMsg": "File Not Uploaded"}, status=200)
            else:
                return Response(
                    {"errorMsg": "Please Provide File To Upload"}, status=200
                )
        except Exception as e:
            return Response({"errorMsg": f"Data Not Found [ {e} ]"}, status=200)


class RepairView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = [
        "Admin",
        "Location Admin",
        "Site Admin",
        "Depot User",
        "MNR Team",
        "Repair",
    ]

    def get(self, request, pk, *args, **kwargs):
        try:
            repair_data = Repair.getById(pk)
            response = download_mnr_file_from_s3(
                process="Repair", process_id=repair_data.parent.parent.pk
            )
            return response
        except Exception as e:
            return Response({"errorMsg": f"Data Not Found [ {e} ]"}, status=200)

    def post(self, request, *args, **kwargs):
        try:
            # General Data
            data = request.data
            app_user = AccountUser.objects.get(username=request.user.username)
            damage_category = None
            man_power = None
            grade = None
            remarks = None
            billing = None
            repair_billing = None
            wash_billing = None
            location = Location.objects.get(name=data["location"])
            site = Site.objects.get(name=data["site"])

            if not len(data["damage_category"]) == 0:
                damage_category = data["damage_category"]

            if not len(data["man_power"]) == 0:
                man_power = data["man_power"]

            if not len(data["grade"]) == 0:
                grade = data["grade"]

            if not len(data["remarks"]) == 0:
                remarks = data["remarks"]

            with transaction.atomic():
                # Getting Estimate Data
                estimate_data = None
                if Estimate.checkExistById(data["estimate_id"]):
                    estimate_data = Estimate.getById(data["estimate_id"])
                else:
                    return Response(
                        {"errorMsg": "Data Not Exist with the given Estimate Id"},
                        status=200,
                    )

                if data["placement"] == "True":
                    if Repair.checkExistByParentId(data["estimate_id"]):
                        repair_data = Repair.getByParentId(data["estimate_id"])
                        repair_data.damage_category = damage_category
                        repair_data.updated_by = app_user
                        repair_data.save()
                        if man_power is not None:
                            repair_data.man_power.clear()
                            repair_data.man_power.set(
                                MnrStaff.objects.filter(pk__in=man_power)
                            )
                        repair_data.save()
                    else:
                        repair_data = Repair.createPlacement(
                            parent=estimate_data,
                            damage_category=damage_category,
                            created_by=app_user,
                        )
                        if man_power is not None:
                            repair_data.man_power.set(
                                MnrStaff.objects.filter(pk__in=man_power)
                            )

                if data["complete"] == "True":
                    if Repair.checkExistByParentId(data["estimate_id"]):
                        repair_data = Repair.getByParentId(data["estimate_id"])
                        if repair_data.parent.is_approved is True:
                            repair_data.workComplete(
                                grade=grade, remarks=remarks, updated_by=app_user
                            )
                            repair_data.createNumber()

                            bill_type = get_bill_type_for_mnr(estimate_data.parent)
                            if bill_type["type"] == "repair and wash":
                                if site.new_billing_module:
                                    repair_bill = create_billing_mnr(
                                        bill_type=bill_type,
                                        obj=repair_data,
                                        repair_type="Repair",
                                        site=site,
                                    )
                                    wash_bill = create_billing_mnr(
                                        bill_type=bill_type,
                                        obj=repair_data,
                                        repair_type="Washing/Cleaning",
                                        site=site,
                                    )

                            else:
                                repair_type = None
                                if bill_type["type"] == "repair":
                                    repair_type = "Repair"
                                elif bill_type["type"] == "wash":
                                    repair_type = "Washing/Cleaning"

                                if site.new_billing_module:
                                    bill = create_billing_mnr(
                                        bill_type=bill_type,
                                        obj=repair_data,
                                        repair_type=repair_type,
                                        site=site,
                                    )

                        else:
                            return Response(
                                {
                                    "errorMsg": "Can't Proceed, Your Estimate is Not Approved Yet"
                                },
                                status=200,
                            )

            return Response({"successMsg": "Data Saved"}, status=200)
        except Exception as e:
            return Response({"errorMsg": str(e)}, status=200)

    def put(self, request, pk, *args, **kwargs):
        try:
            # General Data
            data = request.data
            app_user = AccountUser.objects.get(username=request.user.username)
            damage_category = None
            man_power = None
            grade = None
            remarks = None

            with transaction.atomic():
                if not len(data["damage_category"]) == 0:
                    damage_category = data["damage_category"]

                if not len(data["man_power"]) == 0:
                    man_power = data["man_power"]

                if not len(data["grade"]) == 0:
                    grade = data["grade"]

                if not len(data["remarks"]) == 0:
                    remarks = data["remarks"]

                current_repair_date = data.get("current_repair_date", None)
                current_repair_time = data.get("current_repair_time", None)

                if current_repair_date is not None and current_repair_time is not None:
                    # current_repair_date stuff
                    if len(current_repair_date) == 0:
                        current_repair_date = (
                            datetime.datetime.now()
                            .astimezone(timezone.get_current_timezone())
                            .date()
                        )
                    else:
                        try:
                            current_repair_date = datetime.datetime.strptime(
                                data["current_repair_date"], "%Y-%m-%d"
                            ).date()
                        except:
                            current_repair_date = (
                                datetime.datetime.now()
                                .astimezone(timezone.get_current_timezone())
                                .date()
                            )
                    # current_repair_time stuff
                    if len(current_repair_time) == 0:
                        current_repair_time = (
                            datetime.datetime.now()
                            .astimezone(timezone.get_current_timezone())
                            .time()
                        )
                    else:
                        try:
                            current_repair_time = datetime.datetime.strptime(
                                data["current_repair_time"], "%H:%M"
                            ).time()
                        except:
                            current_repair_time = (
                                datetime.datetime.now()
                                .astimezone(timezone.get_current_timezone())
                                .time()
                            )

                if Repair.checkExistById(pk):
                    repair_data = Repair.getById(pk)
                    if repair_data.parent.current_date > current_repair_date:
                        return Response(
                            {"errorMsg": "Repair date is behind Estimate date"}
                        )
                    repair_data.damage_category = damage_category
                    repair_data.grade = grade
                    repair_data.remarks = remarks
                    repair_data.updated_by = app_user
                    repair_data.save()
                    if man_power is not None:
                        repair_data.man_power.clear()
                        repair_data.man_power.set(
                            MnrStaff.objects.filter(pk__in=man_power)
                        )
                    if (
                        current_repair_date is not None
                        and current_repair_time is not None
                    ):
                        repair_data.updateRepairDateTime(
                            current_repair_date, current_repair_time
                        )
                    repair_data.save()

            return Response({"successMsg": "Data Updated"}, status=200)
        except Exception as e:
            return Response({"errorMsg": str(e)}, status=200)


class RepairUploadSampleFile(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = [
        "Admin",
        "Location Admin",
        "Site Admin",
        "Depot User",
        "MNR Team",
        "Repair",
    ]

    def get(self, request, *args, **kwargs):
        try:
            temp_file_path = os.path.join(
                BASE_DIR, "sample_stock/sample_repair_upload.xlsx"
            )
            with open(temp_file_path, "rb") as temp:
                file_response = HttpResponse(
                    temp.read(), content_type=f"application/xlsx"
                )
                file_response["Content-Disposition"] = (
                    f'attachment; filename="sample_repair_upload.xlsx"'
                )
            return file_response
        except:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return Response({"errorMsg": f"Data Not Found"}, status=200)

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data = request.data["file"]
            location_str = request.data["location"]
            site_str = request.data["site"]
            option = request.data["option"]
            location = Location.objects.get(name=location_str)
            site = Site.objects.get(name=site_str)
            result = repair_extract_excel_data(
                input_excel=data, option=option, location=location, site=site
            )
            if result == "Header Not Found" or result is False:
                return Response(
                    {"errorMsg": "Data file is corrupted, unable to import data"},
                    status=200,
                )
            else:
                return Response(result, status=200)
        except Exception as e:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return Response({"errorMsg": f"Invalid Data Provided [ {e} ]"}, status=200)


class RepairImport(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = [
        "Admin",
        "Location Admin",
        "Site Admin",
        "Depot User",
        "MNR Team",
        "Repair",
    ]

    def create_placement_data(self, estimate, damage_category, man_power, app_user):
        if Repair.checkExistByParentId(estimate.pk):
            repair_data = Repair.getByParentId(estimate.pk)
            repair_data.damage_category = damage_category
            repair_data.updated_by = app_user
            if man_power is not None:
                repair_data.man_power.clear()
                repair_data.man_power.set(MnrStaff.objects.filter(pk__in=man_power))
            repair_data.save()

        else:
            repair_data = Repair.createPlacement(
                parent=estimate,
                damage_category=damage_category,
                created_by=app_user,
            )
            if man_power is not None:
                repair_data.man_power.set(MnrStaff.objects.filter(pk__in=man_power))
            repair_data.save()
        return repair_data

    def create_wash_billing(self, estimate, bill_type, location, site, repair_data):
        if site.new_billing_module:
            bill = create_billing_mnr(
                bill_type=bill_type,
                obj=repair_data,
                repair_type="Washing/Cleaning",
                site=site,
            )
        return True

    def create_billing(self, estimate, bill_type, location, site, repair_data):
        repair_type = None
        if bill_type["type"] == "repair":
            repair_type = "Repair"
        elif bill_type["type"] == "wash":
            repair_type = "Washing/Cleaning"

        if site.new_billing_module:
            bill = create_billing_mnr(
                bill_type=bill_type,
                obj=repair_data,
                repair_type=repair_type,
                site=site,
            )
        return True

    def create_repair_billing(self, estimate, bill_type, location, site, repair_data):
        if site.new_billing_module:
            bill = create_billing_mnr(
                bill_type=bill_type,
                obj=repair_data,
                repair_type="Repair",
                site=site,
            )

        return True

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            app_user = AccountUser.objects.get(username=request.user.username)
            data = request.data["importable_data"]
            location = Location.objects.get(name=request.data["location"])
            site = Site.objects.get(name=request.data["site"])
            with transaction.atomic():
                rejected_data = []
                for each in data:
                    damage_category = each["damage"]
                    man_power = each["man_power"]
                    grade = each["grade"]
                    remarks = each["remarks"]
                    container_no = each["container_no"]
                    current_date = each["current_date"]
                    current_time = each["current_time"]
                    site_type = site.type

                    estimate = (
                        Estimate.objects.filter(
                            parent__depot__container__container_no=container_no
                        ).latest("pk")
                        if site_type == "DEPOT"
                        else Estimate.objects.filter(
                            parent__non_depot__container__container_no=container_no
                        ).latest("pk")
                    )

                    # if not Estimate.checkExistById(estimate.pk):
                    #     return Response(
                    #         {"errorMsg": "Data Not Exist with the given Estimate Id"},
                    #         status=200,
                    #     )
                    # estimate_data = Estimate.getById(estimate.pk)

                    if request.data["option"] == "Placement":
                        if Repair.objects.filter(parent=estimate).exists():
                            if Repair.objects.get(parent=estimate).is_draft:
                                rejected_data.append(container_no)
                                continue
                        self.create_placement_data(
                            estimate, damage_category, man_power, app_user
                        )

                    if request.data["option"] == "Complete and Placement":
                        if Repair.objects.filter(parent=estimate).exists():
                            if Repair.objects.get(parent=estimate).is_proceed:
                                rejected_data.append(container_no)
                                continue
                        self.create_placement_data(
                            estimate, damage_category, man_power, app_user
                        )
                        if Repair.checkExistByParentId(estimate.pk):
                            repair_data = Repair.getByParentId(estimate.pk)
                            # if not repair_data.parent.is_approved:
                            #     return Response(
                            #         {
                            #             "errorMsg": "Can't Proceed, Your Estimate is Not Approved Yet"
                            #         },
                            #         status=200,
                            #     )
                            repair_data.RepairWorkComplete(
                                current_date=current_date,
                                current_time=current_time,
                                grade=grade,
                                remarks=remarks,
                                updated_by=app_user,
                            )
                            repair_data.createNumber()

                            bill_type = get_bill_type_for_mnr(estimate.parent)
                            if bill_type["type"] == "repair and wash":
                                self.create_repair_billing(
                                    estimate, bill_type, location, site, repair_data
                                )

                                self.create_wash_billing(
                                    estimate, bill_type, location, site, repair_data
                                )

                            else:
                                self.create_billing(
                                    estimate, bill_type, location, site, repair_data
                                )

                if rejected_data:
                    return Response(
                        {
                            "errorMsg": f"All other data imported except this{rejected_data} containers"
                        },
                        status=200,
                    )

            return Response({"successMsg": "Repair Data Imported"}, status=200)
        except Exception as e:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return Response({"errorMsg": f"Invalid Data Provided [ {e} ]"}, status=200)


class RejectedRepairFile(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = [
        "Admin",
        "Location Admin",
        "Site Admin",
        "Depot User",
        "MNR Team",
        "Repair",
    ]

    def get_repair_df_data(self, data):
        container_no = []
        damage = []
        man_power = []
        current_date = []
        current_time = []
        grade = []
        remarks = []
        for each in data:
            container_no.append(each["container_no"])
            damage.append(each["damage"])
            man_power.append(each.get("man_power", "_"))
            current_date.append(each.get("current_date", "_"))
            current_time.append(each.get("current_time", "_"))
            grade.append(each.get("grade", "_"))
            remarks.append(each.get("remarks", "_"))

        return {
            "container_no": container_no,
            "damage": damage,
            "man_power": man_power,
            "current_date": current_date,
            "current_time": current_time,
            "grade": grade,
            "remarks": remarks,
        }

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data = request.data["rejected_data"]
            faults = request.data["faults"]

            reapir_df_data = self.get_repair_df_data(data)
            if not os.path.exists(os.path.join(BASE_DIR, "temp/sample_stock/")):
                os.makedirs(os.path.join(BASE_DIR, "temp/sample_stock/"))
            new_temp_file_path = os.path.join(
                BASE_DIR, f"temp/sample_stock/sample_repair_upload.xlsx"
            )
            temp_file_path = os.path.join(
                BASE_DIR, "sample_stock/sample_repair_upload.xlsx"
            )
            old_temp_data = None
            with open(temp_file_path, "rb") as temp:
                old_temp_data = temp.read()
            with open(new_temp_file_path, "wb") as f:
                f.write(old_temp_data)

            # New Version code
            book = load_workbook(new_temp_file_path)
            with pd.ExcelWriter(
                new_temp_file_path,
                engine="openpyxl",
                mode="a",
                if_sheet_exists="replace",
            ) as writer:
                repair_df = pd.DataFrame(reapir_df_data)
                repair_df.to_excel(
                    writer,
                    sheet_name="repair_upload",
                    startrow=0,
                    startcol=0,
                    index=False,
                )

                faults_df = pd.DataFrame(faults)
                faults_df.to_excel(
                    writer, sheet_name="faults", startrow=0, startcol=0, index=False
                )

            # Olde Version
            # book = load_workbook(new_temp_file_path)
            # writer = pd.ExcelWriter(new_temp_file_path, engine="openpyxl")
            # writer.book = book
            # writer.sheets = dict((ws.title, ws) for ws in book.worksheets)
            # repair_df = pd.DataFrame(reapir_df_data)
            # faults_df = pd.DataFrame(faults)
            # repair_df.to_excel(
            #     writer, sheet_name="repair_upload", startrow=0, startcol=0, index=False
            # )
            # faults_df.to_excel(
            #     writer, sheet_name="faults", startrow=0, startcol=0, index=False
            # )
            # fault_sheet = book.get_sheet_by_name("faults")
            # writer.save()

            with open(new_temp_file_path, "rb") as temp:
                file_response = HttpResponse(
                    temp.read(), content_type=f"application/xlsx"
                )
                file_response["Content-Disposition"] = (
                    f'attachment; filename="rejected_repair_file.xlsx"'
                )
                os.remove(new_temp_file_path)
            return file_response
        except Exception as e:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return Response({"errorMsg": f"Invalid Data Provided [ {e} ]"}, status=200)


class RepairImageDeleteView(APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = [
        "Admin",
        "Location Admin",
        "Site Admin",
        "Depot User",
        "MNR Team",
        "Repair",
    ]

    def post(self, request, pk, *args, **kwargs):
        try:
            repair_data = Repair.getById(pk)
            data = request.data
            image_id_list = data["image_id_list"]
            response = delete_mnr_repair_img_from_s3(
                process="Repair",
                process_id=repair_data.parent.parent.pk,
                image_id_list=image_id_list,
            )
            if response:
                return Response({"successMsg": "Successfully deleted images"})
            return Response({"errorMsg": "Images not found"})
        except Exception as e:
            return Response({"errorMsg": f"Data Not Found [ {e} ]"}, status=200)
