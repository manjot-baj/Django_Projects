from http import client
import os
from rest_framework import views
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from ..bulk_upload_functions import (
    getBulkTariffExcelData,
    make_bulk_tariff_excel,
    extract_survey_upload_excel_data,
    extract_survey_data_for_rejected_file,
    getSpecificLocationData,
    getTariffLineData,
    createSurveyLineObject,
    checkContainerExistence,
)
import random

from django.http import HttpResponse
import logging, traceback
from ..models import (
    TariffMaster,
    Survey,
    MnrStaff,
    Estimate,
)
from master.models import Site
from SNP_DMS.settings.base import BASE_DIR
from django.db import transaction
import shutil
import datetime
from depot.models import ContainerStock
from non_depot.models import NonDepotContainerStock
from account.models import AccountUser
from master.models import Location, Site
import pandas as pd
from openpyxl import load_workbook
from common.functions import get_location_site

from account.permissions import HasAllowedRoles


class DowloadTariffForBulkSurvey(views.APIView):
    """Post Function will download tariff data"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User", "MNR Team"]

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=200)
        try:
            data = request.data
            site_object = Site.objects.select_related("location").get(
                location__name=data["location"], name=data["site"]
            )
            tariff = TariffMaster.objects.get(site=site_object, client="MSC")
            main_data = getBulkTariffExcelData(parent=tariff)
            specific_location_data = getSpecificLocationData()
            temp_file_path = make_bulk_tariff_excel(main_data, specific_location_data)

            with open(temp_file_path, "rb") as temp:
                file_response = HttpResponse(
                    temp.read(), content_type=f"application/xlsx"
                )
                file_response["Content-Disposition"] = (
                    f'attachment; filename="msc_tariff.xlsx"'
                )
                os.remove(temp_file_path)
                return file_response
        except:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return Response({"errorMsg": "Data Not Found"}, status=200)


class SurveyBulkUploadSampleFile(views.APIView):
    """Get Function will give the sample file"""

    """ Post Function will give the extracted data """

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User", "MNR Team"]

    def get(self, request, *args, **kwargs):
        try:
            temp_file_path = os.path.join(
                BASE_DIR, "sample_stock/sample_washing_survey_bulk_upload.xlsx"
            )
            with open(temp_file_path, "rb") as temp:
                file_response = HttpResponse(
                    temp.read(), content_type=f"application/xlsx"
                )
                file_response["Content-Disposition"] = (
                    f'attachment; filename="sample_washing_survey_bulk_upload.xlsx"'
                )
            return file_response
        except:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return Response({"errorMsg": f"Data Not Found"}, status=200)

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=422)
        try:
            data = request.data["file"]
            # location = request.data["location"]
            # site = request.data["site"]
            location, site = get_location_site(
                request.data["location"], request.data["site"]
            )
            result = extract_survey_upload_excel_data(
                input_excel=data, location=location, site=site
            )
            if result == "Header Not Found" or result is False:
                return Response(
                    {"errorMsg": "Data file is corrupted, unable to import data"},
                    status=422,
                )
            else:
                return Response(result)
        except Exception as e:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())


class ImportWashingSurveyView(views.APIView):
    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User", "MNR Team"]

    def randomNumber(self):
        random_no = f"E{random.randint(1000000, 9999999)}"
        while (
            Estimate.objects.filter(number=random_no).exists()
            or Survey.objects.filter(estimate_number=random_no).exists()
        ):
            random_no = f"E{random.randint(1000000, 9999999)}"
            if (
                not Estimate.objects.filter(number=random_no).exists()
                and not Survey.objects.filter(estimate_number=random_no).exists()
            ):
                break
        return random_no

    def convertNone(self, dict_data):
        for key in dict_data:
            if len(str(dict_data[key])) == 0:
                dict_data[key] = None
        return dict_data

    def post(self, request, *args, **kwargs):
        try:
            data = request.data
            location = Location.objects.get(name=data["location"])
            site = Site.objects.get(name=data["site"])
            importable_data = data["importable_data"]
            accountUser = AccountUser.objects.get(username=request.user)
            labour_rate = TariffMaster.objects.get(
                client="MSC", location=location, site=site
            ).labour_rate

            if any(checkContainerExistence(site, importable_data)):
                return Response({"errorMsg": "Data already exists"})

            with transaction.atomic():
                for each in importable_data:
                    name = each["survey_by"].split()
                    if len(name) == 1:
                        param = {
                            "firstName": each["survey_by"],
                            "role": "Surveyor",
                            "location": location,
                            "site": site,
                        }
                    else:
                        param = {
                            "firstName": name[0],
                            "lastName": name[1],
                            "role": "Surveyor",
                            "location": location,
                            "site": site,
                        }
                    surveyBy = MnrStaff.objects.get(**param)
                    if site.type == "DEPOT":
                        depot = ContainerStock.objects.filter(
                            container__container_no=each["container_no"],
                            container__location=location,
                            container__site=site,
                            container__client__ref_code="MSC",
                            stage="Survey",
                        ).latest("pk")
                        parent = Survey.objects.create(
                            depot=depot,
                            non_depot=None,
                            survey_by=surveyBy,
                            created_by=accountUser,
                            updated_by=accountUser,
                            is_draft=True,
                            is_proceed=False,
                            original_date=datetime.datetime.now().date(),
                            original_time=datetime.datetime.now().time(),
                            current_date=datetime.datetime.now().date(),
                            current_time=datetime.datetime.now().time(),
                            labour_rate=labour_rate,
                            estimate_number=self.randomNumber(),
                        )
                    else:
                        non_depot = NonDepotContainerStock.objects.filter(
                            container__container_no=each["container_no"],
                            container__location=location,
                            container__site=site,
                            container__client__ref_code="MSC",
                            stage="Survey",
                        ).latest("pk")

                        parent = Survey.objects.create(
                            depot=None,
                            non_depot=non_depot,
                            survey_by=surveyBy,
                            created_by=accountUser,
                            updated_by=accountUser,
                            is_draft=True,
                            is_proceed=False,
                            original_date=datetime.datetime.now().date(),
                            original_time=datetime.datetime.now().time(),
                            current_date=datetime.datetime.now().date(),
                            current_time=datetime.datetime.now().time(),
                            labour_rate=labour_rate,
                            estimate_number=self.randomNumber(),
                        )

                    tarrif_line = getTariffLineData(location, site, each)
                    createSurveyLineObject(parent, tarrif_line, each)

            return Response({"successMsg": "Data Saved"}, status=200)
        except Exception as e:
            return Response({"errorMsg": str(e)}, status=200)


class RejectedSurveyDataFile(views.APIView):
    """Post Function will download and return the rejected data in a xls file"""

    permission_classes = (IsAuthenticated, HasAllowedRoles)
    allowed_roles = ["Admin", "Location Admin", "Site Admin", "Depot User", "MNR Team"]

    def post(self, request, *args, **kwargs):
        if not request.data:
            return Response({"errorMsg": "Please provide Data"}, status=422)
        try:
            data = request.data["rejected_data"]
            faults = request.data["faults"]
            tool_room_df_data = extract_survey_data_for_rejected_file(data=data)

            if not os.path.exists(os.path.join(BASE_DIR, "temp/sample_stock/")):
                os.makedirs(os.path.join(BASE_DIR, "temp/sample_stock/"))

            new_temp_file_path = os.path.join(
                BASE_DIR, "temp/sample_stock/sample_washing_survey_bulk_upload.xlsx"
            )
            temp_file_path = os.path.join(
                BASE_DIR, "sample_stock/sample_washing_survey_bulk_upload.xlsx"
            )

            shutil.copyfile(temp_file_path, new_temp_file_path)

            # New Version code
            book = load_workbook(new_temp_file_path)
            with pd.ExcelWriter(
                new_temp_file_path,
                engine="openpyxl",
                mode="a",
                if_sheet_exists="replace",
            ) as writer:
                tool_room_df_data.to_excel(
                    writer,
                    sheet_name="washing_survey",
                    startrow=0,
                    startcol=0,
                    index=False,
                )
                pd.DataFrame(faults).to_excel(
                    writer, sheet_name="faults", startrow=0, startcol=0, index=False
                )

            # Old Version
            # book = load_workbook(new_temp_file_path)
            # with pd.ExcelWriter(new_temp_file_path, engine="openpyxl") as writer:
            #     writer.book = book
            #     writer.sheets = dict((ws.title, ws) for ws in book.worksheets)
            #     tool_room_df_data.to_excel(
            #         writer,
            #         sheet_name="washing_survey",
            #         startrow=0,
            #         startcol=0,
            #         index=False,
            #     )
            #     pd.DataFrame(faults).to_excel(
            #         writer, sheet_name="faults", startrow=0, startcol=0, index=False
            #     )

            with open(new_temp_file_path, "rb") as temp:
                file_response = HttpResponse(
                    temp.read(), content_type="application/xlsx"
                )
                file_response["Content-Disposition"] = (
                    'attachment; filename="rejected_survey_file.xlsx"'
                )

            os.remove(new_temp_file_path)

            return file_response
        except Exception as e:
            error_log = logging.getLogger("error_log")
            error_log.error(traceback.format_exc())
            return Response({"errorMsg": f"Invalid Data Provided [ {e} ]"})
