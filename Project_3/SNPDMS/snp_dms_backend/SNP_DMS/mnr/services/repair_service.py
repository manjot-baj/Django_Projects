from mnr.models import (
    MnrStaff,
    SurveyLine,
    AfterRepairImage,
    Estimate,
    Survey,
    Approval,
)
from non_depot.models import NonDepotContainerStock
from decouple import config
from datetime import datetime
from mnr.models import Repair
from django.utils import timezone
from common.exceptions import ValidationError, ResourceNotFound
from ..functions import get_bill_type_for_mnr
from billing_invoice.functions import create_billing_mnr
from depot.functions_two import check_char_digit
import openpyxl
from depot.models import ContainerStock
import os
from SNP_DMS.settings.base import BASE_DIR
import xlsxwriter
from mnr.services.mnr_service import MNRService

AWS_REPAIR_IMAGE_BUCKET_NAME = config("AWS_REPAIR_IMAGE_BUCKET_NAME")


class RepairService:

    def staffData(self, location, site):
        return [
            {
                "name": f"{staff['firstName']} {staff['lastName'] or ''}",
                "pk": staff["pk"],
            }
            for staff in MnrStaff.objects.filter(
                location=location, site=site, role="Worker"
            ).values("firstName", "lastName", "pk")
        ]

    def getRepairData(self, repair, repair_container_no, stock, staff_data):
        man_hours = sum(
            list(
                SurveyLine.objects.filter(parent=repair.parent.parent).values_list(
                    "labour_hrs_tariff", flat=True
                )
            )
        )

        return {
            "pk": repair.pk,
            "container_no": repair_container_no or "",
            "repair_id": repair.pk,
            "is_draft": repair.is_draft,
            "is_proceed": repair.is_proceed,
            "estimate_id": repair.parent.pk,
            "number": repair.number or "",
            "created_by": repair.created_by.username or "",
            "updated_by": repair.updated_by.username or "",
            "line": stock.container.shipping_line.name or "",
            "size_type": f"{stock.container.size.name} / {stock.container.type.name}"
            or "",
            "in_date": stock.gate_in.in_date or "",
            "condition": stock.get_mnr_container_header_data()["condition"],
            "placement": repair.placement,
            "complete": repair.complete,
            "damage_category": repair.damage_category or "",
            "placement_date": repair.placement_date or "",
            "placement_time": (
                repair.placement_time.strftime("%H:%M") if repair.placement_time else ""
            ),
            "repair_date": repair.repair_date or "",
            "repair_time": (
                repair.repair_time.strftime("%H:%M") if repair.repair_time else ""
            ),
            "current_repair_date": repair.current_repair_date or "",
            "current_repair_time": (
                repair.current_repair_time.strftime("%H:%M")
                if repair.current_repair_time
                else ""
            ),
            "grade": repair.grade or "",
            "remarks": repair.remarks or "",
            "location": stock.container.location.name or "",
            "site": stock.container.site.name or "",
            "man_power": [each.pk for each in repair.man_power.all()],
            "is_img_uploaded": repair.is_img_uploaded,
            "man_hours": man_hours,
            "image_data": [
                {
                    "pk": image.pk,
                    "s3_object_name": image.s3_object_name,
                    "s3_file_name": image.s3_file_name,
                    "s3_image_link": f"https://{AWS_REPAIR_IMAGE_BUCKET_NAME}.s3.amazonaws.com/{image.s3_object_name}",
                }
                for image in AfterRepairImage.objects.filter(parent=repair)
            ],
            "staff_data": staff_data,
        }

    def getApprovalData(self, approval_container_no, approval, stock, staff_data):
        return {
            "container_no": approval_container_no or "",
            "estimate_id": approval.parent.pk,
            "line": stock.container.shipping_line.name or "",
            "size_type": f"{stock.container.size.name} / {stock.container.type.name}"
            or "",
            "in_date": stock.gate_in.in_date or "",
            "condition": stock.get_mnr_container_header_data()["condition"],
            "location": stock.container.location.name or "",
            "site": stock.container.site.name or "",
            "staff_data": staff_data,
            "man_hours": "",
        }

    def updateRepair(self, data, pk, app_user):
        damage_category = None
        man_power = None
        grade = None
        remarks = None
        if not len(data["damage_category"]) == 0:
            damage_category = data["damage_category"]

        if not len(data["man_power"]) == 0:
            man_power = data["man_power"]

        if not len(data["grade"]) == 0:
            grade = data["grade"]

        if not len(data["remarks"]) == 0:
            remarks = data["remarks"]

        current_repair_date = data.get("current_repair_date", None)
        current_repair_time = data.get("current_repair_time", None)

        if current_repair_date is not None and current_repair_time is not None:
            # current_repair_date stuff
            if len(current_repair_date) == 0:
                current_repair_date = (
                    datetime.now()
                    .astimezone(timezone.get_current_timezone())
                    .date()
                )
            else:
                try:
                    current_repair_date = datetime.strptime(
                        data["current_repair_date"], "%Y-%m-%d"
                    ).date()
                except:
                    current_repair_date = (
                        datetime.now()
                        .astimezone(timezone.get_current_timezone())
                        .date()
                    )
            # current_repair_time stuff
            if len(current_repair_time) == 0:
                current_repair_time = (
                    datetime.now()
                    .astimezone(timezone.get_current_timezone())
                    .time()
                )
            else:
                try:
                    current_repair_time = datetime.strptime(
                        data["current_repair_time"], "%H:%M"
                    ).time()
                except:
                    current_repair_time = (
                        datetime.now()
                        .astimezone(timezone.get_current_timezone())
                        .time()
                    )

        if Repair.checkExistById(pk):
            repair_data = Repair.getById(pk)
            if repair_data.parent.current_date > current_repair_date:
                raise ValidationError(
                    "Current repair date cannot be before the estimate date."
                )
            repair_data.damage_category = damage_category
            repair_data.grade = grade
            repair_data.remarks = remarks
            repair_data.updated_by = app_user
            repair_data.save()
            if man_power is not None:
                repair_data.man_power.clear()
                repair_data.man_power.set(MnrStaff.objects.filter(pk__in=man_power))
            if current_repair_date is not None and current_repair_time is not None:
                repair_data.updateRepairDateTime(
                    current_repair_date, current_repair_time
                )
            repair_data.save()

    def createRepair(
        self, data, app_user, site, damage_category, man_power, grade, remarks
    ):
        estimate_data = None
        if Estimate.checkExistById(data["estimate_id"]):
            estimate_data = Estimate.getById(data["estimate_id"])
        else:
            raise ResourceNotFound(
                f"Estimate with ID {data['estimate_id']} does not exist."
            )

        if data["placement"] == "True":
            if Repair.checkExistByParentId(data["estimate_id"]):
                repair_data = Repair.getByParentId(data["estimate_id"])
                repair_data.damage_category = damage_category
                repair_data.updated_by = app_user
                repair_data.save()
                if man_power is not None:
                    repair_data.man_power.clear()
                    repair_data.man_power.set(MnrStaff.objects.filter(pk__in=man_power))
                repair_data.save()
            else:
                repair_data = Repair.createPlacement(
                    parent=estimate_data,
                    damage_category=damage_category,
                    created_by=app_user,
                )
                if man_power is not None:
                    repair_data.man_power.set(MnrStaff.objects.filter(pk__in=man_power))

        if data["complete"] == "True":
            if Repair.checkExistByParentId(data["estimate_id"]):
                repair_data = Repair.getByParentId(data["estimate_id"])
                if repair_data.parent.is_approved is True:
                    repair_data.workComplete(
                        grade=grade, remarks=remarks, updated_by=app_user
                    )
                    repair_data.createNumber()

                    bill_type = get_bill_type_for_mnr(estimate_data.parent)
                    if bill_type["type"] == "repair and wash":
                        if site.new_billing_module:
                            repair_bill = create_billing_mnr(
                                bill_type=bill_type,
                                obj=repair_data,
                                repair_type="Repair",
                                site=site,
                            )
                            wash_bill = create_billing_mnr(
                                bill_type=bill_type,
                                obj=repair_data,
                                repair_type="Washing/Cleaning",
                                site=site,
                            )

                    else:
                        repair_type = None
                        if bill_type["type"] == "repair":
                            repair_type = "Repair"
                        elif bill_type["type"] == "wash":
                            repair_type = "Washing/Cleaning"

                        if site.new_billing_module:
                            bill = create_billing_mnr(
                                bill_type=bill_type,
                                obj=repair_data,
                                repair_type=repair_type,
                                site=site,
                            )

                else:
                    raise ValidationError(
                        "Cannot complete the repair as the estimate is not approved."
                    )

    def extractDataList(self, input_excel):
        ps = openpyxl.load_workbook(input_excel)
        sheet = ps["repair_upload"]
        container_no_raw = [
            sheet["A" + str(row)].value for row in range(1, sheet.max_row + 1)
        ]
        damage_raw = [
            sheet["B" + str(row)].value for row in range(1, sheet.max_row + 1)
        ]
        first_name_raw = [
            sheet["C" + str(row)].value for row in range(1, sheet.max_row + 1)
        ]
        last_name_raw = [
            sheet["D" + str(row)].value for row in range(1, sheet.max_row + 1)
        ]
        current_date_raw = [
            sheet["E" + str(row)].value for row in range(1, sheet.max_row + 1)
        ]
        current_time_raw = [
            sheet["F" + str(row)].value for row in range(1, sheet.max_row + 1)
        ]
        grade_raw = [sheet["G" + str(row)].value for row in range(1, sheet.max_row + 1)]
        remarks_raw = [
            sheet["H" + str(row)].value for row in range(1, sheet.max_row + 1)
        ]

        container_no = [each for each in container_no_raw if not each is None]
        damage = [each for each in damage_raw if not each is None]
        first_name = [each for each in first_name_raw if not each is None]
        last_name = [each for each in last_name_raw if not each is None]
        current_date = [each for each in current_date_raw if not each is None]
        current_time = [each for each in current_time_raw if not each is None]
        grade = [each for each in grade_raw if not each is None]
        remarks = [each for each in remarks_raw if not each is None]

        popped = [
            container_no.pop(0),
            damage.pop(0),
            first_name.pop(0),
            last_name.pop(0),
            current_date.pop(0),
            current_time.pop(0),
            grade.pop(0),
            remarks.pop(0),
        ]

        headers = [
            "container_no",
            "damage",
            "first_name",
            "last_name",
            "current_date",
            "current_time",
            "grade",
            "remarks",
        ]

        if not popped == headers:
            return "Header Not Found"

        extracted_data_list = []

        for i in range(len(container_no)):
            extracted_data_list.append(
                {
                    "sr_no": str(i),
                    "container_no": str(container_no[i]),
                    "damage": str(damage[i]),
                    "first_name": str(first_name[i]),
                    "last_name": str(last_name[i]),
                    "current_date": str(current_date[i]),
                    "current_time": str(current_time[i]),
                    "grade": str(grade[i]),
                    "remarks": str(remarks[i]),
                }
            )
        ps.close()
        return extracted_data_list

    def validateDepotContainer(self, container_no_str, location, site, error_msg, each):
        if not ContainerStock.objects.filter(
            container__container_no=container_no_str,
            container__location=location,
            container__site=site,
        ).exists():
            error_msg.append(
                f"In row {str(int(each['sr_no']) + 2)} there is problem in container_no column, "
                f"this({container_no_str}) Container does not exists"
            )
        else:
            container_stock = ContainerStock.objects.get(
                container__container_no=container_no_str,
                container__location=location,
                container__site=site,
                container_status="IN",
            )
            if not Survey.objects.filter(depot=container_stock).exists():
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 2)} there is problem in container_no column, "
                    f"this({container_no_str}) in Survey Pending stage"
                )
            elif not Estimate.objects.filter(parent__depot=container_stock).exists():
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 2)} there is problem in container_no column, "
                    f"this({container_no_str}) Estimate data does not exists"
                )
            elif not Approval.objects.filter(
                parent__parent__depot=container_stock
            ).exists():
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 2)} there is problem in container_no column, "
                    f"this({container_no_str}) Approval data does not exists"
                )
            elif not Approval.objects.get(
                parent__parent__depot=container_stock
            ).is_approved:
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 2)} there is problem in container_no column, "
                    f"this({container_no_str}) Estimate is Not Approved Yet"
                )
            elif Repair.objects.filter(parent__parent__depot=container_stock).exists():
                if Repair.objects.get(parent__parent__depot=container_stock).is_proceed:
                    error_msg.append(
                        f"In row {str(int(each['sr_no']) + 2)} there is problem in container_no column, "
                        f"this({container_no_str}) Repair Already exists"
                    )
                elif Repair.objects.get(parent__parent__depot=container_stock).is_draft:
                    error_msg.append("")
            else:
                error_msg.append("")

        return error_msg

    def validateNonDepotContainer(
        self, container_no_str, location, site, error_msg, each
    ):
        if not NonDepotContainerStock.objects.filter(
            container__container_no=container_no_str,
            container__location=location,
            container__site=site,
        ).exists():
            error_msg.append(
                f"In row {str(int(each['sr_no']) + 2)} there is problem in container_no column, "
                f"this({container_no_str}) Container does not exists"
            )
        else:
            container_stock = NonDepotContainerStock.objects.get(
                container__container_no=container_no_str,
                container__location=location,
                container__site=site,
                container_status="IN",
            )
            if not Survey.objects.filter(non_depot=container_stock).exists():
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 2)} there is problem in container_no column, "
                    f"this({container_no_str}) in Survey Pending stage"
                )
            elif not Estimate.objects.filter(
                parent__non_depot=container_stock
            ).exists():
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 2)} there is problem in container_no column, "
                    f"this({container_no_str}) Estimate data does not exists"
                )
            elif not Approval.objects.filter(
                parent__parent__non_depot=container_stock
            ).exists():
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 2)} there is problem in container_no column, "
                    f"this({container_no_str}) Approval data does not exists"
                )
            elif not Approval.objects.get(
                parent__parent__non_depot=container_stock
            ).is_approved:
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 2)} there is problem in container_no column, "
                    f"this({container_no_str}) Estimate is Not Approved Yet"
                )
            elif Repair.objects.filter(
                parent__parent__non_depot=container_stock
            ).exists():
                if Repair.objects.get(
                    parent__parent__non_depot=container_stock
                ).is_proceed:
                    error_msg.append(
                        f"In row {str(int(each['sr_no']) + 2)} there is problem in container_no column, "
                        f"this({container_no_str}) Repair Already exists"
                    )
                elif Repair.objects.get(
                    parent__parent__non_depot=container_stock
                ).is_draft:
                    error_msg.append("")

            else:
                error_msg.append("")

        return error_msg

    def extractRepairData(self, input_excel, option, location, site):
        extracted_data_list = self.extractDataList(input_excel)
        error_data = []
        error_data_msg = {}
        correct_data = []
        site_type = site.type

        for each in extracted_data_list:
            error_msg = []

            container_no_str = each["container_no"]
            if len(container_no_str) == 11 and check_char_digit(container_no_str):
                is_container_repeated = any(
                    obj_data["container_no"] == container_no_str
                    for obj_data in correct_data
                )
                if is_container_repeated:
                    error_msg.append(
                        f"In row {str(int(each['sr_no']) + 2)} there is problem in container_no column, "
                        f"this({container_no_str}) container is repeated in your file"
                    )
                else:
                    if site_type == "DEPOT":
                        self.validateDepotContainer(
                            container_no_str, location, site, error_msg, each
                        )
                    else:
                        self.validateNonDepotContainer(
                            container_no_str, location, site, error_msg, each
                        )

            else:
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 2)} there is problem in container_no column,"
                    f"container_no ({container_no_str}) is invalid"
                )

            damage_str = each["damage"]
            damage_list = ["OK", "CLEANING", "LD", "MD", "HD", "AV", "AR", "DM"]
            if not damage_str in damage_list:
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 2)} there is problem in damage column,"
                    f"damage provided is not correct, should be among these {', '.join(damage_list)}"
                )
            else:
                error_msg.append("")

            first_name_str = each.pop("first_name")
            last_name_str = each.pop("last_name")
            if last_name_str == "_":
                staff_filter = {
                    "firstName": first_name_str,
                    "role": "Worker",
                    "location": location,
                    "site": site,
                }
            else:
                staff_filter = {
                    "firstName": first_name_str,
                    "lastName": last_name_str,
                    "role": "Worker",
                    "location": location,
                    "site": site,
                }

            if MnrStaff.objects.filter(**staff_filter).exists():
                error_msg.append("")
                staff = MnrStaff.objects.get(**staff_filter)
                each["man_power"] = [staff.pk]
            else:
                error_msg.append(
                    f"In row {int(each['sr_no']) + 2} there is a problem in the first_name and last_name column, "
                    f"man_power not found in the system database"
                )

            if option == "Complete and Placement":
                current_date_str = each["current_date"]
                if current_date_str == "_" or current_date_str == "":
                    error_msg.append(
                        f"In row {str(int(each['sr_no']) + 2)} there is problem in current_date column,"
                        f"current date column cannot be empty"
                    )
                else:
                    try:
                        datetime.strptime(current_date_str, "%d_%m_%Y").date()
                        error_msg.append("")
                    except:
                        error_msg.append(
                            f"In row {str(int(each['sr_no']) + 2)} there is problem in current_date column,"
                            f"date is not in dd_mm_yyyy format"
                        )

                current_time_str = each["current_time"]
                if current_time_str == "_" or current_time_str == "":
                    error_msg.append(
                        f"In row {str(int(each['sr_no']) + 2)} there is problem in current time column,"
                        f"current time column cannot be empty"
                    )
                else:
                    try:
                        datetime.strptime(current_time_str, "%H_%M").time()
                        error_msg.append("")
                    except:
                        error_msg.append(
                            f"In row {str(int(each['sr_no']) + 2)} there is problem in current_time column,"
                            f"date is not in HH_MM format"
                        )

                grade_str = each["grade"]
                if grade_str == "_" or grade_str == "":
                    error_msg.append(
                        f"In row {str(int(each['sr_no']) + 2)} there is problem in grade column,"
                        f"grade column cannot be empty"
                    )
                else:
                    grade_list = ["A", "B", "C", "D", "E", "F"]
                    if not grade_str in grade_list:
                        error_msg.append(
                            f"In row {str(int(each['sr_no']) + 2)} there is problem in grade column,"
                            f"grade provided is not correct, should be among these {', '.join(grade_list)}"
                        )
                    else:
                        error_msg.append("")

                remarks_str = each["remarks"]
                try:
                    if remarks_str == "_":
                        each["remarks"] = ""

                    error_msg.append("")
                except:
                    error_msg.append(
                        f"In row {str(int(each['sr_no']) + 2)} there is problem in remarks column"
                    )
            else:
                each["current_date"] = ""
                each["current_time"] = ""
                each["grade"] = ""
                each["remarks"] = ""
                error_msg.extend([""] * 4)

            error_data_msg[f"row {str(int(each['sr_no']) + 2)}"] = error_msg
            if any(error_data_msg[f"row {str(int(each['sr_no']) + 2)}"]) is True:
                error_data.append(each)
            if not each in error_data:
                correct_data.append(each)
        correct_data_count = str(len(correct_data))
        error_data_count = str(len(error_data))
        main_data = {
            "importable_data": correct_data,
            "importable_data_count": correct_data_count,
            "rejected_data": error_data,
            "rejected_data_count": error_data_count,
            "faults": error_data_msg,
        }
        return main_data

    def createPlacementData(self, estimate, damage_category, man_power, app_user):
        if Repair.checkExistByParentId(estimate.pk):
            repair_data = Repair.getByParentId(estimate.pk)
            repair_data.damage_category = damage_category
            repair_data.updated_by = app_user
            if man_power is not None:
                repair_data.man_power.clear()
                repair_data.man_power.set(MnrStaff.objects.filter(pk__in=man_power))
            repair_data.save()

        else:
            repair_data = Repair.createPlacement(
                parent=estimate,
                damage_category=damage_category,
                created_by=app_user,
            )
            if man_power is not None:
                repair_data.man_power.set(MnrStaff.objects.filter(pk__in=man_power))
            repair_data.save()
        return repair_data

    def createWashBilling(self, estimate, bill_type, location, site, repair_data):
        if site.new_billing_module:
            bill = create_billing_mnr(
                bill_type=bill_type,
                obj=repair_data,
                repair_type="Washing/Cleaning",
                site=site,
            )
        return True

    def createBilling(self, estimate, bill_type, location, site, repair_data):
        repair_type = None
        if bill_type["type"] == "repair":
            repair_type = "Repair"
        elif bill_type["type"] == "wash":
            repair_type = "Washing/Cleaning"

        if site.new_billing_module:
            bill = create_billing_mnr(
                bill_type=bill_type,
                obj=repair_data,
                repair_type=repair_type,
                site=site,
            )
        return True

    def createRepairBilling(self, estimate, bill_type, location, site, repair_data):
        if site.new_billing_module:
            bill = create_billing_mnr(
                bill_type=bill_type,
                obj=repair_data,
                repair_type="Repair",
                site=site,
            )

        return True

    def getRepairDfData(self, data):
        container_no = []
        damage = []
        man_power = []
        current_date = []
        current_time = []
        grade = []
        remarks = []
        for each in data:
            container_no.append(each["container_no"])
            damage.append(each["damage"])
            man_power.append(each.get("man_power", "_"))
            current_date.append(each.get("current_date", "_"))
            current_time.append(each.get("current_time", "_"))
            grade.append(each.get("grade", "_"))
            remarks.append(each.get("remarks", "_"))

        return {
            "container_no": container_no,
            "damage": damage,
            "man_power": man_power,
            "current_date": current_date,
            "current_time": current_time,
            "grade": grade,
            "remarks": remarks,
        }

    def getGenericJobSheetRowData(self, pk):
        try:
            repair_object = Repair.objects.get(pk=pk)
            if repair_object.parent.parent.depot is not None:
                container_header_data = (
                    repair_object.parent.parent.depot.get_mnr_container_header_data()
                )
            else:
                container_header_data = (
                    repair_object.parent.parent.non_depot.get_mnr_container_header_data()
                )
            main_data = {
                "container_no": container_header_data["container_no"],
                "size_type": container_header_data["size_type"],
                "line": container_header_data["line"],
                "aging": container_header_data["aging"],
            }
            survey_object = repair_object.parent.parent
            survey_line_objects = SurveyLine.objects.filter(
                parent=survey_object, is_rejected=False
            )
            survey_line_data = [
                {
                    "main_component": (
                        "" if line.main_component is None else line.main_component
                    ),
                    "component_description": (
                        ""
                        if line.component_description is None
                        else line.component_description
                    ),
                    "location_description": (
                        ""
                        if line.location_description is None
                        else line.location_description
                    ),
                    "specific_location_description": (
                        ""
                        if line.specific_location_description is None
                        else line.specific_location_description
                    ),
                    "material_description": (
                        ""
                        if line.material_description is None
                        else line.material_description
                    ),
                    "damage_description": (
                        ""
                        if line.damage_description is None
                        else line.damage_description
                    ),
                    "repair_description": (
                        ""
                        if line.repair_description is None
                        else line.repair_description
                    ),
                    "unit": "" if line.unit is None else line.unit,
                    "measurement": "" if line.measurement is None else line.measurement,
                    "length_and_width": (
                        ""
                        if line.length_and_width is None
                        else MNRService().getLengthWidthValueString(
                            line.length_and_width
                        )
                    ),
                    "quantity": "" if line.quantity is None else str(line.quantity),
                }
                for line in survey_line_objects
            ]
            main_data["survey_lines"] = survey_line_data
            return main_data
        except Exception as e:
            return None

    def getGenericJobSheetMainData(self, pk_list):
        try:
            main_data = []
            sr_no_count = 0
            for pk in pk_list:
                sr_no_count += 1
                row_data = self.getGenericJobSheetRowData(pk)
                row_data["sr_no"] = str(sr_no_count)
                main_data.append(row_data)

            sr_no = []
            container_no = []
            size_type = []
            line = []
            aging = []
            seq_no = []
            main_component = []
            component_description = []
            location_description = []
            specific_location_description = []
            damage_description = []
            material_description = []
            repair_description = []
            measurement = []
            unit = []
            length_and_width = []
            quantity = []

            for each in main_data:
                sr_no.append(each["sr_no"])
                container_no.append(each["container_no"])
                size_type.append(each["size_type"])
                line.append(each["line"])
                aging.append(each["aging"])
                seq_no_count = 0
                for line_data in each["survey_lines"]:
                    seq_no_count += 1
                    if seq_no_count > 1:
                        sr_no.append("")
                        container_no.append("")
                        size_type.append("")
                        line.append("")
                        aging.append("")
                    seq_no.append(seq_no_count)
                    main_component.append(line_data["main_component"])
                    component_description.append(line_data["component_description"])
                    location_description.append(line_data["location_description"])
                    specific_location_description.append(
                        line_data["specific_location_description"]
                    )
                    damage_description.append(line_data["damage_description"])
                    material_description.append(line_data["material_description"])
                    repair_description.append(line_data["repair_description"])
                    measurement.append(line_data["measurement"])
                    unit.append(line_data["unit"])
                    length_and_width.append(line_data["length_and_width"])
                    quantity.append(line_data["quantity"])
                sr_no.append("")
                container_no.append("")
                size_type.append("")
                line.append("")
                aging.append("")
                seq_no.append("")
                main_component.append("")
                component_description.append("")
                location_description.append("")
                specific_location_description.append("")
                damage_description.append("")
                material_description.append("")
                repair_description.append("")
                measurement.append("")
                unit.append("")
                length_and_width.append("")
                quantity.append("")

            excel_data = [
                [
                    sr_no[i],
                    container_no[i],
                    size_type[i],
                    line[i],
                    aging[i],
                    seq_no[i],
                    main_component[i],
                    component_description[i],
                    location_description[i],
                    specific_location_description[i],
                    damage_description[i],
                    material_description[i],
                    repair_description[i],
                    measurement[i],
                    unit[i],
                    length_and_width[i],
                    quantity[i],
                ]
                for i in range(len(sr_no))
            ]
            return excel_data
        except Exception as e:
            return None

    def createGenericJobSheetWb(self, df_data):
        try:
            line = df_data[0][3]
            if not os.path.exists(os.path.join(BASE_DIR, "temp/")):
                os.makedirs(os.path.join(BASE_DIR, "temp/"))
            dt = datetime.now().astimezone(timezone.get_current_timezone())
            date = dt.date().strftime("%Y%m%d")
            time = dt.time().strftime("%H%M")
            temp_file_path = os.path.join(
                BASE_DIR, f"temp/{line.lower()}_job_sheet_{date}{time}.xlsx"
            )
            generic_workbook = xlsxwriter.Workbook(temp_file_path)
            job_sheet = generic_workbook.add_worksheet("JOB")
            job_sheet.add_table(
                f"A1:Q{1 + len(df_data)}",
                {
                    "data": df_data,
                    "columns": [
                        {"header": "Sr.No"},
                        {"header": "Container No"},
                        {"header": "Size Type"},
                        {"header": "Line"},
                        {"header": "Aging"},
                        {"header": "Seq No"},
                        {"header": "Main Component"},
                        {"header": "Component"},
                        {"header": "Location"},
                        {"header": "Specific Location"},
                        {"header": "Damage"},
                        {"header": "Material"},
                        {"header": "Repair"},
                        {"header": "Measurement"},
                        {"header": "Unit"},
                        {"header": "L x W"},
                        {"header": "QTY"},
                    ],
                },
            )
            generic_workbook.close()
            return temp_file_path
        except Exception as e:
            return None
