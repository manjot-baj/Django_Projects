from mnr.models import TariffMaster, TariffMasterLine, TariffSpecificLocation
import os
from SNP_DMS.settings.base import BASE_DIR
from datetime import datetime
from django.utils import timezone
import xlsxwriter
from common.exceptions import AlreadyExists, ValidationError
import openpyxl
import random


class TariffService:

    def getTariffData(self, payload):
        parent = TariffMaster.objects.get(pk=payload["parent_id"])
        main_component = payload.get("main_component", None)
        component_code = payload.get("component_code", None)
        component_description = payload.get("component_description", None)
        location_code = payload.get("location_code", None)
        location_description = payload.get("location_description", None)
        specific_location_code = payload.get("specific_location_code", None)
        specific_location_description = payload.get(
            "specific_location_description", None
        )
        damage_code = payload.get("damage_code", None)
        damage_description = payload.get("damage_description", None)
        material_code = payload.get("material_code", None)
        material_description = payload.get("material_description", None)
        repair_code = payload.get("repair_code", None)
        repair_description = payload.get("repair_description", None)
        measurement = payload.get("measurement", None)
        unit = payload.get("unit", None)
        repair_type = payload.get("repair_type", None)

        if (
            main_component is not None
            and component_code is not None
            and component_description is not None
            and location_code is not None
            and location_description is not None
            and specific_location_code is not None
            and specific_location_description is not None
            and damage_code is not None
            and damage_description is not None
            and material_code is not None
            and material_description is not None
            and repair_code is not None
            and repair_description is not None
            and unit is not None
            and measurement is not None
        ):
            queryset = TariffMasterLine.objects.filter(
                parent=parent,
                main_component=main_component,
                component_code=component_code,
                component_description=component_description,
                location_code=location_code,
                location_description=location_description,
                damage_code=damage_code,
                damage_description=damage_description,
                material_code=material_code,
                material_description=material_description,
                repair_code=repair_code,
                repair_description=repair_description,
                unit=unit,
                measurement=measurement,
            ).exclude(length_and_width=None)

            if repair_type == "Washing":
                queryset = queryset.exclude(wash_clean_tariff=float(0))
            else:
                queryset.filter(wash_clean_tariff=float(0))

            queryset = queryset.values("length_and_width").distinct()

            main_list = list(queryset)

        elif (
            main_component is not None
            and component_code is not None
            and component_description is not None
            and location_code is not None
            and location_description is not None
            and specific_location_code is not None
            and specific_location_description is not None
            and damage_code is not None
            and damage_description is not None
            and material_code is not None
            and material_description is not None
            and repair_code is not None
            and repair_description is not None
        ):
            queryset = TariffMasterLine.objects.filter(
                parent=parent,
                main_component=main_component,
                component_code=component_code,
                component_description=component_description,
                location_code=location_code,
                location_description=location_description,
                damage_code=damage_code,
                damage_description=damage_description,
                material_code=material_code,
                material_description=material_description,
                repair_code=repair_code,
                repair_description=repair_description,
            ).exclude(unit=None, measurement=None)

            if repair_type == "Washing":
                queryset = queryset.exclude(wash_clean_tariff=float(0))
            else:
                queryset.filter(wash_clean_tariff=float(0))

            queryset = queryset.values("unit", "measurement").distinct()

            main_list = list(queryset)

        elif (
            main_component is not None
            and component_code is not None
            and component_description is not None
            and location_code is not None
            and location_description is not None
            and specific_location_code is not None
            and specific_location_description is not None
            and damage_code is not None
            and damage_description is not None
            and material_code is not None
            and material_description is not None
        ):
            queryset = TariffMasterLine.objects.filter(
                parent=parent,
                main_component=main_component,
                component_code=component_code,
                component_description=component_description,
                location_code=location_code,
                location_description=location_description,
                damage_code=damage_code,
                damage_description=damage_description,
                material_code=material_code,
                material_description=material_description,
            ).exclude(repair_code=None, repair_description=None)

            if repair_type == "Washing":
                queryset = queryset.exclude(wash_clean_tariff=float(0))
            else:
                queryset.filter(wash_clean_tariff=float(0))

            queryset = queryset.values("repair_code", "repair_description").distinct()

            main_list = list(queryset)

        elif (
            main_component is not None
            and component_code is not None
            and component_description is not None
            and location_code is not None
            and location_description is not None
            and specific_location_code is not None
            and specific_location_description is not None
            and damage_code is not None
            and damage_description is not None
        ):
            queryset = TariffMasterLine.objects.filter(
                parent=parent,
                main_component=main_component,
                component_code=component_code,
                component_description=component_description,
                location_code=location_code,
                location_description=location_description,
                damage_code=damage_code,
                damage_description=damage_description,
            ).exclude(material_code=None, material_description=None)

            if repair_type == "Washing":
                queryset = queryset.exclude(wash_clean_tariff=float(0))
            else:
                queryset.filter(wash_clean_tariff=float(0))

            queryset = queryset.values(
                "material_code", "material_description"
            ).distinct()

            main_list = list(queryset)

        elif (
            main_component is not None
            and component_code is not None
            and component_description is not None
            and location_code is not None
            and location_description is not None
            and specific_location_code is not None
            and specific_location_description is not None
        ):
            queryset = TariffMasterLine.objects.filter(
                parent=parent,
                main_component=main_component,
                component_code=component_code,
                component_description=component_description,
                location_code=location_code,
                location_description=location_description,
            ).exclude(damage_code=None, damage_description=None)

            if repair_type == "Washing":
                queryset = queryset.exclude(wash_clean_tariff=float(0))
            else:
                queryset.filter(wash_clean_tariff=float(0))

            queryset = queryset.values("damage_code", "damage_description").distinct()

            main_list = list(queryset)

        elif (
            main_component is not None
            and component_code is not None
            and component_description is not None
            and location_code is not None
            and location_description is not None
        ):
            queryset = (
                TariffSpecificLocation.objects.filter(
                    location_code=location_code.split("_")[0],
                )
                .values("specific_location_code", "specific_location_description")
                .distinct()
            )

            main_list = list(queryset)
            main_list.append(
                {
                    "specific_location_code": "NA",
                    "specific_location_description": "NA",
                }
            )

        elif (
            main_component is not None
            and component_code is not None
            and component_description is not None
        ):
            queryset = TariffMasterLine.objects.filter(
                parent=parent,
                main_component=main_component,
                component_code=component_code,
                component_description=component_description,
            ).exclude(location_code=None, location_description=None)

            if repair_type == "Washing":
                queryset = queryset.exclude(wash_clean_tariff=float(0))
            else:
                queryset.filter(wash_clean_tariff=float(0))

            queryset = queryset.values(
                "location_code", "location_description"
            ).distinct()

            main_list = list(queryset)

        elif main_component is not None:

            queryset = TariffMasterLine.objects.filter(
                parent=parent,
                main_component=main_component,
            ).exclude(component_code=None, component_description=None)

            if repair_type == "Washing":
                queryset = queryset.exclude(wash_clean_tariff=float(0))
            else:
                queryset.filter(wash_clean_tariff=float(0))

            queryset = queryset.values(
                "component_code", "component_description"
            ).distinct()

            main_list = list(queryset)

        return main_list

    def getListOfTariff(self, params):
        queryset = TariffMaster.objects.select_related("location", "site").filter(
            **params
        )

        return [
            {
                "pk": each.pk,
                "client": each.client,
                "location": each.location.name,
                "site": each.site.name,
                "labour_rate": str(each.labour_rate),
            }
            for each in queryset
        ]

    def getTariffExcelRowData(self, data_object):
        data = {}
        tariff_code = "_"
        main_component = "_"
        component_code = "_"
        component_description = "_"
        location_code = "_"
        location_description = "_"
        damage_code = "_"
        damage_description = "_"
        material_code = "_"
        material_description = "_"
        repair_code = "_"
        repair_description = "_"
        unit = "_"
        measurement = "_"
        length_and_width = "_"
        quantity = "0"
        labour_hrs_tariff = "0"
        wash_clean_tariff = "0"
        material_tariff = "0"

        if data_object.tariff_code is not None:
            tariff_code = data_object.tariff_code
        data["tariff_code"] = tariff_code

        if data_object.main_component is not None:
            main_component = data_object.main_component
        data["main_component"] = main_component

        if data_object.component_code is not None:
            component_code = str(data_object.component_code).split("_")[0]
        data["component_code"] = component_code

        if data_object.component_description is not None:
            component_description = data_object.component_description
        data["component_description"] = component_description

        if data_object.location_code is not None:
            location_code = str(data_object.location_code).split("_")[0]
        data["location_code"] = location_code

        if data_object.location_description is not None:
            location_description = data_object.location_description
        data["location_description"] = location_description

        if data_object.damage_code is not None:
            damage_code = str(data_object.damage_code).split("_")[0]
        data["damage_code"] = damage_code

        if data_object.damage_description is not None:
            damage_description = data_object.damage_description
        data["damage_description"] = damage_description

        if data_object.material_code is not None:
            material_code = str(data_object.material_code).split("_")[0]
        data["material_code"] = material_code

        if data_object.material_description is not None:
            material_description = data_object.material_description
        data["material_description"] = material_description

        if data_object.repair_code is not None:
            repair_code = str(data_object.repair_code).split("_")[0]
        data["repair_code"] = repair_code

        if data_object.repair_description is not None:
            repair_description = data_object.repair_description
        data["repair_description"] = repair_description

        if data_object.unit is not None:
            unit = data_object.unit
        data["unit"] = unit

        if data_object.measurement is not None:
            measurement = data_object.measurement
        data["measurement"] = measurement

        if data_object.length_and_width is not None:
            length_and_width = data_object.length_and_width
        data["length_and_width"] = length_and_width

        if data_object.quantity is not None:
            quantity = str(data_object.quantity)
        data["quantity"] = quantity

        if data_object.labour_hrs_tariff is not None:
            labour_hrs_tariff = str(data_object.labour_hrs_tariff)
        data["labour_hrs_tariff"] = labour_hrs_tariff

        if data_object.wash_clean_tariff is not None:
            wash_clean_tariff = str(data_object.wash_clean_tariff)
        data["wash_clean_tariff"] = wash_clean_tariff

        if data_object.material_tariff is not None:
            material_tariff = str(data_object.material_tariff)
        data["material_tariff"] = material_tariff

        return data

    def getTariffExcelData(self, pk):

        parent = TariffMaster.objects.get(pk=pk)
        queryset = TariffMasterLine.objects.filter(parent=parent)
        count = 0
        data = []
        for each in queryset:
            count += 1
            each_data = self.getTariffExcelRowData(each)
            each_data["sr_no"] = count
            data.append(each_data)
        sr_no = [each.get("sr_no") for each in data]
        tariff_code = [each.get("tariff_code") for each in data]
        main_component = [each.get("main_component") for each in data]
        component_code = [each.get("component_code") for each in data]
        component_description = [each.get("component_description") for each in data]
        location_code = [each.get("location_code") for each in data]
        location_description = [each.get("location_description") for each in data]
        damage_code = [each.get("damage_code") for each in data]
        damage_description = [each.get("damage_description") for each in data]
        material_code = [each.get("material_code") for each in data]
        material_description = [each.get("material_description") for each in data]
        repair_code = [each.get("repair_code") for each in data]
        repair_description = [each.get("repair_description") for each in data]
        unit = [each.get("unit") for each in data]
        measurement = [each.get("measurement") for each in data]
        length_and_width = [each.get("length_and_width") for each in data]
        quantity = [each.get("quantity") for each in data]
        labour_hrs_tariff = [each.get("labour_hrs_tariff") for each in data]
        wash_clean_tariff = [each.get("wash_clean_tariff") for each in data]
        material_tariff = [each.get("material_tariff") for each in data]
        df_data = [
            [
                tariff_code[i],
                main_component[i],
                component_code[i],
                component_description[i],
                location_code[i],
                location_description[i],
                damage_code[i],
                damage_description[i],
                material_code[i],
                material_description[i],
                repair_code[i],
                repair_description[i],
                unit[i],
                measurement[i],
                length_and_width[i],
                quantity[i],
                labour_hrs_tariff[i],
                wash_clean_tariff[i],
                material_tariff[i],
            ]
            for i in range(len(sr_no))
        ]

        if len(df_data) == 0:
            df_data.append(["" for i in range(0, 18)])
        main_response = {
            "client": parent.client,
            "labour_rate": parent.labour_rate,
            "tariff_lines": df_data,
        }
        return main_response

    def makeTariffExcel(self, data):

        if not os.path.exists(os.path.join(BASE_DIR, "temp/")):
            os.makedirs(os.path.join(BASE_DIR, "temp/"))
        dt = datetime.now().astimezone(timezone.get_current_timezone())
        date = dt.date().strftime("%Y%m%d")
        time = dt.time().strftime("%H%M")
        temp_file_path = os.path.join(
            BASE_DIR, f"temp/tariff_{data['client']}_{date}{time}.xlsx"
        )
        workbook = xlsxwriter.Workbook(temp_file_path)
        sheet = workbook.add_worksheet("tariff")
        sheet.write("D1", "Labour Rate")
        sheet.write("E1", f"{data['labour_rate']}")
        sheet.add_table(
            f"A2:S{2 + len(data['tariff_lines'])}",
            {
                "data": data["tariff_lines"],
                "columns": [
                    {"header": "TARIFF CODE"},
                    {"header": "MAIN COMPONENT"},
                    {"header": "COMPONENT CODE"},
                    {"header": "COMPONENT DESCRIPTION"},
                    {"header": "LOCATION CODE"},
                    {"header": "LOCATION DESCRIPTION"},
                    {"header": "DAMAGE CODE"},
                    {"header": "DAMAGE DESCRIPTION"},
                    {"header": "MATERIAL CODE"},
                    {"header": "MATERIAL DESCRIPTION"},
                    {"header": "REPAIR CODE"},
                    {"header": "REPAIR DESCRIPTION"},
                    {"header": "UNIT CODE"},
                    {"header": "UNIT DESCRIPTION"},
                    {"header": "L/W"},
                    {"header": "QTY"},
                    {"header": "Labour Hours Tarrif"},
                    {"header": "Wash/Clean Tarrif"},
                    {"header": "Material Rate Tarrif"},
                ],
            },
        )
        workbook.close()
        return temp_file_path

    def getCleanLW(self, string):
        if "*" in str(string):
            string_list = string.split("*")
            l = str(int(float(string_list[0])))
            w = str(int(float(string_list[1])))
            return f"{l}*{w}"
        else:
            return str(int(string))

    def extractTariffExcelData(self, input_excel):
        ps = openpyxl.load_workbook(input_excel, data_only=True)
        sheet = ps["tariff"]
        # Data Extraction
        # tariff data
        labour_rate = sheet["E1"].value

        tariff_code_raw1 = [
            sheet["A" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        tariff_code_raw2 = [each for each in tariff_code_raw1 if not each is None]
        tariff_code = [None if each == "_" else each for each in tariff_code_raw2]

        main_component_raw1 = [
            sheet["B" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        main_component_raw2 = [each for each in main_component_raw1 if not each is None]
        main_component = [None if each == "_" else each for each in main_component_raw2]

        component_code_raw1 = [
            sheet["C" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        component_code_raw2 = [each for each in component_code_raw1 if not each is None]
        component_code = [None if each == "_" else each for each in component_code_raw2]

        component_description_raw1 = [
            sheet["D" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        component_description_raw2 = [
            each for each in component_description_raw1 if not each is None
        ]
        component_description = [
            None if each == "_" else each for each in component_description_raw2
        ]

        location_code_raw1 = [
            sheet["E" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        location_code_raw2 = [each for each in location_code_raw1 if not each is None]
        location_code = [None if each == "_" else each for each in location_code_raw2]

        location_description_raw1 = [
            sheet["F" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        location_description_raw2 = [
            each for each in location_description_raw1 if not each is None
        ]
        location_description = [
            None if each == "_" else each for each in location_description_raw2
        ]

        damage_code_raw1 = [
            sheet["G" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        damage_code_raw2 = [each for each in damage_code_raw1 if not each is None]
        damage_code = [None if each == "_" else each for each in damage_code_raw2]

        damage_description_raw1 = [
            sheet["H" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        damage_description_raw2 = [
            each for each in damage_description_raw1 if not each is None
        ]
        damage_description = [
            None if each == "_" else each for each in damage_description_raw2
        ]

        material_code_raw1 = [
            sheet["I" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        material_code_raw2 = [each for each in material_code_raw1 if not each is None]
        material_code = [None if each == "_" else each for each in material_code_raw2]

        material_description_raw1 = [
            sheet["J" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        material_description_raw2 = [
            each for each in material_description_raw1 if not each is None
        ]
        material_description = [
            None if each == "_" else each for each in material_description_raw2
        ]

        repair_code_raw1 = [
            sheet["K" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        repair_code_raw2 = [each for each in repair_code_raw1 if not each is None]
        repair_code = [None if each == "_" else each for each in repair_code_raw2]

        repair_description_raw1 = [
            sheet["L" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        repair_description_raw2 = [
            each for each in repair_description_raw1 if not each is None
        ]
        repair_description = [
            None if each == "_" else each for each in repair_description_raw2
        ]

        unit_code_raw1 = [
            sheet["M" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        unit_code_raw2 = [each for each in unit_code_raw1 if not each is None]
        unit_code = [None if each == "_" else each for each in unit_code_raw2]

        unit_description_raw1 = [
            sheet["N" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        unit_description_raw2 = [
            each for each in unit_description_raw1 if not each is None
        ]
        unit_description = [
            None if each == "_" else each for each in unit_description_raw2
        ]

        length_width_raw1 = [
            sheet["O" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        length_width_raw2 = [each for each in length_width_raw1 if not each is None]
        length_width_1 = [None if each == "_" else each for each in length_width_raw2]

        quantity_raw1 = [
            sheet["P" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        quantity_raw2 = [each for each in quantity_raw1 if not each is None]
        quantity = [0 if each == "_" else each for each in quantity_raw2]

        labour_hrs_tariff_raw1 = [
            sheet["Q" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        labour_hrs_tariff_raw2 = [
            each for each in labour_hrs_tariff_raw1 if not each is None
        ]
        labour_hrs_tariff = [
            0 if each == "_" else each for each in labour_hrs_tariff_raw2
        ]

        wash_clean_tariff_raw1 = [
            sheet["R" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        wash_clean_tariff_raw2 = [
            each for each in wash_clean_tariff_raw1 if not each is None
        ]
        wash_clean_tariff = [
            0 if each == "_" else each for each in wash_clean_tariff_raw2
        ]

        material_rate_tariff_raw1 = [
            sheet["S" + str(row)].value for row in range(2, sheet.max_row + 1)
        ]
        material_rate_tariff_raw2 = [
            each for each in material_rate_tariff_raw1 if not each is None
        ]
        material_rate_tariff = [
            0 if each == "_" else each for each in material_rate_tariff_raw2
        ]

        ps.close()

        # Checking excel data if valid
        popped = [
            tariff_code.pop(0),
            main_component.pop(0),
            component_code.pop(0),
            component_description.pop(0),
            location_code.pop(0),
            location_description.pop(0),
            damage_code.pop(0),
            damage_description.pop(0),
            material_code.pop(0),
            material_description.pop(0),
            repair_code.pop(0),
            repair_description.pop(0),
            unit_code.pop(0),
            unit_description.pop(0),
            length_width_1.pop(0),
            quantity.pop(0),
            labour_hrs_tariff.pop(0),
            wash_clean_tariff.pop(0),
            material_rate_tariff.pop(0),
        ]

        headers = [
            "TARIFF CODE",
            "MAIN COMPONENT",
            "COMPONENT CODE",
            "COMPONENT DESCRIPTION",
            "LOCATION CODE",
            "LOCATION DESCRIPTION",
            "DAMAGE CODE",
            "DAMAGE DESCRIPTION",
            "MATERIAL CODE",
            "MATERIAL DESCRIPTION",
            "REPAIR CODE",
            "REPAIR DESCRIPTION",
            "UNIT CODE",
            "UNIT DESCRIPTION",
            "L/W",
            "QTY",
            "Labour Hours Tarrif",
            "Wash/Clean Tarrif",
            "Material Rate Tarrif",
        ]

        if not popped == headers:
            return "Header Not Found"
        # Data Gatheration

        length_width = [
            None if each == "_" else self.getCleanLW(each) for each in length_width_1
        ]

        helper_tariff_code = []
        helper_main_component = []
        helper_component_code = []
        helper_component_description = []
        helper_location_code = []
        helper_location_description = []
        helper_damage_code = []
        helper_damage_description = []
        helper_material_code = []
        helper_material_description = []
        helper_repair_code = []
        helper_repair_description = []
        helper_unit_code = []
        helper_unit_description = []
        helper_length_width = []
        helper_quantity = []
        helper_labour_hrs_tariff = []
        helper_wash_clean_tariff = []
        helper_material_rate_tariff = []

        for i in range(len(main_component)):
            if len(damage_code[i].split("/")) > 1:
                damage_code_list = damage_code[i].split("/")
                damage_description_list = damage_description[i].split("/")
                for j in range(len(damage_code_list)):
                    helper_tariff_code.append(tariff_code[i])
                    helper_main_component.append(main_component[i])
                    helper_component_code.append(component_code[i])
                    helper_component_description.append(component_description[i])
                    helper_location_code.append(location_code[i])
                    helper_location_description.append(location_description[i])
                    helper_damage_code.append(damage_code_list[j])
                    helper_damage_description.append(damage_description_list[j])
                    helper_material_code.append(material_code[i])
                    helper_material_description.append(material_description[i])
                    helper_repair_code.append(repair_code[i])
                    helper_repair_description.append(repair_description[i])
                    helper_unit_code.append(unit_code[i])
                    helper_unit_description.append(unit_description[i])
                    helper_length_width.append(length_width[i])
                    helper_quantity.append(quantity[i])
                    helper_labour_hrs_tariff.append(labour_hrs_tariff[i])
                    helper_wash_clean_tariff.append(wash_clean_tariff[i])
                    helper_material_rate_tariff.append(material_rate_tariff[i])
            else:
                helper_tariff_code.append(tariff_code[i])
                helper_main_component.append(main_component[i])
                helper_component_code.append(component_code[i])
                helper_component_description.append(component_description[i])
                helper_location_code.append(location_code[i])
                helper_location_description.append(location_description[i])
                helper_damage_code.append(damage_code[i])
                helper_damage_description.append(damage_description[i])
                helper_material_code.append(material_code[i])
                helper_material_description.append(material_description[i])
                helper_repair_code.append(repair_code[i])
                helper_repair_description.append(repair_description[i])
                helper_unit_code.append(unit_code[i])
                helper_unit_description.append(unit_description[i])
                helper_length_width.append(length_width[i])
                helper_quantity.append(quantity[i])
                helper_labour_hrs_tariff.append(labour_hrs_tariff[i])
                helper_wash_clean_tariff.append(wash_clean_tariff[i])
                helper_material_rate_tariff.append(material_rate_tariff[i])

        tariff_code = helper_tariff_code
        main_component = helper_main_component
        component_code = helper_component_code
        component_description = helper_component_description
        location_code = helper_location_code
        location_description = helper_location_description
        damage_code = helper_damage_code
        damage_description = helper_damage_description
        material_code = helper_material_code
        material_description = helper_material_description
        repair_code = helper_repair_code
        repair_description = helper_repair_description
        unit_code = helper_unit_code
        unit_description = helper_unit_description
        length_width = helper_length_width
        quantity = helper_quantity
        labour_hrs_tariff = helper_labour_hrs_tariff
        wash_clean_tariff = helper_wash_clean_tariff
        material_rate_tariff = helper_material_rate_tariff

        component = []
        component_dict = {}
        component_new_element = None
        for i in range(len(main_component)):
            if component_description[i] in component:
                component_code[i] = component_dict[component_description[i]]
            else:
                component.append(component_description[i])
                component_new_element = (
                    f"{component_code[i]}_{random.randint(1000, 9999)}"
                )
                component_dict[component_description[i]] = component_new_element
                component_code[i] = component_new_element

        location = []
        location_dict = {}
        location_new_element = None
        for i in range(len(main_component)):
            if location_description[i] in location:
                location_code[i] = location_dict[location_description[i]]
            else:
                location.append(location_description[i])
                location_new_element = (
                    f"{location_code[i]}_{random.randint(1000, 9999)}"
                )
                location_dict[location_description[i]] = location_new_element
                location_code[i] = location_new_element

        damage = []
        damage_dict = {}
        damage_new_element = None
        for i in range(len(main_component)):
            if damage_description[i] in damage:
                damage_code[i] = damage_dict[damage_description[i]]
            else:
                damage.append(damage_description[i])
                damage_new_element = f"{damage_code[i]}_{random.randint(1000, 9999)}"
                damage_dict[damage_description[i]] = damage_new_element
                damage_code[i] = damage_new_element

        material = []
        material_dict = {}
        material_new_element = None
        for i in range(len(main_component)):
            if material_description[i] in material:
                material_code[i] = material_dict[material_description[i]]
            else:
                material.append(material_description[i])
                material_new_element = (
                    f"{material_code[i]}_{random.randint(1000, 9999)}"
                )
                material_dict[material_description[i]] = material_new_element
                material_code[i] = material_new_element

        repair = []
        repair_dict = {}
        repair_new_element = None
        for i in range(len(main_component)):
            if repair_description[i] in repair:
                repair_code[i] = repair_dict[repair_description[i]]
            else:
                repair.append(repair_description[i])
                repair_new_element = f"{repair_code[i]}_{random.randint(1000, 9999)}"
                repair_dict[repair_description[i]] = repair_new_element
                repair_code[i] = repair_new_element

        main_data = [
            {
                "labour_rate": labour_rate,
                "tariff_code": tariff_code[i],
                "main_component": main_component[i],
                "component_code": component_code[i],
                "component_description": component_description[i],
                "location_code": location_code[i],
                "location_description": location_description[i],
                "damage_code": damage_code[i],
                "damage_description": damage_description[i],
                "material_code": material_code[i],
                "material_description": material_description[i],
                "repair_code": repair_code[i],
                "repair_description": repair_description[i],
                "unit_code": unit_code[i],
                "unit_description": unit_description[i],
                "length_width": length_width[i],
                "quantity": quantity[i],
                "labour_hrs_tariff": labour_hrs_tariff[i],
                "wash_clean_tariff": wash_clean_tariff[i],
                "material_rate_tariff": material_rate_tariff[i],
            }
            for i in range(len(main_component))
        ]
        return main_data

    def uploadTariff(self, file, client, location, site):
        if TariffMaster.objects.filter(
            client=client, location=location, site=site
        ).exists():
            raise AlreadyExists("Tariff Already Uploaded")
        extracted_data = self.extractTariffExcelData(input_excel=file)
        if extracted_data == "Header Not Found":
            raise ValidationError("Data file is corrupted, unable to extract data")
        else:
            try:
                labour_rate = float(extracted_data[0]["labour_rate"])
            except Exception as e:
                raise ValidationError(
                    "Please Provide Valid Data for Tariff Calculation"
                )

            parent = TariffMaster(
                client=client, location=location, site=site, labour_rate=labour_rate
            )
            parent.save()

            objList = []
            for each in extracted_data:
                try:
                    labour_hrs_tariff = float(each["labour_hrs_tariff"])
                    quantity = int(each["quantity"])
                    material_tariff = float(each["material_rate_tariff"])
                    wash_clean_tariff = float(each["wash_clean_tariff"])
                    total_cost = float(0)
                    labour_cost = float(0)
                    material_cost = float(0)

                    labour_cost = labour_rate * labour_hrs_tariff * quantity
                    if not wash_clean_tariff == float(0):
                        material_cost = wash_clean_tariff * quantity
                    else:
                        material_cost = material_tariff * quantity

                    total_cost = float(labour_cost) + float(material_cost)

                except:
                    TariffMaster.objects.filter(
                        client=client, location=location, site=site
                    ).delete()
                    raise ValidationError(
                        "Please Provide Valid Data for Tariff Calculation"
                    )

                objList.append(
                    TariffMasterLine(
                        parent=parent,
                        size=None,
                        type=None,
                        tariff_code=each["tariff_code"],
                        main_component=each["main_component"],
                        component_code=each["component_code"],
                        component_description=each["component_description"],
                        location_code=each["location_code"],
                        location_description=each["location_description"],
                        damage_code=each["damage_code"],
                        damage_description=each["damage_description"],
                        material_code=each["material_code"],
                        material_description=each["material_description"],
                        repair_code=each["repair_code"],
                        repair_description=each["repair_description"],
                        unit=each["unit_code"],
                        measurement=each["unit_description"],
                        length_and_width=each["length_width"],
                        quantity=quantity,
                        labour_hrs_tariff=labour_hrs_tariff,
                        wash_clean_tariff=wash_clean_tariff,
                        material_tariff=material_tariff,
                        labour_cost=labour_cost,
                        material_cost=material_cost,
                        total_cost=total_cost,
                    )
                )
            TariffMasterLine.objects.bulk_create(objList)

    def extractSpecificLocationExcelData(self, input_excel):

        ps = openpyxl.load_workbook(input_excel, data_only=True)
        sheet = ps["Sheet1"]
        # Data Extraction
        main_location_code_raw1 = [
            sheet["A" + str(row)].value for row in range(1, sheet.max_row + 1)
        ]
        main_location_code_raw2 = [
            each for each in main_location_code_raw1 if not each is None
        ]
        main_location_code = [
            None if each == "_" else each for each in main_location_code_raw2
        ]

        specific_location_code_raw1 = [
            sheet["B" + str(row)].value for row in range(1, sheet.max_row + 1)
        ]
        specific_location_code_raw2 = [
            each for each in specific_location_code_raw1 if not each is None
        ]
        specific_location_code = [
            None if each == "_" else each for each in specific_location_code_raw2
        ]

        description_raw1 = [
            sheet["C" + str(row)].value for row in range(1, sheet.max_row + 1)
        ]
        description_raw2 = [each for each in description_raw1 if not each is None]
        description = [None if each == "_" else each for each in description_raw2]
        ps.close()

        # Checking excel data if valid
        popped = [
            main_location_code.pop(0),
            specific_location_code.pop(0),
            description.pop(0),
        ]
        headers = ["Main Location Code", "Specific Location Code", "Description"]

        if not popped == headers:
            return "Header Not Found"

        main_data = []
        for i in range(len(main_location_code)):
            main_data.append(
                TariffSpecificLocation(
                    location_code=main_location_code[i],
                    specific_location_code=specific_location_code[i],
                    specific_location_description=description[i],
                )
            )
        TariffSpecificLocation.objects.bulk_create(main_data)
