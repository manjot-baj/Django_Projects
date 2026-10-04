import pytz
from master.functions import pagination_func
from datetime import datetime, timezone
from SNP_DMS.settings.base import BASE_DIR
import os
import openpyxl
import pandas as pd
import xlsxwriter
from openpyxl import load_workbook

# error handling
from common.exceptions import ResourceNotFound, ValidationError

# models
from procurement.models import (
    ToolRoom,
    ToolRateHistory,
    RequisitionLine,
    ConsumptionLine,
    ToolCategory,
    SharedClass,
    Requisition,
)


class ToolRoomService:
    def getTool(self, pk):
        tool = ToolRoom.manager.getToolById(pk)
        return self.toolData(tool)

    def toolData(self, tool):

        return {
            "pk": tool.pk,
            "category": tool.category.name if tool.category else "",
            "category_id": tool.category.pk,
            "name": tool.name,
            "rate": tool.rate,
            "sku_code": tool.sku_code,
            "in_stock": tool.in_stock,
            "unit": tool.unit,
            "location": tool.location_id,
            "site": tool.site_id,
        }

    def updateTool(self, pk, params, user):
        tool = ToolRoom.manager.getToolById(pk)
        user = SharedClass().getAccountUser(user)
        if float(tool.rate) != float(params["rate"]):
            tool_history_params = {
                "tool_id": tool.pk,
                "previous_rate": tool.rate,
                "modified_rate": params["rate"],
                "changed_in": "ToolRoom",
                "requisition_order_no": "",
                "location_id": params["location_id"],
                "site_id": params["site_id"],
            }
            ToolRateHistory.manager.createToolRateHistory(tool_history_params)

        if user.role.name == "Admin":
            if float(params.get("in_stock")) > float(tool.in_stock):
                value = float(params.get("in_stock")) - float(tool.in_stock)
                self.automateRequisition(params, tool, value)
            ToolRoom.manager.updateToolById(pk, params)
        else:
            params.pop("in_stock")
            ToolRoom.manager.updateToolById(pk, params)

    def automateOrderNo(self, location, site):
        site_obj = SharedClass().getSiteById(site)
        label = f"ATMT{site_obj.code}-"
        params = {"status": "AUTOMATE", "location_id": location, "site_id": site}
        queryset = Requisition.manager.getRequisitions(params)
        count = Requisition.manager.aggregateMaxOrderNoCount(queryset)
        if count:
            maxSerialNo = int(count.split("-")[1]) + 1
        else:
            maxSerialNo = 1

        return label + str(maxSerialNo)

    def automateRequisition(self, params, tool, value):
        amount = float(tool.rate) * value
        order_no = self.automateOrderNo(
            params.get("location_id"), params.get("site_id")
        )
        requisition_params = {
            "order_no": order_no,
            "date": datetime.now(pytz.timezone("Asia/Kolkata")).date(),
            "status": "AUTOMATE",
            "total_amount": amount,
            "location_id": params.get("location_id"),
            "site_id": params.get("site_id"),
        }

        requisition = Requisition.manager.createRequisition(requisition_params)
        line_params = {
            "parent": requisition,
            "tool_id": tool.pk,
            "required_qty": value,
            "received_qty": value,
            "remaining_qty": 0,
            "amount": amount,
            "rate": tool.rate,
        }
        RequisitionLine.manager.createRequisitionLine(line_params)

    def toolList(self, payload):
        category = payload.get("category")
        name = payload.get("name")
        sku_code = payload.get("sku_code")
        location = payload.get("location")
        site = payload.get("site")
        from_date = payload.get("from_date")
        to_date = payload.get("to_date")
        on_page_data = payload.get("on_page_data")
        pg_no = payload.get("pg_no")
        process = payload.get("process")
        param = {}

        if from_date and to_date:
            try:
                from_date_obj = datetime.strptime(from_date, "%Y-%m-%d")
                to_date_obj = datetime.strptime(to_date, "%Y-%m-%d")
            except:
                raise ValidationError("Wrong date format")
            param["parent__date__range"] = (from_date_obj, to_date_obj)
            if category:
                param["tool__category_id"] = category
            if name:
                param["tool__name"] = name
            if sku_code:
                param["tool__sku_code"] = sku_code
            if location:
                param["tool__location_id"] = location
            if site:
                param["tool__site_id"] = site

            if process == "Requisition":
                filtered_data = RequisitionLine.manager.requisitionDataForToolRoom(
                    param
                )
            else:
                filtered_data = ConsumptionLine.manager.consumptionDataForToolRoom(
                    param
                )

        else:
            if category:
                param["category_id"] = category
            if name:
                param["name"] = name
            if sku_code:
                param["sku_code"] = sku_code
            if location:
                param["location_id"] = location
            if site:
                param["site_id"] = site
            filtered_data = ToolRoom.manager.getToolList(param)

        (
            no_of_data_count,
            on_page_data_count,
            no_of_pages,
            prev_page,
            next_page,
            current_page,
        ) = pagination_func(filtered_data, on_page_data, pg_no)
        if process == "":
            data = [self.toolData(each) for each in current_page.object_list]
        elif process == "Requisition":
            data = self.requisitionConsumptionFilteredData(
                current_page.object_list, "Requisition"
            )
        elif process == "Consumption":
            data = self.requisitionConsumptionFilteredData(
                current_page.object_list, "Consumption"
            )
        return {
            "no_of_data": no_of_data_count,
            "on_page_data": on_page_data_count,
            "total_pages": no_of_pages,
            "prev_page": prev_page,
            "next_page": next_page,
            "data": data,
        }

    def requisitionConsumptionFilteredData(self, object_list, type):
        if type == "Requisition":
            return [
                {
                    "pk": each["tool__pk"],
                    "category": each["tool__category__name"],
                    "name": each["tool__name"],
                    "in_stock": each["tool__in_stock"],
                    "quantity_purchased": each["quantity_purchased"],
                    "purchased_amount": each["purchased_amount"],
                }
                for each in object_list
            ]
        else:
            return [
                {
                    "pk": each["tool__pk"],
                    "category": each["tool__category__name"],
                    "name": each["tool__name"],
                    "in_stock": each["tool__in_stock"],
                    "quantity_consumed": each["quantity_consumed"],
                }
                for each in object_list
            ]

    def deleteTool(self, data):
        not_deleted = []
        for pk in data:
            try:
                obj = ToolRoom.manager.getToolById(pk)
            except:
                raise ResourceNotFound("Tool not Found")
            dependent_list = [
                obj.tool_room_consumption_line_rel.exists(),
                obj.tool_room_requisition_line_rel.exists(),
            ]
            if not any(dependent_list):
                obj.delete()
            else:
                not_deleted.append(obj.name)
        return not_deleted

    def extractDataList(self, input_excel):
        ps = openpyxl.load_workbook(input_excel, read_only=True)

        sheet = ps["tool_master"]

        category_raw, name_raw, sku_code_raw, rate_raw, unit_raw, quantity_raw = (
            [],
            [],
            [],
            [],
            [],
            [],
        )

        for row in sheet.iter_rows():

            category_raw.append(row[0].value or "")
            name_raw.append(row[1].value or "")
            sku_code_raw.append(row[2].value or "")

            # rate
            rate_value = row[3].value
            if rate_value == "RATE":
                rate_raw.append("RATE")
            elif rate_value is not None:
                try:
                    rate_raw.append(float(rate_value))
                except (ValueError, TypeError):
                    rate_raw.append("")
            else:
                rate_raw.append("")

            unit_raw.append(row[4].value or "")

            # quantity
            quantity_value = row[5].value
            if quantity_value == "QUANTITY":
                quantity_raw.append("QUANTITY")
            elif quantity_value is not None:
                try:
                    quantity_raw.append(float(quantity_value))
                except (ValueError, TypeError):
                    quantity_raw.append("")
            else:
                quantity_raw.append("")

        headers = [
            category_raw.pop(0),
            name_raw.pop(0),
            sku_code_raw.pop(0),
            rate_raw.pop(0),
            unit_raw.pop(0),
            quantity_raw.pop(0),
        ]
        if headers != ["CATEGORY", "ITEM NAME", "SKU CODE", "RATE", "UNIT", "QUANTITY"]:
            return "Header Not Found"

        data = [
            (category, name, sku_code, rate, unit, quantity)
            for category, name, sku_code, rate, unit, quantity in zip(
                category_raw, name_raw, sku_code_raw, rate_raw, unit_raw, quantity_raw
            )
            if category or name or sku_code or rate or unit or quantity
        ]

        extracted_data_list = [
            {
                "sr_no": str(i),
                "category": str(category),
                "name": str(name),
                "sku_code": str(sku_code),
                "rate": str(rate),
                "unit": str(unit),
                "quantity": str(quantity),
            }
            for i, (category, name, sku_code, rate, unit, quantity) in enumerate(
                data, start=1
            )
        ]

        ps.close()

        return extracted_data_list

    def checkErrors(self, extracted_data_list, location, site):
        error_data_msg = {}
        correct_data = []
        error_data = []

        for each in extracted_data_list:
            error_msg = []
            tool_name = each["name"]

            if each["category"]:
                error_msg.append("")
            else:
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 1)} there is problem in CATEGORY column, "
                    f"CATEGORY is Mandatory"
                )
            params = {"name": tool_name, "location_id": location, "site_id": site}
            if not tool_name:
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 1)} there is problem in ITEM NAME column, "
                    f"ITEM NAME is Mandatory"
                )
            elif ToolRoom.manager.toolExists(params):
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 1)} there is problem in NAME column, "
                    f"({tool_name}) already exists in system"
                )

            elif not SharedClass().getSiteById(site).procurement_admin:
                params.pop("site")
                if ToolRoom.manager.toolExists(params):
                    error_msg.append(
                        f"In row {str(int(each['sr_no']) + 1)} there is problem in NAME column, "
                        f"({tool_name}) already exists in parent location"
                    )
            else:
                error_msg.append("")

            if each["rate"]:
                error_msg.append("")
            else:
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 1)} there is problem in RATE column, "
                    f"RATE is Mandatory"
                )
            if each["sku_code"]:
                error_msg.append("")
            else:
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 1)} there is problem in SKU CODE column, "
                    f"SKU_CODE is Mandatory"
                )

            if each["unit"]:
                error_msg.append("")
            else:
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 1)} there is problem in UNIT column, "
                    f"UNIT is Mandatory"
                )

            if each["quantity"]:
                error_msg.append("")
            else:
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 1)} there is problem in QUANTITY column, "
                    f"QUANTITY is Mandatory"
                )

            if float(each["quantity"]) < 0:
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 1)} there is problem in QUANTITY column, "
                    f"Negative QUANTITY found"
                )
            else:
                error_msg.append("")

            if float(each["rate"]) < 0:
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 1)} there is problem in RATE column, "
                    f"Negative RATE found"
                )
            else:
                error_msg.append("")

            if (each["quantity"] == "0" and each["rate"] != "0") or (
                each["quantity"] != "0" and each["rate"] == "0"
            ):
                error_msg.append(
                    f"In row {str(int(each['sr_no']) + 1)} there is problem in QUANTITY and RATE column, "
                    f"Either both Quantity and Rate can be zero, or both can have some value"
                )
            else:
                error_msg.append("")

            error_data_msg[f"row {str(int(each['sr_no']) + 1)}"] = error_msg
            if any(error_data_msg[f"row {str(int(each['sr_no']) + 1)}"]) is True:
                error_data.append(each)
            if not each in error_data:
                correct_data.append(each)

        return correct_data, error_data, error_data_msg

    def collectMainData(self, correct_data, error_data, error_data_msg):
        correct_data_count = str(len(correct_data))
        error_data_count = str(len(error_data))
        return {
            "importable_data": correct_data,
            "importable_data_count": correct_data_count,
            "rejected_data": error_data,
            "rejected_data_count": error_data_count,
            "faults": error_data_msg,
        }

    def checkToolExistence(self, payload, location, site):
        count = 0
        while count < len(payload):
            params = {
                "name": payload[count]["name"],
                "location_id": location,
                "site_id": site,
            }
            if ToolRoom.manager.toolExists(params):
                return True
            count += 1
        return False

    def uploadBulkToolData(self, importable_data, location, site):
        for each in importable_data:
            category = ToolCategory.manager.getOrCreateCategoryByName(
                each.get("category")
            )
            params = {
                "location_id": location,
                "site_id": site,
                "category": category,
                "name": each.get("name"),
                "sku_code": each.get("sku_code"),
                "rate": each.get("rate"),
                "unit": each.get("unit"),
                "in_stock": each.get("quantity"),
            }

            tool = ToolRoom.manager.createTool(params)

            if float(params.get("in_stock")) > float(0):
                value = float(tool.in_stock)
                self.automateRequisition(params, tool, value)

    def extractDataForRejectedFile(self, data):
        tool_room_df_data = pd.DataFrame(data)[
            ["category", "name", "sku_code", "rate", "unit", "quantity"]
        ]
        tool_room_df_data.columns = [
            "CATEGORY",
            "ITEM NAME",
            "SKU CODE",
            "RATE",
            "UNIT",
            "QUANTITY",
        ]
        return tool_room_df_data

    def writeDataToExcel(self, new_temp_file_path, tool_room_df_data, faults):
        book = load_workbook(new_temp_file_path)

        with pd.ExcelWriter(new_temp_file_path, engine="openpyxl") as writer:

            writer.workbook = book

            writer.worksheets = dict((ws.title, ws) for ws in book.worksheets)

            tool_room_df_data.to_excel(
                writer,
                sheet_name="tool_master",
                startrow=0,
                startcol=0,
                index=False,
            )
            pd.DataFrame(faults).to_excel(
                writer, sheet_name="faults", startrow=0, startcol=0, index=False
            )

    def toolListForDropdownByCategory(self, id, location, site):

        params = {"category_id": id, "location_id": location, "site_id": site}
        tool_list = ToolRoom.manager.getToolNameList(params)

        return [{"pk": each["pk"], "name": each["name"]} for each in tool_list]

    def toolInfo(self, payload):

        data = ToolRoom.manager.getToolById(payload.get("id"))
        return {
            "sku_code": data.sku_code,
            "rate": data.rate,
            "in_stock": data.in_stock,
        }

    def toolCategories(self, location, site):

        return ToolRoom.manager.getToolCategoriesByLocationSite(location, site)

    def toolListForDropdown(self, location, site):
        params = {"location_id": location, "site_id": site}

        tool_list = ToolRoom.manager.getToolNameList(params)

        return [{"pk": each["pk"], "name": each["name"]} for each in tool_list]

    def skuCodes(self, location, site):
        return ToolRoom.manager.getSKUCodes(location, site)

    def toolRateHistory(self, params, payload):
        filtered_data = ToolRateHistory.manager.getToolRateHistory(params)
        (
            no_of_data_count,
            on_page_data_count,
            no_of_pages,
            prev_page,
            next_page,
            current_page,
        ) = pagination_func(
            filtered_data, payload.get("on_page_data"), payload.get("pg_no")
        )
        main_data = self.getToolHistoryData(current_page.object_list)
        return {
            "no_of_data": no_of_data_count,
            "on_page_data": on_page_data_count,
            "total_pages": no_of_pages,
            "prev_page": prev_page,
            "next_page": next_page,
            "data": main_data,
        }

    def getToolHistoryData(self, object_list):
        return [
            {
                "created_at": each["created_at"].date().strftime("%Y-%m-%d"),
                "item": each["tool__name"],
                "category": each["tool__category__name"],
                "modified_rate": each["modified_rate"],
                "previous_rate": each["previous_rate"],
                "location": each["location__name"],
                "site": each["site__name"],
            }
            for each in object_list
        ]

    def inventoryDfData(self, data_object_list):
        df_data = [
            [
                i + 1,
                each.get("item_name", ""),
                each.get("sku_code", ""),
                each.get("unit", ""),
                each.get("opening_stock", ""),
                each.get("received_quantity", ""),
                each.get("delivered_quantity", ""),
                each.get("closing_stock", ""),
                each.get("rate", ""),
                each.get("approx_value", ""),
            ]
            for i, each in enumerate(data_object_list, start=0)
        ]
        # Make sure that df data has at least one row
        df_data = df_data or [["" for _ in range(9)]]

        return df_data

    def createInventoryReport(self, from_date, to_date, df_data):
        base_temp_dir = os.path.join(BASE_DIR, "temp")
        os.makedirs(base_temp_dir, exist_ok=True)

        temp_file_path = os.path.join(
            base_temp_dir, f"inventory_report_{from_date}_to_{to_date}.xlsx"
        )

        workbook = xlsxwriter.Workbook(temp_file_path)
        inventory_sheet = workbook.add_worksheet("INVENTORY REPORT")

        report_name = "REPORT NAME : INVENTORY REPORT(* Closing and Opening Balance is based on Requisition and Consumption and not on manually added stock quantity)"
        report_date = f"REPORT DATE : {from_date}  to {to_date} "
        merge_format, merge_format3 = SharedClass().getMergeFormat(workbook)

        inventory_sheet.merge_range(
            "A1:U2", "Golden Horn Containers Service", merge_format
        )
        inventory_sheet.merge_range("A3:U3", report_name, merge_format3)
        inventory_sheet.merge_range("A4:U4", report_date, merge_format3)
        inventory_sheet.merge_range("A5:U5", None, merge_format3)

        inventory_sheet.add_table(
            f"A7:J{7 + len(df_data)}",
            {
                "data": df_data,
                "columns": [
                    {"header": "Sl.No."},
                    {"header": "Item Name"},
                    {"header": "SKU Code"},
                    {"header": "Usage Unit"},
                    {"header": "Opening Stock"},
                    {"header": "Received"},
                    {"header": "Delivered"},
                    {"header": "Closing Stock"},
                    {"header": "Rate"},
                    {"header": "Approximate Value In Rupee"},
                ],
            },
        )
        workbook.close()
        return temp_file_path

    def stockDfData(self, data_object_list):
        df_data = [
            [
                i + 1,
                each.get("tool", ""),
                each.get("sku_code", ""),
                each.get("unit", ""),
                each.get("stock_on_hand", ""),
                each.get("avg_price", ""),
                each.get("rate", ""),
            ]
            for i, each in enumerate(data_object_list, start=0)
        ]
        # Make sure that df data has at least one row
        df_data = df_data or [["" for _ in range(7)]]

        return df_data

    def createStockReport(self, from_date, to_date, df_data):
        base_temp_dir = os.path.join(BASE_DIR, "temp")
        os.makedirs(base_temp_dir, exist_ok=True)

        temp_file_path = os.path.join(
            base_temp_dir, f"stock_report_{datetime.today().date()}.xlsx"
        )

        workbook = xlsxwriter.Workbook(temp_file_path)
        stock_sheet = workbook.add_worksheet("STOCK REPORT")

        report_name = "REPORT NAME : STOCK REPORT"
        report_date = f"REPORT DATE : {from_date}  to {to_date} "
        merge_format, merge_format3 = SharedClass().getMergeFormat(workbook)

        stock_sheet.merge_range("A1:U2", "Golden Horn Containers Service", merge_format)
        stock_sheet.merge_range("A3:U3", report_name, merge_format3)
        stock_sheet.merge_range("A4:U4", report_date, merge_format3)
        stock_sheet.merge_range("A5:U5", None, merge_format3)

        stock_sheet.add_table(
            f"A7:G{7 + len(df_data)}",
            {
                "data": df_data,
                "columns": [
                    {"header": "Sr.No."},
                    {"header": "Item Name"},
                    {"header": "SKU Code"},
                    {"header": "Unit"},
                    {"header": "Stock on Hand"},
                    {"header": "Amount"},
                    {"header": "Average Price"},
                ],
            },
        )
        workbook.close()
        return temp_file_path

    def getMasterStockTableData(self, params, on_page_data, pg_no):
        filtered_data = ToolRoom.manager.getMasterStockData(params)
        (
            no_of_data_count,
            on_page_data_count,
            no_of_pages,
            prev_page,
            next_page,
            current_page,
        ) = pagination_func(filtered_data, on_page_data, pg_no)
        main_data = [
            {
                "pk": each["pk"],
                "category": each["category__name"],
                "name": each["name"],
                "sku_code": each["sku_code"],
                "rate": each["rate"],
                "in_stock": each["in_stock"],
                "unit": each["unit"],
                "location": each["location__name"],
                "site": each["site__name"],
            }
            for each in current_page.object_list
        ]

        return {
            "no_of_data": no_of_data_count,
            "on_page_data": on_page_data_count,
            "total_pages": no_of_pages,
            "prev_page": prev_page,
            "next_page": next_page,
            "data": main_data,
        }

    def masterStockReportObjectList(self, params):
        procurement_enabled_sites = SharedClass().procurementEnabledSites(params)
        object_list = []
        for each in procurement_enabled_sites:
            object_list.append(
                (f"{each}", ToolRoom.manager.dataForMasterStock(each, params))
            )
        return object_list

    def masterStockDfData(self, data_object_list):
        df_data = []
        for key, value in data_object_list:
            data = [
                [
                    i + 1,
                    each.get("name", ""),
                    each.get("unit", ""),
                    each.get("in_stock", ""),
                    each.get("rate", ""),
                ]
                for i, each in enumerate(value, start=0)
            ]
            data = data or [["" for _ in range(5)]]
            df_data.append((f"{key}", data))

        return df_data

    def createMasterStockReport(self, df_data):
        base_temp_dir = os.path.join(BASE_DIR, "temp")
        os.makedirs(base_temp_dir, exist_ok=True)

        temp_file_path = os.path.join(
            base_temp_dir, f"master_stock_report_{datetime.today().date()}.xlsx"
        )

        workbook = xlsxwriter.Workbook(temp_file_path)
        for key, value in df_data:
            stock_sheet = workbook.add_worksheet(f"{key}")

            report_name = "REPORT NAME : MASTER STOCK REPORT"
            report_date = f"REPORT DATE : As on {datetime.today().date()} "
            merge_format, merge_format3 = SharedClass().getMergeFormat(workbook)

            stock_sheet.merge_range(
                "A1:U2", "Golden Horn Containers Service", merge_format
            )
            stock_sheet.merge_range("A3:U3", report_name, merge_format3)
            stock_sheet.merge_range("A4:U4", report_date, merge_format3)
            stock_sheet.merge_range("A5:U5", None, merge_format3)

            stock_sheet.add_table(
                f"A7:E{7 + len(value)}",
                {
                    "data": value,
                    "columns": [
                        {"header": "Sr.No."},
                        {"header": "Item Name"},
                        {"header": "Unit"},
                        {"header": "Stock on Hand"},
                        {"header": "Rate"},
                    ],
                },
            )
        workbook.close()
        return temp_file_path
