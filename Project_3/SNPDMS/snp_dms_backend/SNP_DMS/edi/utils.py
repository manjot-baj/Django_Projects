from datetime import datetime, timedelta
import openpyxl, xlsxwriter, os
from SNP_DMS.settings.base import BASE_DIR
from depot.models import ContainerStock
from mnr.models import Approval, Repair
from django.utils import timezone


def getStartAndEndDate(month):
    current_year = datetime.now().year

    start_date = datetime(current_year, month, 1)
    if month == 12:
        end_date = datetime(current_year, 12, 31)
    else:
        next_month = month + 1
        start_next_month = datetime(current_year, next_month, 1)
        end_date = start_next_month - timedelta(days=1)
    return start_date, end_date


def appendData(queryset, key, data):
    for each in queryset:
        date_str = each["created_at"].strftime("%Y-%m-%d")
        if date_str in data.keys():
            data[date_str][key] += each["count"]


def calculateSuccessRate(data):
    for _, value in data.items():
        if value["in_count"] == 0 or value["in_edi_sent_count"] == 0:
            value["in_success_rate"] = "0%"
        else:
            value["in_success_rate"] = (
                str(
                    round(
                        (value["in_edi_sent_count"] / value["in_count"]) * 100,
                        2,
                    )
                )
                + "%"
            )
        if value["out_count"] == 0 or value["out_edi_sent_count"] == 0:
            value["out_success_rate"] = "0%"
        else:
            value["out_success_rate"] = (
                str(
                    round(
                        (value["out_edi_sent_count"] / value["out_count"]) * 100,
                        2,
                    )
                )
                + "%"
            )


def getInEDIMailDfData(data_object_list):
    df_data = [
        [
            i + 1,
            each.container.container_no,
            f"{each.container.size.name}{each.container.type.name}",
            each.container.client.name if each.container.client else "",
            each.lolo.customer_name.name if each.lolo.customer_name else "",
            each.container.client.ref_code if each.container.client else "",
            each.gate_in.arrived,
            each.gate_in.condition,
            "Yes" if each.is_email_sent else "No",
            each.container.gross_wt,
            each.container.tare_wt,
            each.container.payload,
            (
                ""
                if each.container.manufacturing_date == ""
                else each.container.manufacturing_date.strftime("%d/%m/%Y")
            ),
            each.gate_in.grade or "",
            each.gate_in.remarks or "",
            timezone.localtime(each.created_at).date().strftime("%b %d %Y")
            + " "
            + timezone.localtime(each.created_at).time().strftime("%I:%M %p"),
            each.gate_in.in_date.strftime("%b %d %Y")
            + " "
            + each.gate_in.in_time.strftime("%I:%M %p"),
            each.gate_in.bl_no if each.gate_in.bl_no else "",
            each.gate_in.shipper or "",
            each.gate_in.vessel_name or "",
            each.gate_in.voyage_no or "",
            each.gate_in.source or "",
            each.gate_in.transporter_name.name if each.gate_in.transporter_name else "",
            each.gate_in.vehicle_no,
            each.gate_in.cargo or "",
        ]
        for i, each in enumerate(data_object_list, start=0)
    ]
    df_data = df_data or [["" for _ in range(25)]]

    return df_data


def getOutEDIMailDfData(data_object_list):
    df_data = []
    for i, each in enumerate(data_object_list, start=0):
        stock = ContainerStock.objects.get(
            container=each.container, gate_out=each.gate_out
        )
        df_data.append(
            [
                i + 1,
                each.container.container_no,
                f"{each.container.size.name}{each.container.type.name}",
                each.container.client.name if each.container.client else "",
                each.lolo.customer_name.name if each.lolo.customer_name else "",
                each.container.client.ref_code if each.container.client else "",
                each.gate_out.departed,
                each.gate_out.condition or "",
                "Yes" if each.is_email_sent else "No",
                each.container.gross_wt or "",
                each.container.tare_wt or "",
                each.container.payload or "",
                (
                    ""
                    if each.container.manufacturing_date == ""
                    else each.container.manufacturing_date.strftime("%d/%m/%Y")
                ),
                each.gate_out.grade or "",
                each.gate_out.remarks or "",
                timezone.localtime(each.created_at).date().strftime("%b %d %Y")
                + " "
                + timezone.localtime(each.created_at).time().strftime("%I:%M %p"),
                stock.gate_in.in_date.strftime("%b %d %Y")
                + " "
                + stock.gate_in.in_time.strftime("%I:%M %p"),
                each.gate_out.out_date.strftime("%b %d %Y")
                + " "
                + each.gate_out.out_time.strftime("%I:%M %p"),
                each.gate_out.shipper or "",
                each.gate_out.vessel_name,
                each.gate_out.voyage_no,
                each.gate_out.destination or "",
                (
                    each.gate_out.transporter_name.name
                    if each.gate_out.transporter_name
                    else ""
                ),
                each.gate_out.vehicle_no or "",
                each.gate_out.export_cargo or "",
            ]
        )

    df_data = df_data or [["" for _ in range(25)]]

    return df_data


def get_merge_format(workbook):
    merge_format = workbook.add_format(
        {
            "bold": 2,
            "align": "center",
            "valign": "vcenter",
            "font_color": "#e0a92a",
            "font_size": 18,
        }
    )
    merge_format3 = workbook.add_format(
        {
            "bold": 1,
            "font_size": 11,
        }
    )
    return merge_format, merge_format3


def createEdiReport(
    inward_df_data, outward_df_data, start_date, end_date, name, location, site, line
):
    base_temp_dir = os.path.join(BASE_DIR, "temp")
    os.makedirs(base_temp_dir, exist_ok=True)

    temp_file_path = os.path.join(
        base_temp_dir, f"edi_report_{start_date}_to_{end_date}.xlsx"
    )

    workbook = xlsxwriter.Workbook(temp_file_path)
    inward_sheet = workbook.add_worksheet("INWARD")

    report_name = f"REPORT NAME : {name}"
    report_date = f"REPORT DATE : {start_date}  to {end_date} "
    merge_format, merge_format3 = get_merge_format(workbook)

    inward_sheet.merge_range("A1:U2", "Golden Horn Containers Service", merge_format)
    inward_sheet.merge_range("A3:U3", report_name, merge_format3)
    inward_sheet.merge_range("A4:U4", report_date, merge_format3)
    inward_sheet.merge_range("A5:U5", f"SHIPPING LINE: {line}", merge_format3)
    inward_sheet.merge_range(
        "A6:U6", f"DEPOT NAME: {site}, LOCATION: {location}", merge_format3
    )
    inward_sheet.merge_range("A7:U7", None, merge_format3)

    inward_sheet.add_table(
        f"A9:Y{9 + len(inward_df_data)}",
        {
            "data": inward_df_data,
            "columns": [
                {"header": "Sl.No"},
                {"header": "ContainerNo"},
                {"header": "Size Type"},
                {"header": "Line"},
                {"header": "Customer"},
                {"header": "OPR"},
                {"header": "Arrived"},
                {"header": "Status"},
                {"header": "EDI Sent"},
                {"header": "Gross Wt"},
                {"header": "Tare Wt"},
                {"header": "Payload"},
                {"header": "MfgDate"},
                {"header": "Grade"},
                {"header": "Remarks"},
                {"header": "Created At"},
                {"header": "GateInDate"},
                {"header": "BL No"},
                {"header": "Shipper"},
                {"header": "Vessel"},
                {"header": "Voyage"},
                {"header": "Place"},
                {"header": "Transporter"},
                {"header": "VehicleNo"},
                {"header": "ImportCargo"},
            ],
        },
    )
    outward_sheet = workbook.add_worksheet("OUTWARD")

    report_name = f"REPORT NAME : {name}"
    report_date = f"REPORT DATE : {start_date}  to {end_date} "
    merge_format, merge_format3 = get_merge_format(workbook)

    outward_sheet.merge_range("A1:U2", "Golden Horn Containers Service", merge_format)
    outward_sheet.merge_range("A3:U3", report_name, merge_format3)
    outward_sheet.merge_range("A4:U4", report_date, merge_format3)
    outward_sheet.merge_range("A5:U5", f"SHIPPING LINE: {line}", merge_format3)
    outward_sheet.merge_range(
        "A6:U6", f"DEPOT NAME: {site}, LOCATION: {location}", merge_format3
    )

    outward_sheet.merge_range("A7:U7", None, merge_format3)

    outward_sheet.add_table(
        f"A9:Y{9 + len(outward_df_data)}",
        {
            "data": outward_df_data,
            "columns": [
                {"header": "Sl.No"},
                {"header": "ContainerNo"},
                {"header": "Size Type"},
                {"header": "Line"},
                {"header": "Customer"},
                {"header": "OPR"},
                {"header": "Departed"},
                {"header": "Status"},
                {"header": "EDI Sent"},
                {"header": "GrossWt"},
                {"header": "TareWt"},
                {"header": "Payload"},
                {"header": "MfgDate"},
                {"header": "Grade"},
                {"header": "Remarks"},
                {"header": "Created At"},
                {"header": "GateInDate"},
                {"header": "GateOutDate"},
                {"header": "Shipper"},
                {"header": "Vessel"},
                {"header": "Voyage"},
                {"header": "Place"},
                {"header": "Transporter"},
                {"header": "VehicleNo"},
                {"header": "ExportCargo"},
            ],
        },
    )
    workbook.close()
    return temp_file_path


def getTextExcelMoveCodeDFData(sent_data, waiting_data, edi_type):
    sent_df_data = [
        [
            i + 1,
            each.container_no,
            each.move_code if edi_type == "text" else each.excel_edi_move_code,
            each.email_date.date().strftime("%b %d %Y")
            + " "
            + each.email_date.time().strftime("%I:%M %p"),
            each.process_date.date().strftime("%b %d %Y")
            + " "
            + each.process_date.time().strftime("%I:%M %p"),
            each.time_diff or "",
        ]
        for i, each in enumerate(sent_data, start=0)
    ]
    sent_df_data = sent_df_data or [["" for _ in range(7)]]

    waiting_df_data = [
        [
            i + 1,
            each.container_no,
            each.move_code,
            timezone.localtime(each.created_at).date().strftime("%b %d %Y")
            + " "
            + timezone.localtime(each.created_at).time().strftime("%I:%M %p"),
            each.date.date().strftime("%b %d %Y")
            + " "
            + each.date.time().strftime("%I:%M %p"),
        ]
        for i, each in enumerate(waiting_data, start=0)
    ]
    waiting_df_data = waiting_df_data or [["" for _ in range(5)]]

    return sent_df_data, waiting_df_data


def createMoveCodeReport(
    sent_df_data, waiting_df_data, start_date, end_date, location, site
):
    base_temp_dir = os.path.join(BASE_DIR, "temp")
    os.makedirs(base_temp_dir, exist_ok=True)

    temp_file_path = os.path.join(
        base_temp_dir, f"edi_report_{start_date}_to_{end_date}.xlsx"
    )

    workbook = xlsxwriter.Workbook(temp_file_path)
    sent_edi_sheet = workbook.add_worksheet("SENT")
    waiting_edi_sheet = workbook.add_worksheet("WAITING")

    report_name = f"REPORT NAME : SENT EDI"
    report_date = f"REPORT DATE : {start_date}  to {end_date} "

    merge_format, merge_format3 = get_merge_format(workbook)

    sent_edi_sheet.merge_range("A1:U2", "Golden Horn Containers Service", merge_format)
    sent_edi_sheet.merge_range("A3:U3", report_name, merge_format3)
    sent_edi_sheet.merge_range("A4:U4", report_date, merge_format3)
    sent_edi_sheet.merge_range("A5:U5", f"SHIPPING LINE: MSC", merge_format3)
    sent_edi_sheet.merge_range(
        "A6:U6", f"DEPOT NAME: {site}, LOCATION: {location}", merge_format3
    )

    sent_edi_sheet.merge_range("A7:U7", None, merge_format3)

    sent_edi_sheet.add_table(
        f"A9:F{9 + len(sent_df_data)}",
        {
            "data": sent_df_data,
            "columns": [
                {"header": "Sl.No"},
                {"header": "ContainerNo"},
                {"header": "Move Code"},
                {"header": "Email Date"},
                {"header": "Process Date"},
                {"header": "Time Difference"},
            ],
        },
    )

    report_name = f"REPORT NAME : WAITING EDI"
    report_date = f"REPORT DATE : {start_date}  to {end_date} "

    merge_format, merge_format3 = get_merge_format(workbook)

    waiting_edi_sheet.merge_range(
        "A1:U2", "Golden Horn Containers Service", merge_format
    )
    waiting_edi_sheet.merge_range("A3:U3", report_name, merge_format3)
    waiting_edi_sheet.merge_range("A4:U4", report_date, merge_format3)

    waiting_edi_sheet.merge_range("A5:U5", f"SHIPPING LINE: MSC", merge_format3)
    waiting_edi_sheet.merge_range(
        "A6:U6", f"DEPOT NAME: {site}, LOCATION: {location}", merge_format3
    )
    waiting_edi_sheet.merge_range("A7:U7", None, merge_format3)

    waiting_edi_sheet.add_table(
        f"A9:E{9 + len(waiting_df_data)}",
        {
            "data": waiting_df_data,
            "columns": [
                {"header": "Sl.No"},
                {"header": "ContainerNo"},
                {"header": "Move Code"},
                {"header": "Created Date"},
                {"header": "Date"},
            ],
        },
    )
    workbook.close()
    return temp_file_path


def getEstimateWistimDfData(data_object_list, site_type):
    df_data = []

    for i, each in enumerate(data_object_list, start=0):
        if site_type == "DEPOT":
            param = {"parent__parent__depot__pk": each[0]}
        else:
            param = {"parent__parent__non_depot__pk": each[0]}

        try:
            approval = Approval.objects.get(**param)
            if approval.is_approved:
                status = "Approved"
            elif approval.is_denied:
                if approval.denial_reason == "PARTIALLY":
                    status = "Partially Approved"
                elif approval.denial_reason == "REJECTED":
                    status = "Rejected"
                elif approval.denial_reason == "CANCEL":
                    status = "Cancelled"
        except:
            status = ""

        df_data.append(
            [
                i + 1,
                each[1],
                f"{each[2]}{each[3]}",
                each[4],
                each[5],
                each[6],
                each[7].strftime("%b %d %Y") + " " + each[8].strftime("%I:%M %p"),
                "Yes" if each[9] else "No",
                "Yes" if each[10] else "No",
                status,
            ]
        )

    df_data = df_data or [["" for _ in range(9)]]

    return df_data


def createEstimateWistimReport(df_data, start_date, end_date, location, site, line):
    base_temp_dir = os.path.join(BASE_DIR, "temp")
    os.makedirs(base_temp_dir, exist_ok=True)

    temp_file_path = os.path.join(
        base_temp_dir, f"estimate_edi_report_{start_date}_to_{end_date}.xlsx"
    )

    workbook = xlsxwriter.Workbook(temp_file_path)
    estimate_wistim_sheet = workbook.add_worksheet("ESTIMATE_WISTIM")

    report_name = f"REPORT NAME : ESTIMATE WISTIM AND REPAIR DISTIM REPORT"
    report_date = f"REPORT DATE : {start_date}  to {end_date} "
    merge_format, merge_format3 = get_merge_format(workbook)

    estimate_wistim_sheet.merge_range(
        "A1:U2", "Golden Horn Containers Service", merge_format
    )
    estimate_wistim_sheet.merge_range("A3:U3", report_name, merge_format3)
    estimate_wistim_sheet.merge_range("A4:U4", report_date, merge_format3)
    estimate_wistim_sheet.merge_range("A5:U5", f"SHIPPING LINE: {line}", merge_format3)
    estimate_wistim_sheet.merge_range(
        "A6:U6", f"LOCATION: {location}, DEPOT NAME: {site}", merge_format3
    )
    estimate_wistim_sheet.merge_range("A7:U7", None, merge_format3)

    estimate_wistim_sheet.add_table(
        f"A9:J{9 + len(df_data)}",
        {
            "data": df_data,
            "columns": [
                {"header": "Sl.No"},
                {"header": "ContainerNo"},
                {"header": "Size Type"},
                {"header": "Client"},
                {"header": "Line"},
                {"header": "Condition"},
                {"header": "Gate In Date Time"},
                {"header": "Estimate Westim Sent"},
                {"header": "Repair Destim Sent"},
                {"header": "Approval Status"},
            ],
        },
    )
    workbook.close()
    return temp_file_path


def getRepairDistimDfData(data_object_list, site_type):
    df_data = []

    for i, each in enumerate(data_object_list, start=0):
        if site_type == "DEPOT":
            param = {"parent__parent__depot__pk": each["pk"]}
        else:
            param = {"parent__parent__non_depot__pk": each["pk"]}
        gate_in_date_time = datetime.combine(
            each["gate_in__in_date"],
            each["gate_in__in_time"],
        )
        try:
            approval = Approval.objects.get(**param)

            approval_date_time = datetime.combine(
                approval.approved_date, approval.approved_time
            )
            time_difference = approval_date_time - gate_in_date_time

            approval_span = (
                "Within 48 Hours"
                if time_difference <= timedelta(hours=48)
                else "Above 48 Hours"
            )
        except:
            approval_span = ""

        try:
            repair = Repair.objects.get(**param)
            repair_date_time = datetime.combine(repair.repair_date, repair.repair_time)
            time_difference = repair_date_time - gate_in_date_time

            if repair.damage_category == "LD":
                if time_difference <= timedelta(days=3):
                    repair_span = "Within 3 days(LD)"
                else:
                    repair_span = "Above 3 days(LD)"

            elif repair.damage_category == "MD":
                if time_difference <= timedelta(days=5):
                    repair_span = "Within 5 days(MD)"
                else:
                    repair_span = "Above 5 days(MD)"

            elif repair.damage_category == "HD":
                if time_difference <= timedelta(days=10):
                    repair_span = "Within 10 days(HD)"
                else:
                    repair_span = "Above 10 days(HD)"
        except:
            repair_span = ""

        if each["is_estimate_westim_sent"]:

            estimate_date_time = datetime.combine(
                each["estimate_date"], each["estimate_time"]
            )
            time_difference = estimate_date_time - gate_in_date_time
            wistim_span = (
                "Within 24 hours"
                if time_difference <= timedelta(hours=24)
                else "Above 24 hours"
            )
        else:
            wistim_span = ""

        df_data.append(
            [
                i + 1,
                each["container__container_no"],
                f"{each['container__size__name']}{each['container__type__name']}",
                each["container__client__name"],
                each["container__client__ref_code"],
                each["gate_in__condition"],
                each["gate_in__in_date"].strftime("%b %d %Y")
                + " "
                + each["gate_in__in_time"].strftime("%I:%M %p"),
                wistim_span,
                approval_span,
                repair_span,
            ]
        )

    df_data = df_data or [["" for _ in range(10)]]

    return df_data


def createRepairDistimReport(df_data, start_date, end_date, location, site, line):
    base_temp_dir = os.path.join(BASE_DIR, "temp")
    os.makedirs(base_temp_dir, exist_ok=True)

    temp_file_path = os.path.join(
        base_temp_dir, f"repair_edi_report_{start_date}_to_{end_date}.xlsx"
    )

    workbook = xlsxwriter.Workbook(temp_file_path)
    repair_distim_sheet = workbook.add_worksheet("REPAIR_DISTIM")

    report_name = f"REPORT NAME : Westim and Destim HR EDI Report"
    report_date = f"REPORT DATE : {start_date}  to {end_date} "
    merge_format, merge_format3 = get_merge_format(workbook)

    repair_distim_sheet.merge_range(
        "A1:U2", "Golden Horn Containers Service", merge_format
    )
    repair_distim_sheet.merge_range("A3:U3", report_name, merge_format3)
    repair_distim_sheet.merge_range("A4:U4", report_date, merge_format3)
    repair_distim_sheet.merge_range("A5:U5", f"SHIPPING LINE: {line}", merge_format3)
    repair_distim_sheet.merge_range(
        "A6:U6", f"LOCATION: {location}, DEPOT NAME: {site}", merge_format3
    )
    repair_distim_sheet.merge_range("A7:U7", None, merge_format3)

    repair_distim_sheet.add_table(
        f"A9:J{9 + len(df_data)}",
        {
            "data": df_data,
            "columns": [
                {"header": "Sl.No"},
                {"header": "ContainerNo"},
                {"header": "Size Type"},
                {"header": "Client"},
                {"header": "Line"},
                {"header": "Condition"},
                {"header": "Gate In Date Time"},
                {"header": "Wistim Sent In"},
                {"header": "Approval Done In"},
                {"header": "Repair Done In"},
            ],
        },
    )
    workbook.close()
    return temp_file_path
