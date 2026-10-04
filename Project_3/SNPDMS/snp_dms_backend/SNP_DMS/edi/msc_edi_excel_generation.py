import logging, traceback, datetime, os, shutil, xlsxwriter
import pandas as pd
from openpyxl import load_workbook
from SNP_DMS.settings.base import BASE_DIR
from django.utils import timezone
from master.models import LocationCodeDetail
from depot.models import GateInHistory, GateOutHistory
from edi.data_collecter_functions import detect_arrived_by, detect_repair_type
from edi.models import MscExcelEdiMoveCodeInfo


def get_msc_inward_data_object(
    from_date_str, to_date_str, from_time_str, to_time_str, location, site, container_no
):
    try:
        tz = timezone.get_current_timezone()
        in_query = (
            GateInHistory.objects.select_related(
                "container",
                "container__client",
                "container__location",
                "container__site",
                "gate_in",
            )
            .filter(
                container__client__ref_code="MSC",
                container__location=location,
                container__site=site,
            )
            .order_by("date")
        )

        if not len(from_date_str) == 0:
            if not len(from_time_str) == 0:
                from_date = datetime.datetime.strptime(from_date_str, "%Y-%m-%d").date()
                to_date = datetime.datetime.strptime(to_date_str, "%Y-%m-%d").date()
                from_time = datetime.datetime.strptime(from_time_str, "%H:%M").time()
                to_time = datetime.datetime.strptime(to_time_str, "%H:%M").time()
                from_date_time = datetime.datetime.combine(
                    from_date, from_time
                ).astimezone(tz)
                to_date_time = datetime.datetime.combine(to_date, to_time).astimezone(
                    tz
                )
                inward_data = in_query.filter(
                    date__range=(from_date_time, to_date_time)
                )

                if container_no is not None:
                    if not len(container_no) == 0:
                        inward_data = inward_data.filter(
                            container__container_no__in=container_no
                        )

                in_list = list(inward_data)
                # to sort ascending order by date
                in_list.sort(key=lambda x: x.date, reverse=False)
                return in_list
            else:
                inward_data = in_query.filter(
                    gate_in__in_date__range=[from_date_str, to_date_str]
                )

                if container_no is not None:
                    if not len(container_no) == 0:
                        inward_data = inward_data.filter(
                            container__container_no__in=container_no
                        )

                in_list = list(inward_data)
                # to sort ascending order by date
                in_list.sort(key=lambda x: x.date, reverse=False)
                return in_list
        else:

            if container_no is not None:
                if not len(container_no) == 0:
                    in_list = [
                        in_query.filter(container__container_no=ctn).latest("pk")
                        for ctn in container_no
                    ]
                    # to sort ascending order by date
                    in_list.sort(key=lambda x: x.date, reverse=False)
                    return in_list

            to_date_time = datetime.datetime.now().astimezone(tz)
            from_date_time = to_date_time - datetime.timedelta(hours=24)
            from_date_time = from_date_time.astimezone(tz)
            inward_data = in_query.filter(date__range=(from_date_time, to_date_time))
            in_list = list(inward_data)
            # to sort ascending order by date
            in_list.sort(key=lambda x: x.date, reverse=False)
            return in_list
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return []


def get_msc_outward_data_object(
    from_date_str, to_date_str, from_time_str, to_time_str, location, site, container_no
):
    try:
        tz = timezone.get_current_timezone()
        out_query = (
            GateOutHistory.objects.select_related(
                "container",
                "container__client",
                "container__location",
                "container__site",
                "gate_out",
            )
            .filter(
                container__client__ref_code="MSC",
                container__location=location,
                container__site=site,
            )
            .order_by("date")
        )

        if not len(from_date_str) == 0:
            if not len(from_time_str) == 0:
                from_date = datetime.datetime.strptime(from_date_str, "%Y-%m-%d").date()
                to_date = datetime.datetime.strptime(to_date_str, "%Y-%m-%d").date()
                from_time = datetime.datetime.strptime(from_time_str, "%H:%M").time()
                to_time = datetime.datetime.strptime(to_time_str, "%H:%M").time()
                from_date_time = datetime.datetime.combine(
                    from_date, from_time
                ).astimezone(tz)
                to_date_time = datetime.datetime.combine(to_date, to_time).astimezone(
                    tz
                )
                outward_data = out_query.filter(
                    date__range=(from_date_time, to_date_time)
                )

                if container_no is not None:
                    if not len(container_no) == 0:
                        outward_data = outward_data.filter(
                            container__container_no__in=container_no
                        )

                out_list = list(outward_data)
                # to sort ascending order by date
                out_list.sort(key=lambda x: x.date, reverse=False)
                return out_list
            else:
                outward_data = out_query.filter(
                    gate_out__out_date__range=[from_date_str, to_date_str]
                )

                if container_no is not None:
                    if not len(container_no) == 0:
                        outward_data = outward_data.filter(
                            container__container_no__in=container_no
                        )

                out_list = list(outward_data)
                # to sort ascending order by date
                out_list.sort(key=lambda x: x.date, reverse=False)
                return out_list
        else:
            if container_no is not None:
                if not len(container_no) == 0:
                    out_list = [
                        out_query.filter(container__container_no=ctn).latest("pk")
                        for ctn in container_no
                    ]
                    # to sort ascending order by date
                    out_list.sort(key=lambda x: x.date, reverse=False)
                    return out_list

            to_date_time = datetime.datetime.now().astimezone(tz)
            from_date_time = to_date_time - datetime.timedelta(hours=24)
            from_date_time = from_date_time.astimezone(tz)
            outward_data = out_query.filter(date__range=(from_date_time, to_date_time))
            out_list = list(outward_data)
            # to sort ascending order by date
            out_list.sort(key=lambda x: x.date, reverse=False)
            return out_list
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return []


CONTAINER_MOVE_CODE_STATUS = {
    "DMG_MT": "DAM",
    "DVAN": "ICO",
    "MTIN": "MCY",
    "REP_OUT_MT": "TBR",
    "REP_RET_MT": "REP",
    "VAN": "MSH",
    "MIR": "MPI",
    "MOR": "MPO",
    "RET": "ERM",
}


ARRIVED_PROCESS_MAPPING = {
    "Factory": "factory_in",
    "Road/Rail": "road_in",
    "FS RETURN": "fs_return_in",
    "CFS/ICD": "cfs_in",
    "Port/Vessel": "vessel_in",
}

DEPARTED_PROCESS_MAPPING = {
    "Factory": "factory_out",
    "Road/Rail": "road_out",
    "FS RETURN": "fs_return_out",
    "CFS/ICD": "cfs_out",
    "Port/Vessel": "vessel_out",
}


def msc_edi_excel_cfs_in_first_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        # edi_partner_code = obj_data.container.site.vendor_name
        event_location = obj_data.container.site.event_location_msc_code

        try:
            if obj_data.container.site.name.lower() == "tuticorin":
                location_code_obj = LocationCodeDetail.objects.get(
                    name_code=obj_data.gate_in.from_location_name_code
                )
                depot_name = location_code_obj.depot_name
                depot_msc_code = location_code_obj.ova_code
                event_location = location_code_obj.ova_code[:5]
        except:
            pass

        move_code = "DVAN"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_in.in_date
        time = obj_data.gate_in.in_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time - datetime.timedelta(minutes=30)
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        # row_dict["edi_partner_code"] = edi_partner_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=ARRIVED_PROCESS_MAPPING[obj_data.gate_in.arrived],
                move_code=container_status,
                is_deleted=False,
                in_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_cfs_in_second_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "MTIN"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_in.in_date
        time = obj_data.gate_in.in_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=ARRIVED_PROCESS_MAPPING[obj_data.gate_in.arrived],
                move_code=container_status,
                is_deleted=False,
                in_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_cfs_in_third_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "DMG_MT"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_in.in_date
        time = obj_data.gate_in.in_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time + datetime.timedelta(minutes=15)
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=ARRIVED_PROCESS_MAPPING[obj_data.gate_in.arrived],
                move_code=container_status,
                is_deleted=False,
                in_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_cfs_in_fourth_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "REP_OUT_MT"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_in.in_date
        time = obj_data.gate_in.in_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time + datetime.timedelta(minutes=30)
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=ARRIVED_PROCESS_MAPPING[obj_data.gate_in.arrived],
                move_code=container_status,
                is_deleted=False,
                in_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_cfs_out_first_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "VAN"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_out.out_date
        time = obj_data.gate_out.out_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        # booking_no = (
        #     str(obj_data.gate_out.booking_no).split("-")[0]
        #     if obj_data.gate_out.seal_no is not None
        #     else None
        # )
        booking_no = (
            None
            if "/" in str(obj_data.gate_out.booking_no)
            or obj_data.gate_out.booking_no is None
            else str(obj_data.gate_out.booking_no).split("-")[0]
        )
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = "M" if obj_data.gate_out.seal_no is not None else None
        seal_no_1 = obj_data.gate_out.seal_no
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=DEPARTED_PROCESS_MAPPING[obj_data.gate_out.departed],
                move_code=container_status,
                is_deleted=False,
                out_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_cfs_out_second_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "REP_RET_MT"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_out.out_date
        time = obj_data.gate_out.out_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time - datetime.timedelta(minutes=15)
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=DEPARTED_PROCESS_MAPPING[obj_data.gate_out.departed],
                move_code=container_status,
                is_deleted=False,
                out_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_factory_in_first_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "MTIN"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_in.in_date
        time = obj_data.gate_in.in_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=ARRIVED_PROCESS_MAPPING[obj_data.gate_in.arrived],
                move_code=container_status,
                is_deleted=False,
                in_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_factory_in_second_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "DMG_MT"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_in.in_date
        time = obj_data.gate_in.in_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time + datetime.timedelta(minutes=15)
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=ARRIVED_PROCESS_MAPPING[obj_data.gate_in.arrived],
                move_code=container_status,
                is_deleted=False,
                in_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_factory_in_third_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "REP_OUT_MT"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_in.in_date
        time = obj_data.gate_in.in_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time + datetime.timedelta(minutes=30)
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=ARRIVED_PROCESS_MAPPING[obj_data.gate_in.arrived],
                move_code=container_status,
                is_deleted=False,
                in_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_factory_out_first_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "VAN"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_out.out_date
        time = obj_data.gate_out.out_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        # booking_no = (
        #     str(obj_data.gate_out.booking_no).split("-")[0]
        #     if obj_data.gate_out.seal_no is not None
        #     else None
        # )
        booking_no = (
            None
            if "/" in str(obj_data.gate_out.booking_no)
            or obj_data.gate_out.booking_no is None
            else str(obj_data.gate_out.booking_no).split("-")[0]
        )
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = "M" if obj_data.gate_out.seal_no is not None else None
        seal_no_1 = obj_data.gate_out.seal_no
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=DEPARTED_PROCESS_MAPPING[obj_data.gate_out.departed],
                move_code=container_status,
                is_deleted=False,
                out_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_factory_out_second_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "REP_RET_MT"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_out.out_date
        time = obj_data.gate_out.out_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time - datetime.timedelta(minutes=15)
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=DEPARTED_PROCESS_MAPPING[obj_data.gate_out.departed],
                move_code=container_status,
                is_deleted=False,
                out_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_vessel_in_first_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        edi_partner_code = obj_data.container.site.vendor_name
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "MOR"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_in.in_date
        time = obj_data.gate_in.in_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time - datetime.timedelta(minutes=60)
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["edi_partner_code"] = edi_partner_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=ARRIVED_PROCESS_MAPPING[obj_data.gate_in.arrived],
                move_code=container_status,
                is_deleted=False,
                in_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_vessel_in_second_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "MIR"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_in.in_date
        time = obj_data.gate_in.in_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=ARRIVED_PROCESS_MAPPING[obj_data.gate_in.arrived],
                move_code=container_status,
                is_deleted=False,
                in_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_vessel_in_third_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "DMG_MT"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_in.in_date
        time = obj_data.gate_in.in_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time + datetime.timedelta(minutes=15)
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=ARRIVED_PROCESS_MAPPING[obj_data.gate_in.arrived],
                move_code=container_status,
                is_deleted=False,
                in_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_vessel_in_fourth_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "REP_OUT_MT"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_in.in_date
        time = obj_data.gate_in.in_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time + datetime.timedelta(minutes=30)
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=ARRIVED_PROCESS_MAPPING[obj_data.gate_in.arrived],
                move_code=container_status,
                is_deleted=False,
                in_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_vessel_out_first_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        edi_partner_code = obj_data.container.site.vendor_name
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "MOR"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_out.out_date
        time = obj_data.gate_out.out_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        # booking_no = (
        #     str(obj_data.gate_out.booking_no).split("-")[0]
        #     if obj_data.gate_out.seal_no is not None
        #     else None
        # )
        booking_no = (
            None
            if obj_data.gate_out.seal_no is None
            or "/" in str(obj_data.gate_out.booking_no)
            or obj_data.gate_out.booking_no is None
            else str(obj_data.gate_out.booking_no).split("-")[0]
        )
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = "L" if obj_data.gate_out.seal_no is not None else None
        seal_no_1 = obj_data.gate_out.seal_no

        transporter_name = (
            obj_data.gate_out.transporter_name.name
            if obj_data.gate_out.transporter_name is not None
            else None
        )
        truck_no = obj_data.gate_out.vehicle_no
        next_event_location = obj_data.gate_out.to_location_code
        next_depot_code = obj_data.gate_out.to_depot_code

        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["edi_partner_code"] = edi_partner_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        row_dict["transporter_name"] = transporter_name
        row_dict["truck_no"] = truck_no
        row_dict["next_event_location"] = next_event_location
        row_dict["next_depot_code"] = next_depot_code

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=DEPARTED_PROCESS_MAPPING[obj_data.gate_out.departed],
                move_code=container_status,
                is_deleted=False,
                out_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_vessel_out_second_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "REP_RET_MT"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_out.out_date
        time = obj_data.gate_out.out_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time - datetime.timedelta(minutes=15)
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=DEPARTED_PROCESS_MAPPING[obj_data.gate_out.departed],
                move_code=container_status,
                is_deleted=False,
                out_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_road_in_first_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "MIR"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_in.in_date
        time = obj_data.gate_in.in_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=ARRIVED_PROCESS_MAPPING[obj_data.gate_in.arrived],
                move_code=container_status,
                is_deleted=False,
                in_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_road_in_second_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "DMG_MT"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_in.in_date
        time = obj_data.gate_in.in_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time + datetime.timedelta(minutes=15)
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=ARRIVED_PROCESS_MAPPING[obj_data.gate_in.arrived],
                move_code=container_status,
                is_deleted=False,
                in_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_road_in_third_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "REP_OUT_MT"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_in.in_date
        time = obj_data.gate_in.in_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time + datetime.timedelta(minutes=30)
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=ARRIVED_PROCESS_MAPPING[obj_data.gate_in.arrived],
                move_code=container_status,
                is_deleted=False,
                in_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_road_out_first_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        edi_partner_code = obj_data.container.site.vendor_name
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "MOR"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_out.out_date
        time = obj_data.gate_out.out_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        # booking_no = (
        #     str(obj_data.gate_out.booking_no).split("-")[0]
        #     if obj_data.gate_out.seal_no is not None
        #     else None
        # )
        booking_no = (
            None
            if obj_data.gate_out.seal_no is None
            or "/" in str(obj_data.gate_out.booking_no)
            or obj_data.gate_out.booking_no is None
            else str(obj_data.gate_out.booking_no).split("-")[0]
        )
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = "L" if obj_data.gate_out.seal_no is not None else None
        seal_no_1 = obj_data.gate_out.seal_no

        transporter_name = (
            obj_data.gate_out.transporter_name.name
            if obj_data.gate_out.transporter_name is not None
            else None
        )
        truck_no = obj_data.gate_out.vehicle_no
        next_event_location = obj_data.gate_out.road_rail_to_location_code
        next_depot_code = obj_data.gate_out.to_depot_code

        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["edi_partner_code"] = edi_partner_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        row_dict["transporter_name"] = (
            transporter_name[:35] if transporter_name is not None else transporter_name
        )
        row_dict["truck_no"] = truck_no[:35] if truck_no is not None else truck_no
        row_dict["next_event_location"] = (
            next_event_location[:5]
            if next_event_location is not None
            else next_event_location
        )
        row_dict["next_depot_code"] = (
            next_depot_code[:7] if next_depot_code is not None else next_depot_code
        )

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=DEPARTED_PROCESS_MAPPING[obj_data.gate_out.departed],
                move_code=container_status,
                is_deleted=False,
                out_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_road_out_second_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "REP_RET_MT"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_out.out_date
        time = obj_data.gate_out.out_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time - datetime.timedelta(minutes=15)
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=DEPARTED_PROCESS_MAPPING[obj_data.gate_out.departed],
                move_code=container_status,
                is_deleted=False,
                out_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_fs_return_in_first_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "RET"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_in.in_date
        time = obj_data.gate_in.in_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=ARRIVED_PROCESS_MAPPING[obj_data.gate_in.arrived],
                move_code=container_status,
                is_deleted=False,
                in_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def msc_edi_excel_fs_return_out_first_data(obj_data, row_dict, regenerate):
    try:
        tz = timezone.get_current_timezone()
        depot_name = (
            obj_data.container.site.depot_name
            if obj_data.container.site.name.lower() == "matrix"
            else obj_data.container.site.organization
        )
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "VAN"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.gate_out.out_date
        time = obj_data.gate_out.out_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        # booking_no = (
        #     str(obj_data.gate_out.booking_no).split("-")[0]
        #     if obj_data.gate_out.seal_no is not None
        #     else None
        # )
        booking_no = (
            None
            if "/" in str(obj_data.gate_out.booking_no)
            or obj_data.gate_out.booking_no is None
            else str(obj_data.gate_out.booking_no).split("-")[0]
        )
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = "M" if obj_data.gate_out.seal_no is not None else None
        seal_no_1 = obj_data.gate_out.seal_no
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        if regenerate is False:
            MscExcelEdiMoveCodeInfo.objects.get_or_create(
                date=date_time,
                site=obj_data.container.site.name,
                container_no=container_no,
                process=DEPARTED_PROCESS_MAPPING[obj_data.gate_out.departed],
                move_code=container_status,
                is_deleted=False,
                out_data_id=obj_data.pk,
            )

        return row_dict
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def create_msc_edi_excel_format_file(length, filename):
    try:
        if not os.path.exists(os.path.join(BASE_DIR, "temp/sample_msc_edi_excel/")):
            os.makedirs(os.path.join(BASE_DIR, "temp/sample_msc_edi_excel/"))
        temp_file_path = os.path.join(
            BASE_DIR, f"temp/sample_msc_edi_excel/{filename}.xlsx"
        )
        generic_workbook = xlsxwriter.Workbook(temp_file_path)
        generic_sheet = generic_workbook.add_worksheet("Data_example")
        font_format_1 = generic_workbook.add_format(
            {
                "align": "left",
                "font_name": "Calibri",
                "font_size": 11,
                "num_format": "@",
            }
        )
        font_format_2 = generic_workbook.add_format(
            {
                "align": "left",
                "font_name": "Calibri",
                "font_size": 11,
                "num_format": "0",
            }
        )
        font_format_3 = generic_workbook.add_format(
            {
                "align": "left",
                "font_name": "Calibri",
                "font_size": 11,
                "num_format": "dd/mm/yyyy hh:mm",
            }
        )
        white_font_format = generic_workbook.add_format(
            {
                "bold": 1,
                "align": "left",
                "bg_color": "#000000",
                "font_name": "Calibri",
                "font_color": "#ffffff",
                "font_size": 10,
                "num_format": "@",
            }
        )
        red_font_format = generic_workbook.add_format(
            {
                "bold": 1,
                "align": "left",
                "bg_color": "#000000",
                "font_name": "Calibri",
                "font_color": "#FF0000",
                "font_size": 10,
                "num_format": "@",
            }
        )
        red_2_font_format = generic_workbook.add_format(
            {
                "bold": 1,
                "align": "left",
                "bg_color": "#000000",
                "font_name": "Calibri",
                "font_color": "#FF0000",
                "font_size": 10,
                "num_format": "0",
            }
        )
        red_3_font_format = generic_workbook.add_format(
            {
                "bold": 1,
                "align": "left",
                "bg_color": "#000000",
                "font_name": "Calibri",
                "font_color": "#FF0000",
                "font_size": 10,
                "num_format": "m/d/yy h:mm;@",
            }
        )
        dark_green_font_format = generic_workbook.add_format(
            {
                "bold": 1,
                "align": "left",
                "bg_color": "#000000",
                "font_name": "Calibri",
                "font_color": "#70ad47",
                "font_size": 10,
                "num_format": "@",
            }
        )
        green_font_format = generic_workbook.add_format(
            {
                "bold": 1,
                "align": "left",
                "bg_color": "#000000",
                "font_name": "Calibri",
                "font_color": "#92d050",
                "font_size": 10,
                "num_format": "@",
            }
        )
        yellow_font_format = generic_workbook.add_format(
            {
                "bold": 1,
                "align": "left",
                "bg_color": "#000000",
                "font_name": "Calibri",
                "font_color": "#ffff00",
                "font_size": 10,
                "num_format": "@",
            }
        )
        dark_yellow_font_format = generic_workbook.add_format(
            {
                "bold": 1,
                "align": "left",
                "bg_color": "#000000",
                "font_name": "Calibri",
                "font_color": "#ffc000",
                "font_size": 10,
                "num_format": "@",
            }
        )
        blue_font_format = generic_workbook.add_format(
            {
                "bold": 1,
                "align": "left",
                "bg_color": "#000000",
                "font_name": "Calibri",
                "font_color": "#4472c4",
                "font_size": 10,
                "num_format": "@",
            }
        )

        generic_sheet.write("A1", "Depot Name", white_font_format)
        generic_sheet.write("B1", "Depot MSCCode", red_font_format)
        generic_sheet.write("C1", "EDI Partner Code", white_font_format)
        generic_sheet.write("D1", "Event Location ", red_font_format)
        generic_sheet.write("E1", "Container Status", red_font_format)
        generic_sheet.write("F1", "Full / Empty", red_font_format)
        generic_sheet.write("G1", "Container Number", red_font_format)
        generic_sheet.write("H1", "Move Date/Time", red_3_font_format)
        generic_sheet.write("I1", "Booking Number", dark_green_font_format)
        generic_sheet.write("J1", "BL Number", dark_green_font_format)
        generic_sheet.write("K1", "Vessel", dark_green_font_format)
        generic_sheet.write("L1", "Voyage", dark_green_font_format)
        generic_sheet.write("M1", "POL (UN Code)", yellow_font_format)
        generic_sheet.write("N1", "POD (UN Code)", yellow_font_format)
        generic_sheet.write("O1", "Cargo Weight (Kgs)", red_2_font_format)
        generic_sheet.write("P1", "Seal Type 1", blue_font_format)
        generic_sheet.write("Q1", "Seal Number 1", dark_yellow_font_format)
        generic_sheet.write("R1", "Seal Type 2", blue_font_format)
        generic_sheet.write("S1", "Seal Number 2", dark_yellow_font_format)
        generic_sheet.write("T1", "Seal Type 3", blue_font_format)
        generic_sheet.write("U1", "Seal Number 3", dark_yellow_font_format)
        generic_sheet.write("V1", "Transport Carrier (SCAC Code)", white_font_format)
        generic_sheet.write("W1", "Vehicle Registration", white_font_format)
        generic_sheet.write("X1", "Leasing Company (SCAC Code + 2)", white_font_format)
        generic_sheet.write("Y1", "Remarks", green_font_format)
        generic_sheet.write("Z1", "Leasing_Company_Code", green_font_format)
        generic_sheet.write("AA1", "Pickup_Reference", green_font_format)
        generic_sheet.write("AB1", "Lessor_Code", green_font_format)
        generic_sheet.write("AC1", "Transporter_Name", green_font_format)
        generic_sheet.write("AD1", "Truck_Number", green_font_format)
        generic_sheet.write("AE1", "Next_Event_Location", green_font_format)
        generic_sheet.write("AF1", "Next_Depot_Code", green_font_format)
        for i in range(2, (length + 3)):
            generic_sheet.write(f"A{i}", None, font_format_1)
            generic_sheet.write(f"B{i}", None, font_format_1)
            generic_sheet.write(f"C{i}", None, font_format_1)
            generic_sheet.write(f"D{i}", None, font_format_1)
            generic_sheet.write(f"E{i}", None, font_format_1)
            generic_sheet.write(f"F{i}", None, font_format_1)
            generic_sheet.write(f"G{i}", None, font_format_1)
            generic_sheet.write(f"H{i}", None, font_format_3)
            generic_sheet.write(f"I{i}", None, font_format_1)
            generic_sheet.write(f"J{i}", None, font_format_1)
            generic_sheet.write(f"K{i}", None, font_format_1)
            generic_sheet.write(f"L{i}", None, font_format_1)
            generic_sheet.write(f"M{i}", None, font_format_1)
            generic_sheet.write(f"N{i}", None, font_format_1)
            generic_sheet.write(f"O{i}", None, font_format_2)
            generic_sheet.write(f"P{i}", None, font_format_1)
            generic_sheet.write(f"Q{i}", None, font_format_1)
            generic_sheet.write(f"R{i}", None, font_format_1)
            generic_sheet.write(f"S{i}", None, font_format_1)
            generic_sheet.write(f"T{i}", None, font_format_1)
            generic_sheet.write(f"U{i}", None, font_format_1)
            generic_sheet.write(f"V{i}", None, font_format_1)
            generic_sheet.write(f"W{i}", None, font_format_1)
            generic_sheet.write(f"X{i}", None, font_format_1)
            generic_sheet.write(f"Y{i}", None, font_format_1)
            generic_sheet.write(f"Z{i}", None, font_format_1)
            generic_sheet.write(f"AA{i}", None, font_format_1)
            generic_sheet.write(f"AB{i}", None, font_format_1)
            generic_sheet.write(f"AC{i}", None, font_format_1)
            generic_sheet.write(f"AD{i}", None, font_format_1)
            generic_sheet.write(f"AE{i}", None, font_format_1)
            generic_sheet.write(f"AF{i}", None, font_format_1)
        generic_workbook.close()
        return temp_file_path
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def make_msc_edi_excel_file(context, filename):
    try:
        target = create_msc_edi_excel_format_file(len(context["depot_name"]), filename)
        # filling data in excel file

        # New Version
        book = load_workbook(target)
        with pd.ExcelWriter(
            target,
            engine="openpyxl",
            mode="a",
            if_sheet_exists="overlay",
        ) as writer:
            stock_df = pd.DataFrame(context)
            stock_df.to_excel(
                writer,
                sheet_name="Data_example",
                index=False,
                header=False,
                startrow=1,
                startcol=0,
            )

        # Old Version
        # book = load_workbook(target)
        # writer = pd.ExcelWriter(target, engine="openpyxl")
        # writer.book = book
        # writer.sheets = dict((ws.title, ws) for ws in book.worksheets)
        # billing_df = pd.DataFrame(context)
        # billing_df.to_excel(
        #     writer,
        #     sheet_name="Data_example",
        # startrow=1,
        # startcol=0,
        #     index=False,
        #     header=False,
        # )
        # writer.save()

        return target
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def get_msc_edi_excel_main_data(data_object_list, regenerate=False):
    try:
        main_row_dict = {
            "depot_name": None,
            "depot_msc_code": None,
            "edi_partner_code": None,
            "event_location": None,
            "container_status": None,
            "full_empty": None,
            "container_no": None,
            "move_date_time": None,
            "booking_no": None,
            "bl_no": None,
            "vessel": None,
            "voyage": None,
            "pol": None,
            "pod": None,
            "cargo_wt": None,
            "seal_type_1": None,
            "seal_no_1": None,
            "seal_type_2": None,
            "seal_no_2": None,
            "seal_type_3": None,
            "seal_no_3": None,
            "transport_carrier": None,
            "vehicle": None,
            "lease_company": None,
            "remarks": None,
            "lease_company_code": None,
            "pick_up_ref": None,
            "lesser_code": None,
            "transporter_name": None,
            "truck_no": None,
            "next_event_location": None,
            "next_depot_code": None,
        }
        main_dict_list = []
        eligible_data_list = []
        for each in data_object_list:
            process = detect_arrived_by(each)
            repair_type = detect_repair_type(each)
            if process == "cfs_in":
                tz = timezone.get_current_timezone()
                date = each.gate_in.in_date
                time = each.gate_in.in_time
                date_time = datetime.datetime.combine(date=date, time=time).astimezone(
                    tz
                )
                transaction_date_time = date_time + datetime.timedelta(minutes=30)
                move_date_time = transaction_date_time.astimezone(tz)
                if move_date_time < datetime.datetime.now().astimezone(tz):
                    if repair_type == "Washing":
                        second = msc_edi_excel_cfs_in_second_data(
                            obj_data=each,
                            row_dict=main_row_dict.copy(),
                            regenerate=regenerate,
                        )
                        main_dict_list.append(second)
                        eligible_data_list.append(each)
                    else:
                        first = msc_edi_excel_cfs_in_first_data(
                            obj_data=each,
                            row_dict=main_row_dict.copy(),
                            regenerate=regenerate,
                        )
                        main_dict_list.append(first)
                        second = msc_edi_excel_cfs_in_second_data(
                            obj_data=each,
                            row_dict=main_row_dict.copy(),
                            regenerate=regenerate,
                        )
                        main_dict_list.append(second)
                        third = msc_edi_excel_cfs_in_third_data(
                            obj_data=each,
                            row_dict=main_row_dict.copy(),
                            regenerate=regenerate,
                        )
                        main_dict_list.append(third)
                        fourth = msc_edi_excel_cfs_in_fourth_data(
                            obj_data=each,
                            row_dict=main_row_dict.copy(),
                            regenerate=regenerate,
                        )
                        main_dict_list.append(fourth)
                        eligible_data_list.append(each)
                else:
                    pass
            elif process == "cfs_out":
                if repair_type == "Washing":
                    first = msc_edi_excel_cfs_out_first_data(
                        obj_data=each,
                        row_dict=main_row_dict.copy(),
                        regenerate=regenerate,
                    )
                    main_dict_list.append(first)
                    eligible_data_list.append(each)
                else:
                    second = msc_edi_excel_cfs_out_second_data(
                        obj_data=each,
                        row_dict=main_row_dict.copy(),
                        regenerate=regenerate,
                    )
                    main_dict_list.append(second)
                    first = msc_edi_excel_cfs_out_first_data(
                        obj_data=each,
                        row_dict=main_row_dict.copy(),
                        regenerate=regenerate,
                    )
                    main_dict_list.append(first)
                    eligible_data_list.append(each)
            elif process == "factory_in":
                tz = timezone.get_current_timezone()
                date = each.gate_in.in_date
                time = each.gate_in.in_time
                date_time = datetime.datetime.combine(date=date, time=time).astimezone(
                    tz
                )
                transaction_date_time = date_time + datetime.timedelta(minutes=30)
                move_date_time = transaction_date_time.astimezone(tz)
                if move_date_time < datetime.datetime.now().astimezone(tz):
                    if repair_type == "Washing":
                        first = msc_edi_excel_factory_in_first_data(
                            obj_data=each,
                            row_dict=main_row_dict.copy(),
                            regenerate=regenerate,
                        )
                        main_dict_list.append(first)
                        eligible_data_list.append(each)
                    else:
                        first = msc_edi_excel_factory_in_first_data(
                            obj_data=each,
                            row_dict=main_row_dict.copy(),
                            regenerate=regenerate,
                        )
                        main_dict_list.append(first)
                        second = msc_edi_excel_factory_in_second_data(
                            obj_data=each,
                            row_dict=main_row_dict.copy(),
                            regenerate=regenerate,
                        )
                        main_dict_list.append(second)
                        third = msc_edi_excel_factory_in_third_data(
                            obj_data=each,
                            row_dict=main_row_dict.copy(),
                            regenerate=regenerate,
                        )
                        main_dict_list.append(third)
                        eligible_data_list.append(each)
                else:
                    pass
            elif process == "factory_out":
                if repair_type == "Washing":
                    first = msc_edi_excel_factory_out_first_data(
                        obj_data=each,
                        row_dict=main_row_dict.copy(),
                        regenerate=regenerate,
                    )
                    main_dict_list.append(first)
                    eligible_data_list.append(each)
                else:
                    second = msc_edi_excel_factory_out_second_data(
                        obj_data=each,
                        row_dict=main_row_dict.copy(),
                        regenerate=regenerate,
                    )
                    main_dict_list.append(second)
                    first = msc_edi_excel_factory_out_first_data(
                        obj_data=each,
                        row_dict=main_row_dict.copy(),
                        regenerate=regenerate,
                    )
                    main_dict_list.append(first)
                    eligible_data_list.append(each)
            elif process == "vessel_in":
                tz = timezone.get_current_timezone()
                date = each.gate_in.in_date
                time = each.gate_in.in_time
                date_time = datetime.datetime.combine(date=date, time=time).astimezone(
                    tz
                )
                transaction_date_time = date_time + datetime.timedelta(minutes=30)
                move_date_time = transaction_date_time.astimezone(tz)
                if move_date_time < datetime.datetime.now().astimezone(tz):
                    if repair_type == "Washing":
                        second = msc_edi_excel_vessel_in_second_data(
                            obj_data=each,
                            row_dict=main_row_dict.copy(),
                            regenerate=regenerate,
                        )
                        main_dict_list.append(second)
                        eligible_data_list.append(each)
                    else:
                        if not each.container.site.name.upper() in [
                            "GOLDEN HORN CONTAINER SERVICE MUNDRA",
                            "TUTICORIN",
                            "INGHK",
                            "INGHC",
                            "GOLDEN HORN CONTAINER SERVICES TWO KOLKATA",
                        ]:
                            first = msc_edi_excel_vessel_in_first_data(
                                obj_data=each,
                                row_dict=main_row_dict.copy(),
                                regenerate=regenerate,
                            )
                            main_dict_list.append(first)
                        second = msc_edi_excel_vessel_in_second_data(
                            obj_data=each,
                            row_dict=main_row_dict.copy(),
                            regenerate=regenerate,
                        )
                        main_dict_list.append(second)
                        third = msc_edi_excel_vessel_in_third_data(
                            obj_data=each,
                            row_dict=main_row_dict.copy(),
                            regenerate=regenerate,
                        )
                        main_dict_list.append(third)
                        fourth = msc_edi_excel_vessel_in_fourth_data(
                            obj_data=each,
                            row_dict=main_row_dict.copy(),
                            regenerate=regenerate,
                        )
                        main_dict_list.append(fourth)
                        eligible_data_list.append(each)
                else:
                    pass
            elif process == "vessel_out":
                if repair_type == "Washing":
                    first = msc_edi_excel_vessel_out_first_data(
                        obj_data=each,
                        row_dict=main_row_dict.copy(),
                        regenerate=regenerate,
                    )
                    main_dict_list.append(first)
                    eligible_data_list.append(each)
                else:
                    second = msc_edi_excel_vessel_out_second_data(
                        obj_data=each,
                        row_dict=main_row_dict.copy(),
                        regenerate=regenerate,
                    )
                    main_dict_list.append(second)
                    first = msc_edi_excel_vessel_out_first_data(
                        obj_data=each,
                        row_dict=main_row_dict.copy(),
                        regenerate=regenerate,
                    )
                    main_dict_list.append(first)
                    eligible_data_list.append(each)
            elif process == "road_in":
                tz = timezone.get_current_timezone()
                date = each.gate_in.in_date
                time = each.gate_in.in_time
                date_time = datetime.datetime.combine(date=date, time=time).astimezone(
                    tz
                )
                transaction_date_time = date_time + datetime.timedelta(minutes=30)
                move_date_time = transaction_date_time.astimezone(tz)
                if move_date_time < datetime.datetime.now().astimezone(tz):
                    if repair_type == "Washing":
                        first = msc_edi_excel_road_in_first_data(
                            obj_data=each,
                            row_dict=main_row_dict.copy(),
                            regenerate=regenerate,
                        )
                        main_dict_list.append(first)
                        eligible_data_list.append(each)
                    else:
                        first = msc_edi_excel_road_in_first_data(
                            obj_data=each,
                            row_dict=main_row_dict.copy(),
                            regenerate=regenerate,
                        )
                        main_dict_list.append(first)
                        second = msc_edi_excel_road_in_second_data(
                            obj_data=each,
                            row_dict=main_row_dict.copy(),
                            regenerate=regenerate,
                        )
                        main_dict_list.append(second)
                        third = msc_edi_excel_road_in_third_data(
                            obj_data=each,
                            row_dict=main_row_dict.copy(),
                            regenerate=regenerate,
                        )
                        main_dict_list.append(third)
                        eligible_data_list.append(each)
                else:
                    pass
            elif process == "road_out":
                if repair_type == "Washing":
                    first = msc_edi_excel_road_out_first_data(
                        obj_data=each,
                        row_dict=main_row_dict.copy(),
                        regenerate=regenerate,
                    )
                    main_dict_list.append(first)
                    eligible_data_list.append(each)
                else:
                    second = msc_edi_excel_road_out_second_data(
                        obj_data=each,
                        row_dict=main_row_dict.copy(),
                        regenerate=regenerate,
                    )
                    main_dict_list.append(second)
                    first = msc_edi_excel_road_out_first_data(
                        obj_data=each,
                        row_dict=main_row_dict.copy(),
                        regenerate=regenerate,
                    )
                    main_dict_list.append(first)
                    eligible_data_list.append(each)

            elif process == "fs_return_in":
                first = msc_edi_excel_fs_return_in_first_data(
                    obj_data=each,
                    row_dict=main_row_dict.copy(),
                    regenerate=regenerate,
                )
                main_dict_list.append(first)
                eligible_data_list.append(each)
            elif process == "fs_return_out":
                first = msc_edi_excel_fs_return_out_first_data(
                    obj_data=each,
                    row_dict=main_row_dict.copy(),
                    regenerate=regenerate,
                )
                main_dict_list.append(first)
                eligible_data_list.append(each)
            else:
                pass

        if not len(main_dict_list) == 0:
            main_data = {
                "depot_name": [each.get("depot_name") for each in main_dict_list],
                "depot_msc_code": [
                    each.get("depot_msc_code") for each in main_dict_list
                ],
                "edi_partner_code": [
                    each.get("edi_partner_code") for each in main_dict_list
                ],
                "event_location": [
                    each.get("event_location") for each in main_dict_list
                ],
                "container_status": [
                    each.get("container_status") for each in main_dict_list
                ],
                "full_empty": [each.get("full_empty") for each in main_dict_list],
                "container_no": [each.get("container_no") for each in main_dict_list],
                "move_date_time": [
                    each.get("move_date_time") for each in main_dict_list
                ],
                "booking_no": [each.get("booking_no") for each in main_dict_list],
                "bl_no": [each.get("bl_no") for each in main_dict_list],
                "vessel": [each.get("vessel") for each in main_dict_list],
                "voyage": [each.get("voyage") for each in main_dict_list],
                "pol": [each.get("pol") for each in main_dict_list],
                "pod": [each.get("pod") for each in main_dict_list],
                "cargo_wt": [each.get("cargo_wt") for each in main_dict_list],
                "seal_type_1": [each.get("seal_type_1") for each in main_dict_list],
                "seal_no_1": [each.get("seal_no_1") for each in main_dict_list],
                "seal_type_2": [each.get("seal_type_2") for each in main_dict_list],
                "seal_no_2": [each.get("seal_no_2") for each in main_dict_list],
                "seal_type_3": [each.get("seal_type_3") for each in main_dict_list],
                "seal_no_3": [each.get("seal_no_3") for each in main_dict_list],
                "transport_carrier": [
                    each.get("transport_carrier") for each in main_dict_list
                ],
                "vehicle": [each.get("vehicle") for each in main_dict_list],
                "lease_company": [each.get("lease_company") for each in main_dict_list],
                "remarks": [each.get("remarks") for each in main_dict_list],
                "lease_company_code": [
                    each.get("lease_company_code") for each in main_dict_list
                ],
                "pick_up_ref": [each.get("pick_up_ref") for each in main_dict_list],
                "lesser_code": [each.get("lesser_code") for each in main_dict_list],
                "transporter_name": [
                    each.get("transporter_name") for each in main_dict_list
                ],
                "truck_no": [each.get("truck_no") for each in main_dict_list],
                "next_event_location": [
                    each.get("next_event_location") for each in main_dict_list
                ],
                "next_depot_code": [
                    each.get("next_depot_code") for each in main_dict_list
                ],
            }
            return {"main_data": main_data, "eligible_data_list": eligible_data_list}
        return None
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def faridabad_msc_edi_excel_in_first_data(obj_data, row_dict):
    try:
        tz = timezone.get_current_timezone()
        depot_name = obj_data.container.site.organization
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "DMG_MT"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.in_date
        time = obj_data.in_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time + datetime.timedelta(minutes=15)
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        today = datetime.datetime.now().astimezone(tz)
        if transaction_date_time.astimezone(tz) < today:
            return row_dict
        else:
            return None
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def faridabad_msc_edi_excel_in_second_data(obj_data, row_dict):
    try:
        tz = timezone.get_current_timezone()
        depot_name = obj_data.container.site.organization
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "REP_OUT_MT"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.in_date
        time = obj_data.in_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time + datetime.timedelta(minutes=30)
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        today = datetime.datetime.now().astimezone(tz)
        if transaction_date_time.astimezone(tz) < today:
            return row_dict
        else:
            return None
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def faridabad_msc_edi_excel_out_first_data(obj_data, row_dict):
    try:
        tz = timezone.get_current_timezone()
        depot_name = obj_data.container.site.organization
        depot_msc_code = obj_data.container.site.depot_msc_code
        event_location = obj_data.container.site.event_location_msc_code
        move_code = "REP_RET_MT"
        container_status = CONTAINER_MOVE_CODE_STATUS[move_code]
        full_empty = "F" if move_code == "DVAN" else "E"
        container_no = obj_data.container.container_no
        date = obj_data.out_date
        time = obj_data.out_time
        date_time = datetime.datetime.combine(date=date, time=time).astimezone(tz)
        transaction_date_time = date_time - datetime.timedelta(minutes=15)
        move_date_time = transaction_date_time.astimezone(tz).strftime("%d-%m-%Y %H:%M")
        booking_no = None
        cargo_wt = obj_data.container.tare_wt if full_empty == "F" else 0
        seal_type_1 = None
        seal_no_1 = None
        # row update
        row_dict["depot_name"] = depot_name
        row_dict["depot_msc_code"] = depot_msc_code
        row_dict["event_location"] = event_location
        row_dict["container_status"] = container_status
        row_dict["full_empty"] = full_empty
        row_dict["container_no"] = container_no
        row_dict["move_date_time"] = move_date_time
        row_dict["booking_no"] = booking_no
        row_dict["cargo_wt"] = int("".join(c for c in str(cargo_wt) if c.isdigit()))
        row_dict["seal_type_1"] = seal_type_1
        row_dict["seal_no_1"] = seal_no_1

        today = datetime.datetime.now().astimezone(tz)
        if transaction_date_time.astimezone(tz) < today:
            return row_dict
        else:
            return None
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None


def get_faridabad_msc_edi_excel_main_data(data_object_list):
    try:
        main_row_dict = {
            "depot_name": None,
            "depot_msc_code": None,
            "edi_partner_code": None,
            "event_location": None,
            "container_status": None,
            "full_empty": None,
            "container_no": None,
            "move_date_time": None,
            "booking_no": None,
            "bl_no": None,
            "vessel": None,
            "voyage": None,
            "pol": None,
            "pod": None,
            "cargo_wt": None,
            "seal_type_1": None,
            "seal_no_1": None,
            "seal_type_2": None,
            "seal_no_2": None,
            "seal_type_3": None,
            "seal_no_3": None,
            "transport_carrier": None,
            "vehicle": None,
            "lease_company": None,
            "remarks": None,
            "lease_company_code": None,
            "pick_up_ref": None,
            "lesser_code": None,
            "transporter_name": None,
            "truck_no": None,
            "next_event_location": None,
            "next_depot_code": None,
        }
        main_dict_list = []
        in_pk_list = []
        out_pk_list = []
        for each in data_object_list:
            process = ""
            try:
                each.in_date
                process = "in"
            except:
                each.out_date
                process = "out"

            if process == "in":
                first = faridabad_msc_edi_excel_in_first_data(
                    obj_data=each, row_dict=main_row_dict.copy()
                )
                second = faridabad_msc_edi_excel_in_second_data(
                    obj_data=each, row_dict=main_row_dict.copy()
                )
                if first is not None and second is not None:
                    main_dict_list.append(first)
                    main_dict_list.append(second)
                    in_pk_list.append(each.pk)
            elif process == "out":
                first = faridabad_msc_edi_excel_out_first_data(
                    obj_data=each, row_dict=main_row_dict.copy()
                )
                if first is not None:
                    main_dict_list.append(first)
                    out_pk_list.append(each.pk)
            else:
                pass

        main_data = {
            "depot_name": [each.get("depot_name") for each in main_dict_list],
            "depot_msc_code": [each.get("depot_msc_code") for each in main_dict_list],
            "edi_partner_code": [
                each.get("edi_partner_code") for each in main_dict_list
            ],
            "event_location": [each.get("event_location") for each in main_dict_list],
            "container_status": [
                each.get("container_status") for each in main_dict_list
            ],
            "full_empty": [each.get("full_empty") for each in main_dict_list],
            "container_no": [each.get("container_no") for each in main_dict_list],
            "move_date_time": [each.get("move_date_time") for each in main_dict_list],
            "booking_no": [each.get("booking_no") for each in main_dict_list],
            "bl_no": [each.get("bl_no") for each in main_dict_list],
            "vessel": [each.get("vessel") for each in main_dict_list],
            "voyage": [each.get("voyage") for each in main_dict_list],
            "pol": [each.get("pol") for each in main_dict_list],
            "pod": [each.get("pod") for each in main_dict_list],
            "cargo_wt": [each.get("cargo_wt") for each in main_dict_list],
            "seal_type_1": [each.get("seal_type_1") for each in main_dict_list],
            "seal_no_1": [each.get("seal_no_1") for each in main_dict_list],
            "seal_type_2": [each.get("seal_type_2") for each in main_dict_list],
            "seal_no_2": [each.get("seal_no_2") for each in main_dict_list],
            "seal_type_3": [each.get("seal_type_3") for each in main_dict_list],
            "seal_no_3": [each.get("seal_no_3") for each in main_dict_list],
            "transport_carrier": [
                each.get("transport_carrier") for each in main_dict_list
            ],
            "vehicle": [each.get("vehicle") for each in main_dict_list],
            "lease_company": [each.get("lease_company") for each in main_dict_list],
            "remarks": [each.get("remarks") for each in main_dict_list],
            "lease_company_code": [
                each.get("lease_company_code") for each in main_dict_list
            ],
            "pick_up_ref": [each.get("pick_up_ref") for each in main_dict_list],
            "lesser_code": [each.get("lesser_code") for each in main_dict_list],
            "transporter_name": [
                each.get("transporter_name") for each in main_dict_list
            ],
            "truck_no": [each.get("truck_no") for each in main_dict_list],
            "next_event_location": [
                each.get("next_event_location") for each in main_dict_list
            ],
            "next_depot_code": [each.get("next_depot_code") for each in main_dict_list],
        }
        return main_data, in_pk_list, out_pk_list
    except:
        error_log = logging.getLogger("error_log")
        error_log.error(traceback.format_exc())
        return None
